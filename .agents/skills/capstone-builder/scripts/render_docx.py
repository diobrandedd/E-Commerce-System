"""Render the capstone-builder Markdown subset as an academic DOCX."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_inline(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(0.5)

    for style_name, size, bold in (
        ("Capstone Title", 14, True),
        ("Capstone Heading 1", 14, True),
        ("Capstone Heading 2", 12, True),
        ("Capstone Heading 3", 12, True),
    ):
        style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    document.styles["Capstone Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.styles["Capstone Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if "Table Text" not in [s.name for s in document.styles]:
        style = document.styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Times New Roman"
        style.font.size = Pt(9)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.first_line_indent = Inches(0)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["Table Text"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, value.strip())
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph()


def add_image(document: Document, source: str, alt_text: str) -> None:
    image_path = Path(source)
    if not image_path.is_absolute():
        raise ValueError(f"Image path must be absolute: {source}")
    if not image_path.exists():
        raise FileNotFoundError(f"Image not found: {image_path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Inches(0)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.45))
    if alt_text:
        picture = run._r.xpath(".//pic:cNvPr")
        if picture:
            picture[0].set("descr", alt_text)


def render(markdown_path: Path, output_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)

    title_page = True
    paragraph_buffer: list[str] = []
    table_rows: list[list[str]] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(part.strip() for part in paragraph_buffer).strip()
        paragraph = document.add_paragraph()
        add_inline(paragraph, text)
        if title_page:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.space_after = Pt(12)
        paragraph_buffer = []

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            add_table(document, table_rows)
            table_rows = []

    for raw_line in lines:
        line = raw_line.rstrip()

        if line == r"\newpage":
            flush_paragraph()
            flush_table()
            paragraph = document.add_paragraph()
            paragraph.add_run().add_break(WD_BREAK.PAGE)
            title_page = False
            continue

        if line.startswith("|") and line.endswith("|"):
            flush_paragraph()
            values = [value.strip() for value in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", value) for value in values):
                continue
            table_rows.append(values)
            continue
        flush_table()

        if not line.strip():
            flush_paragraph()
            continue

        image_match = re.fullmatch(r"!\[(.*?)\]\((.+?)\)", line)
        if image_match:
            flush_paragraph()
            add_image(document, image_match.group(2), image_match.group(1))
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading_match:
            flush_paragraph()
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            style_name = "Capstone Title" if title_page and level == 1 else f"Capstone Heading {level}"
            paragraph = document.add_paragraph(style=style_name)
            add_inline(paragraph, text)
            if title_page:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        bullet_match = re.match(r"^(\s*)-\s+(.+)$", line)
        if bullet_match:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.first_line_indent = Inches(0)
            add_inline(paragraph, bullet_match.group(2))
            continue

        numbered_match = re.match(r"^(\s*)\d+\.\s+(.+)$", line)
        if numbered_match:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.first_line_indent = Inches(0)
            add_inline(paragraph, numbered_match.group(2))
            continue

        paragraph_buffer.append(line)

    flush_paragraph()
    flush_table()

    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run()
        field_begin = OxmlElement("w:fldChar")
        field_begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = "PAGE"
        field_end = OxmlElement("w:fldChar")
        field_end.set(qn("w:fldCharType"), "end")
        run._r.extend([field_begin, instruction, field_end])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: render_docx.py INPUT.md OUTPUT.docx")
    render(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())


if __name__ == "__main__":
    main()
