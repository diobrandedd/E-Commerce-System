#!/usr/bin/env python3
"""Courier SOA scan + tracking/waybill highlight export (J&T / SPX)."""

from __future__ import annotations

import json
import re
import sys
import warnings
from decimal import Decimal, InvalidOperation
from typing import Any

import openpyxl
from openpyxl.styles import PatternFill

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

JNT_HEADER_ALIASES = ["waybill number", "waybill no", "waybill no.", "waybill #"]
SPX_HEADER_ALIASES = ["tracking number", "tracking no", "tracking no.", "tracking #"]


def cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_header(text: str) -> str:
    return re.sub(r"\s+", " ", cell_str(text).lower().replace("(*)", "").replace("\n", " "))


def normalize_tracking(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value).upper()
    if isinstance(value, float):
        if value != value:  # NaN
            return None
        if abs(value) >= 1e15 or float(value).is_integer():
            try:
                return str(int(Decimal(str(value)).to_integral_value())).upper()
            except (InvalidOperation, OverflowError, ValueError):
                pass
        text = format(value, ".0f") if float(value).is_integer() else str(value)
        return text.strip().upper() or None
    text = cell_str(value)
    if not text:
        return None
    if re.match(r"^\d+\.0$", text):
        text = text[:-2]
    if re.match(r"^\d+\.?\d*[eE][+]?\d+$", text):
        try:
            return str(int(Decimal(text).to_integral_value())).upper()
        except (InvalidOperation, OverflowError, ValueError):
            pass
    return text.upper()


def find_tracking_header(
    ws: openpyxl.worksheet.worksheet.Worksheet, max_scan: int = 20
) -> tuple[str, int, int]:
    """Return (format_key, header_row, tracking_col). format_key is jnt|spx."""
    max_row = min(max_scan, ws.max_row or max_scan)
    max_col = min(ws.max_column or 1, 30)
    jnt_hit: tuple[int, int] | None = None
    spx_hit: tuple[int, int] | None = None
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            norm = normalize_header(cell_str(ws.cell(row=row, column=col).value))
            if not norm:
                continue
            if jnt_hit is None and any(norm == a or a in norm for a in JNT_HEADER_ALIASES):
                jnt_hit = (row, col)
            if spx_hit is None and any(norm == a or a in norm for a in SPX_HEADER_ALIASES):
                spx_hit = (row, col)
        if jnt_hit is not None or spx_hit is not None:
            break
    # Prefer explicit waybill (J&T) when both somehow present; else SPX tracking.
    if jnt_hit is not None and spx_hit is None:
        return "jnt", jnt_hit[0], jnt_hit[1]
    if spx_hit is not None and jnt_hit is None:
        return "spx", spx_hit[0], spx_hit[1]
    if jnt_hit is not None and spx_hit is not None:
        # Same row: prefer waybill if header is waybill; else if "tracking number" only → spx
        if jnt_hit[0] <= spx_hit[0]:
            return "jnt", jnt_hit[0], jnt_hit[1]
        return "spx", spx_hit[0], spx_hit[1]
    raise ValueError(
        "Could not detect J&T (Waybill Number) or SPX (Tracking Number) column in workbook."
    )


def pick_sheet(wb: openpyxl.Workbook) -> openpyxl.worksheet.worksheet.Worksheet:
    if wb.active is not None:
        return wb.active
    return wb[wb.sheetnames[0]]


def scan_workbook(path: str) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = pick_sheet(wb)
    fmt, header_row, track_col = find_tracking_header(ws)
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    max_row = ws.max_row or header_row
    for row in range(header_row + 1, max_row + 1):
        raw = ws.cell(row=row, column=track_col).value
        tracking = normalize_tracking(raw)
        if tracking is None:
            continue
        rows.append(
            {
                "row": row,
                "col": track_col,
                "tracking": tracking,
                "tracking_raw": cell_str(raw),
                "duplicate": tracking in seen,
            }
        )
        seen.add(tracking)
    return {
        "sheet": ws.title,
        "format": fmt,
        "header_row": header_row,
        "track_col": track_col,
        "rows": rows,
    }


def highlight_workbook(src_path: str, dst_path: str, matches_path: str) -> dict[str, Any]:
    with open(matches_path, encoding="utf-8") as fh:
        payload = json.load(fh)
    row_matches: dict[str, str] = {}
    for entry in payload.get("rows", []):
        if not isinstance(entry, dict):
            continue
        row = int(entry.get("row", 0))
        status = str(entry.get("status", "")).lower()
        if row > 0 and status:
            row_matches[str(row)] = status

    wb = openpyxl.load_workbook(src_path)
    ws = pick_sheet(wb)
    header_row = int(payload.get("header_row", 1))
    track_col = int(payload.get("track_col", 1))
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    for row_key, status in row_matches.items():
        row = int(row_key)
        if row <= header_row:
            continue
        cell = ws.cell(row=row, column=track_col)
        if status == "matched":
            cell.fill = green
        elif status in ("missing", "extra"):
            cell.fill = red

    wb.save(dst_path)
    return {"ok": True, "output": dst_path}


def main() -> None:
    if len(sys.argv) < 3:
        print(json.dumps({"error": "Usage: courier_soa_workbook.py <scan|highlight> <workbook> [args...]"}))
        sys.exit(1)
    mode = sys.argv[1].lower()
    path = sys.argv[2]
    try:
        if mode == "scan":
            result = scan_workbook(path)
        elif mode == "highlight":
            if len(sys.argv) < 5:
                raise ValueError("highlight requires destination and matches JSON paths.")
            result = highlight_workbook(path, sys.argv[3], sys.argv[4])
        else:
            raise ValueError("Unknown mode: " + mode)
        print(json.dumps(result, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False))
        sys.exit(1)


if __name__ == "__main__":
    main()
