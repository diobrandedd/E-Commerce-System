#!/usr/bin/env python3
"""Parse historical order Excels: Snappy (Tracking number) or partner iPYX (Order No./Name/Item/s)."""

from __future__ import annotations

import json
import re
import sys
import warnings
from decimal import Decimal, InvalidOperation
from typing import Any

import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

SNAPPY_REQUIRED = {
    "tracking": ["tracking number", "tracking no", "tracking no.", "waybill", "waybill number"],
    "order_id": ["order id", "order no", "order no.", "order number"],
    "recipient": ["recipient name", "customer", "receiver"],
    "product": ["product name", "product"],
}

IPYX_ORDER_ALIASES = ["order no.", "order no", "order number", "order #"]
IPYX_NAME_ALIASES = ["name", "customer", "receiver", "recipient name"]
IPYX_ITEMS_ALIASES = ["item/s", "items", "item", "remarks"]
IPYX_COURIER_ALIASES = ["courier"]
IPYX_TRACKING_ALIASES = ["tracking number", "tracking no", "tracking no.", "waybill", "waybill number"]
IPYX_NOTE_ALIASES = ["note", "notes", "remark", "remarks"]


def cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_header(text: str) -> str:
    return re.sub(r"\s+", " ", cell_str(text).lower().replace("(*)", ""))


def normalize_id(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value != value:
            return ""
        try:
            # TikTok/marketplace IDs often arrive as scientific floats.
            return format(Decimal(str(value)).to_integral_value(), "f").split(".")[0]
        except (InvalidOperation, OverflowError, ValueError):
            return cell_str(value)
    text = cell_str(value).strip().strip("'\"")
    if re.match(r"^\d+\.0$", text):
        text = text[:-2]
    if re.match(r"^\d+\.?\d*[eE][+]?\d+$", text):
        try:
            return format(Decimal(text).to_integral_value(), "f").split(".")[0]
        except (InvalidOperation, OverflowError, ValueError):
            pass
    return text


def looks_precision_lost_id(text: str) -> bool:
    """True when an ID looks like Excel float corruption (many trailing zeros)."""
    if not text or not text.isdigit():
        return False
    # e.g. 820148000000 / 1097810000000000 — real Lazada TNs/order nos rarely end this way.
    return bool(re.search(r"0{5,}$", text))


def normalize_qty(value: Any) -> int:
    if value is None or value == "":
        return 1
    try:
        q = int(Decimal(str(value)))
        return q if q > 0 else 1
    except (InvalidOperation, OverflowError, ValueError):
        return 1


def find_cols(header_cells: list[str], required: dict[str, list[str]]) -> dict[str, int]:
    norms = [normalize_header(h) for h in header_cells]
    found: dict[str, int] = {}
    for key, aliases in required.items():
        for i, norm in enumerate(norms):
            if not norm:
                continue
            if any(norm == a or a in norm for a in aliases):
                found[key] = i
                break
    return found


def pick_sheet(wb: openpyxl.Workbook, sheet_name: str | None = None):
    if sheet_name:
        wanted = sheet_name.strip().lower()
        for name in wb.sheetnames:
            if name.strip().lower() == wanted:
                return wb[name]
        # soft match contains
        for name in wb.sheetnames:
            if wanted in name.strip().lower() or name.strip().lower() in wanted:
                return wb[name]
        raise ValueError(
            f"Sheet not found: {sheet_name}. Available: " + ", ".join(wb.sheetnames)
        )
    if wb.active is not None:
        return wb.active
    return wb[wb.sheetnames[0]]


def is_ipyx_header(cells: list[Any]) -> dict[str, int] | None:
    norms = [normalize_header(cell_str(c)) for c in cells]
    found: dict[str, int] = {}
    for i, norm in enumerate(norms):
        if not norm:
            continue
        if "order_no" not in found and any(norm == a or norm.startswith(a) for a in IPYX_ORDER_ALIASES):
            found["order_no"] = i
        if "name" not in found and any(norm == a for a in IPYX_NAME_ALIASES):
            found["name"] = i
        if "items" not in found and any(norm == a or a in norm for a in IPYX_ITEMS_ALIASES):
            found["items"] = i
        if "courier" not in found and any(norm == a for a in IPYX_COURIER_ALIASES):
            found["courier"] = i
        if "tracking" not in found and any(norm == a or a in norm for a in IPYX_TRACKING_ALIASES):
            found["tracking"] = i
        if "note" not in found and any(norm == a for a in IPYX_NOTE_ALIASES):
            # Don't steal Item/s remarks column if already bound as items.
            if "items" in found and found["items"] == i:
                continue
            found["note"] = i
    if "order_no" in found and "name" in found and "items" in found:
        return found
    return None


def is_section_noise(order_no: str, name: str) -> bool:
    low = order_no.lower()
    if low in (
        "ipyx",
        "-",
        "order no.",
        "order no",
        "order number",
        "pending",
        "new",
        "additional",
        "fwellness ph",
    ):
        return True
    if low.startswith("order date") or low.startswith("pickup date"):
        return True
    if low.startswith("sir "):
        return True
    if name.lower() in ("name", "customer"):
        return True
    return False


def coerce_excel_date(value: Any, default_year: int | None = None) -> str | None:
    """Convert Excel date/datetime/string into YYYY-MM-DD."""
    from datetime import datetime, date

    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel serial date
        try:
            from openpyxl.utils.datetime import from_excel

            dt = from_excel(value)
            if isinstance(dt, datetime):
                return dt.date().isoformat()
            if isinstance(dt, date):
                return dt.isoformat()
        except Exception:
            pass
    text = cell_str(value)
    if not text:
        return None
    # "2026-04-04 00:00:00"
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%b %d, %Y",
        "%B %d, %Y",
        "%b %d %Y",
        "%B %d %Y",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    if default_year is not None:
        raw = re.sub(r"([A-Za-z])(\d)", r"\1 \2", text)
        for fmt in ("%b %d", "%B %d", "%m/%d"):
            try:
                return datetime.strptime(f"{raw} {default_year}", f"{fmt} %Y").strftime("%Y-%m-%d")
            except ValueError:
                continue
    return None


def parse_pickup_date_from_cells(cells: list[Any], default_year: int | None = None) -> str | None:
    """
    Return YYYY-MM-DD from:
      - 'Pickup Date: Jan 2, 2026' in one cell
      - 'Pickup Date:' in one cell + date/datetime in the next cell (Joven layout)
    """
    for i, cell in enumerate(cells):
        text = cell_str(cell)
        if not text:
            continue
        m = re.search(r"pickup\s*date\s*:?\s*(.*)$", text, flags=re.I)
        if not m:
            continue
        rest = m.group(1).strip()
        if rest:
            parsed = coerce_excel_date(rest, default_year)
            if parsed:
                return parsed
            # Inline "Pickup Date: Mar 6"
            raw = re.sub(r"\s+", " ", rest)
            raw = re.sub(r"([A-Za-z])(\d)", r"\1 \2", raw)
            parsed = coerce_excel_date(raw, default_year)
            if parsed:
                return parsed
        # Adjacent cell holds the date (Joven / FWELLNESS layout).
        if i + 1 < len(cells):
            parsed = coerce_excel_date(cells[i + 1], default_year)
            if parsed:
                return parsed
    return None


def year_hint_from_sheet(sheet_title: str) -> int | None:
    m = re.search(r"(20\d{2})", sheet_title)
    if m:
        return int(m.group(1))
    return None


def parse_partner_ipyx(ws, sheet_title: str) -> dict[str, Any]:
    orders: list[dict[str, Any]] = []
    cols: dict[str, int] | None = None
    current_ship_date: str | None = None
    default_year = year_hint_from_sheet(sheet_title)
    excel_row = 0
    for row in ws.iter_rows(values_only=True):
        excel_row += 1
        if row is None:
            continue
        cells = list(row)

        pickup = parse_pickup_date_from_cells(cells, default_year)
        if pickup is not None:
            current_ship_date = pickup
            continue

        header = is_ipyx_header(cells)
        if header is not None:
            cols = header
            continue
        if cols is None:
            continue

        def get(key: str) -> Any:
            idx = cols.get(key) if cols else None
            if idx is None or idx >= len(cells):
                return None
            return cells[idx]

        order_no = normalize_id(get("order_no"))
        recipient = cell_str(get("name"))
        items = cell_str(get("items")).strip()
        courier = cell_str(get("courier")) if "courier" in cols else ""
        note = cell_str(get("note")) if "note" in cols else ""
        tracking = ""
        if "tracking" in cols:
            # Column H "Tracking Number" (Jun/Jul) — prefer as parcel key when present.
            tracking = re.sub(r"\s+", "", normalize_id(get("tracking")))
            if looks_precision_lost_id(tracking):
                tracking = ""

        if order_no == "" and recipient == "" and items == "":
            continue
        if is_section_noise(order_no, recipient):
            continue
        if order_no == "" or recipient == "" or items == "":
            continue
        # Skip only when Order No. is float-corrupted AND no usable waybill.
        if looks_precision_lost_id(order_no) and tracking == "":
            continue
        if note and re.search(r"\bcancel+ed\b", note, flags=re.I):
            continue

        orders.append(
            {
                "excel_row": excel_row,
                "tracking_number": tracking,
                "order_no": order_no,
                "recipient_name": recipient,
                "product_name": items,
                "quantity": 1,
                "items_raw": items,
                "courier": courier,
                "ship_date": current_ship_date,
                "sheet_status": note,
            }
        )

    if orders == []:
        raise ValueError(
            "No partner iPYX order rows found. Expected repeating headers: Order No., Name, Item/s."
        )

    return {
        "format": "partner_ipyx",
        "sheet": sheet_title,
        "total_rows": len(orders),
        "orders": orders,
        "available_sheets": [],
    }


def parse_snappy(ws, sheet_title: str) -> dict[str, Any]:
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise ValueError("Workbook is empty.")
    header_cells = [cell_str(c) for c in header]
    cols = find_cols(header_cells, SNAPPY_REQUIRED)
    # Prefer Recipient name over Customer
    norms = [normalize_header(h) for h in header_cells]
    for i, norm in enumerate(norms):
        if norm == "recipient name":
            cols["recipient"] = i
            break
    for i, norm in enumerate(norms):
        if norm in ("quantity", "qty"):
            cols["quantity"] = i
            break
    for i, norm in enumerate(norms):
        if norm == "status":
            cols["status"] = i
            break

    missing = [k for k in ("tracking", "order_id", "recipient", "product") if k not in cols]
    if missing:
        raise ValueError("snappy_missing:" + ",".join(missing))

    orders: list[dict[str, Any]] = []
    excel_row = 1
    for row in rows_iter:
        excel_row += 1
        if row is None:
            continue
        cells = list(row)

        def get(key: str) -> Any:
            idx = cols.get(key)
            if idx is None or idx >= len(cells):
                return None
            return cells[idx]

        tn = re.sub(r"\s+", "", cell_str(get("tracking")))
        order_no = normalize_id(get("order_id"))
        recipient = cell_str(get("recipient"))
        product = cell_str(get("product")).rstrip(":").strip()
        qty = normalize_qty(get("quantity")) if "quantity" in cols else 1
        status = cell_str(get("status")) if "status" in cols else ""

        if tn == "" and order_no == "" and recipient == "" and product == "":
            continue

        items = f"{qty} x {product}" if product else ""
        orders.append(
            {
                "excel_row": excel_row,
                "tracking_number": tn,
                "order_no": order_no,
                "recipient_name": recipient,
                "product_name": product,
                "quantity": qty,
                "items_raw": items,
                "courier": "",
                "sheet_status": status,
            }
        )

    return {
        "format": "snappy_shipped",
        "sheet": sheet_title,
        "total_rows": len(orders),
        "orders": orders,
        "available_sheets": [],
    }


def parse_workbook(path: str, sheet_name: str | None = None) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    available = list(wb.sheetnames)
    ws = pick_sheet(wb, sheet_name)
    title = ws.title

    # Try snappy first only when first row looks like snappy headers.
    first_rows: list[list[Any]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        first_rows.append(list(row) if row else [])
        if i >= 8:
            break

    # reopen for full parse (read_only iterator consumed)
    wb.close()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = pick_sheet(wb, sheet_name)

    snappy_look = False
    if first_rows:
        cols = find_cols([cell_str(c) for c in first_rows[0]], SNAPPY_REQUIRED)
        if "tracking" in cols and "product" in cols:
            snappy_look = True

    if snappy_look:
        try:
            data = parse_snappy(ws, title)
            data["available_sheets"] = available
            return data
        except ValueError as e:
            if not str(e).startswith("snappy_missing:"):
                raise

    # Partner iPYX: scan for Order No. / Name / Item/s headers anywhere in sheet.
    wb.close()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = pick_sheet(wb, sheet_name)
    data = parse_partner_ipyx(ws, title)
    data["available_sheets"] = available
    return data


def main() -> None:
    if len(sys.argv) < 3 or sys.argv[1] != "parse":
        print(
            json.dumps(
                {
                    "error": "Usage: historical_orders_workbook.py parse <xlsx> [sheet_name]"
                }
            )
        )
        sys.exit(1)
    path = sys.argv[2]
    sheet = sys.argv[3] if len(sys.argv) > 3 else None
    try:
        print(json.dumps(parse_workbook(path, sheet), ensure_ascii=False))
    except Exception as e:
        print(json.dumps({"error": str(e)}))
        sys.exit(1)


if __name__ == "__main__":
    main()
