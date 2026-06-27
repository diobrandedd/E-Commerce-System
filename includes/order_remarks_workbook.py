#!/usr/bin/env python3
"""Detect courier Excel layout, read rows, write remarks column only."""

from __future__ import annotations

import json
import re
import sys
import warnings
from typing import Any

import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

PLANNING_TABS = {"sheet1", "sheet2", "sheet3", "sheet4"}


def cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", cell_str(text).lower())


def sheet_order(wb: openpyxl.Workbook) -> list[str]:
    names = list(wb.sheetnames)
    ordered: list[str] = []
    if "sheet" in names:
        ordered.append("sheet")
    active = wb.active.title if wb.active else None
    if active and active not in ordered:
        ordered.append(active)
    for name in names:
        if name not in ordered:
            ordered.append(name)
    return ordered


def scan_branding(ws: openpyxl.worksheet.worksheet.Worksheet, max_row: int = 30) -> str:
    parts: list[str] = []
    for row in range(1, min(max_row, ws.max_row or 0) + 1):
        for col in range(1, min(ws.max_column or 1, 24) + 1):
            val = normalize(ws.cell(row=row, column=col).value)
            if val:
                parts.append(val)
    return " ".join(parts)


def row_values(ws: openpyxl.worksheet.worksheet.Worksheet, row: int) -> list[str]:
    values: list[str] = []
    max_col = ws.max_column or 1
    for col in range(1, max_col + 1):
        values.append(cell_str(ws.cell(row=row, column=col).value))
    return values


def _norm_header(text: str) -> str:
    return normalize(text).replace("(*)", "").strip()


def _find_col_by_aliases(header_cells: list[str], aliases: list[str]) -> int | None:
    """Match header cell by normalized label (exact, then contains)."""
    norms = [_norm_header(c) for c in header_cells]
    for alias in aliases:
        an = _norm_header(alias)
        if not an:
            continue
        for idx, val in enumerate(norms):
            if val and val == an:
                return idx
    for alias in aliases:
        an = _norm_header(alias)
        if not an:
            continue
        for idx, val in enumerate(norms):
            if val and (an in val or val in an):
                return idx
    return None


def resolve_columns(
    header_cells: list[str],
    profile: dict[str, Any],
) -> dict[str, int]:
    defaults = profile.get("default_columns") or {}

    recv_col = _find_col_by_aliases(
        header_cells, profile.get("receiver_header_aliases", [])
    )
    phone_col = _find_col_by_aliases(
        header_cells, profile.get("phone_header_aliases", [])
    )
    addr_col = _find_col_by_aliases(
        header_cells, profile.get("address_header_aliases", [])
    )
    prov_col = _find_col_by_aliases(
        header_cells, profile.get("province_header_aliases", [])
    )
    city_col = _find_col_by_aliases(
        header_cells, profile.get("city_header_aliases", [])
    )
    region_col = _find_col_by_aliases(
        header_cells, profile.get("region_header_aliases", [])
    )
    rem_col = _find_col_by_aliases(
        header_cells, profile.get("remarks_header_aliases", [])
    )
    item_col = _find_col_by_aliases(
        header_cells, profile.get("item_header_aliases", [])
    )
    track_col = _find_col_by_aliases(
        header_cells, profile.get("tracking_header_aliases", [])
    )
    ref_col = _find_col_by_aliases(
        header_cells,
        profile.get("ref_header_aliases", [])
        or profile.get("parcel_ref_header_aliases", []),
    )
    parcel_col = _find_col_by_aliases(
        header_cells, profile.get("parcel_ref_header_aliases", [])
    )
    if ref_col is None and parcel_col is not None:
        ref_col = parcel_col

    if recv_col is None:
        recv_col = int(defaults.get("receiver", 0))
    if rem_col is None and item_col is not None:
        rem_col = item_col
    if rem_col is None:
        rem_col = int(defaults.get("remarks", 12))
    if phone_col is None and defaults.get("phone") is not None:
        phone_col = int(defaults["phone"])
    if addr_col is None and defaults.get("address") is not None:
        addr_col = int(defaults["address"])
    if city_col is None and defaults.get("city") is not None:
        city_col = int(defaults["city"])
    if prov_col is None and defaults.get("province") is not None:
        prov_col = int(defaults["province"])
    if region_col is None and defaults.get("region") is not None:
        region_col = int(defaults["region"])

    cols: dict[str, int] = {"receiver": recv_col, "phone": phone_col, "remarks": rem_col}
    if item_col is not None:
        cols["item"] = item_col
    if track_col is not None:
        cols["tracking"] = track_col
    if ref_col is not None:
        cols["ref"] = ref_col
        cols["parcel_ref"] = ref_col
    if addr_col is not None:
        cols["address"] = addr_col
    if city_col is not None:
        cols["city"] = city_col
    if prov_col is not None:
        cols["province"] = prov_col
    if region_col is not None:
        cols["region"] = region_col
    return cols


def find_header_row(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    profile: dict[str, Any],
    max_scan: int = 45,
) -> tuple[int | None, dict[str, int] | None]:
    markers = profile.get("header_markers") or []
    for row in range(1, min(max_scan, ws.max_row or 0) + 1):
        cells = row_values(ws, row)
        norms = [normalize(c) for c in cells]

        for marker in markers:
            mn = normalize(marker)
            for val in norms:
                if mn in val or val == mn:
                    cols = resolve_columns(cells, profile)
                    return row, cols

        recv_col = _find_col_by_aliases(cells, profile.get("receiver_header_aliases", []))
        rem_col = _find_col_by_aliases(cells, profile.get("remarks_header_aliases", []))
        item_col = _find_col_by_aliases(cells, profile.get("item_header_aliases", []))
        if recv_col is not None and (rem_col is not None or item_col is not None):
            return row, resolve_columns(cells, profile)

    return None, None


def score_sheet(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    profile: dict[str, Any],
    branding: str,
) -> tuple[int, int | None, dict[str, int] | None]:
    score = 0
    for token in profile.get("branding") or []:
        if normalize(token) in branding:
            score += 12
    header_row, cols = find_header_row(ws, profile)
    if header_row is not None and cols is not None:
        score += 25
        for marker in profile.get("header_markers") or []:
            if normalize(marker):
                score += 4
        if cols.get("ref") is not None or cols.get("parcel_ref") is not None:
            score += 8
        if profile.get("profile_kind") == "order_list" and cols.get("tracking") is None:
            score += 10
        if profile.get("profile_kind") == "tracking_export" and cols.get("tracking") is not None:
            score += 10
    return score, header_row, cols


def profiles_for_mode(profiles: list[dict[str, Any]], mode: str) -> list[dict[str, Any]]:
    if mode == "tracking_map":
        tracking = [
            p
            for p in profiles
            if p.get("profile_kind") == "tracking_export"
            or p.get("id") in ("jnt_tracking", "spx", "jnt")
        ]
        return tracking if tracking else profiles
    # Arranger / write — skip tracking-only exports when order-list exists
    order_list = [p for p in profiles if p.get("profile_kind") == "order_list"]
    legacy = [p for p in profiles if p.get("profile_kind") not in ("tracking_export",)]
    if order_list:
        return order_list + [p for p in legacy if p.get("profile_kind") != "order_list"]
    return legacy if legacy else profiles


def detect_layout(wb: openpyxl.Workbook, profiles: list[dict[str, Any]], override: str | None, mode: str = "detect") -> dict[str, Any] | None:
    candidates = profiles_for_mode(profiles, mode)
    if override:
        candidates = [p for p in profiles if p.get("id") == override]
        if not candidates:
            return None

    best: dict[str, Any] | None = None

    for sheet_name in sheet_order(wb):
        if normalize(sheet_name) in PLANNING_TABS:
            ws_probe = wb[sheet_name]
            found_header = False
            for probe_profile in profiles:
                hr_probe, _cols_probe = find_header_row(ws_probe, probe_profile)
                if hr_probe is not None:
                    found_header = True
                    break
            if not found_header:
                continue

        ws = wb[sheet_name]
        branding = scan_branding(ws)

        for profile in candidates:
            score, header_row, cols = score_sheet(ws, profile, branding)
            if header_row is None or cols is None:
                continue
            if best is None or score > best["score"]:
                rem_idx = cols["remarks"]
                rem_letter = openpyxl.utils.get_column_letter(rem_idx + 1)
                best = {
                    "score": score,
                    "profileId": profile["id"],
                    "profileLabel": profile.get("label", profile["id"]),
                    "sheetName": sheet_name,
                    "headerRow": header_row,
                    "columns": cols,
                    "remarksColumnLetter": rem_letter,
                    "brandingSample": branding[:240],
                }

    return best


def profile_by_id(profiles: list[dict[str, Any]], profile_id: str) -> dict[str, Any] | None:
    for profile in profiles:
        if profile.get("id") == profile_id:
            return profile
    return None


def read_orders(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    layout: dict[str, Any],
    profile: dict[str, Any],
) -> list[dict[str, Any]]:
    header_row = int(layout["headerRow"])
    cols = layout["columns"]
    skip = set(profile.get("skip_receiver_values") or [])

    orders: list[dict[str, Any]] = []
    for row_num in range(header_row + 1, (ws.max_row or 0) + 1):
        recv = cell_str(ws.cell(row=row_num, column=cols["receiver"] + 1).value)
        if not recv or recv in skip:
            continue
        phone_col_idx = cols.get("phone")
        phone = (
            cell_str(ws.cell(row=row_num, column=phone_col_idx + 1).value)
            if phone_col_idx is not None
            else ""
        )
        remarks_raw = cell_str(ws.cell(row=row_num, column=cols["remarks"] + 1).value)
        item_raw = ""
        if "item" in cols and cols["item"] != cols["remarks"]:
            item_raw = cell_str(ws.cell(row=row_num, column=cols["item"] + 1).value)
        item_text = remarks_raw if remarks_raw else item_raw
        if remarks_raw and item_raw and len(item_raw) > len(remarks_raw):
            if not re.search(r"textarea_input|radio_input", remarks_raw, re.I):
                item_text = item_raw
        tracking = ""
        if "tracking" in cols:
            tracking = cell_str(ws.cell(row=row_num, column=cols["tracking"] + 1).value)
        entry: dict[str, Any] = {
            "idx": len(orders),
            "rowNum": row_num,
            "receiver": recv,
            "phone": phone,
            "remarks": item_text,
            "itemText": item_text,
            "tracking": tracking,
        }
        if "address" in cols:
            entry["address"] = cell_str(
                ws.cell(row=row_num, column=cols["address"] + 1).value
            )
        if "city" in cols:
            entry["city"] = cell_str(ws.cell(row=row_num, column=cols["city"] + 1).value)
        if "province" in cols:
            entry["province"] = cell_str(
                ws.cell(row=row_num, column=cols["province"] + 1).value
            )
        if "region" in cols:
            entry["region"] = cell_str(
                ws.cell(row=row_num, column=cols["region"] + 1).value
            )
        ref_idx = cols.get("parcel_ref")
        if ref_idx is None:
            ref_idx = cols.get("ref")
        if ref_idx is not None:
            entry["parcel_ref"] = cell_str(
                ws.cell(row=row_num, column=ref_idx + 1).value
            )
        orders.append(entry)
    return orders


def normalize_ref_value(value: str, profile: dict[str, Any]) -> str:
    ref = cell_str(value).upper()
    if profile.get("ref_normalize") == "strip_pke" and ref.startswith("PKE"):
        ref = ref[3:]
    return ref


def read_tracking_refs(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    layout: dict[str, Any],
    profile: dict[str, Any],
) -> list[dict[str, Any]]:
    header_row = int(layout["headerRow"])
    cols = layout["columns"]
    skip = set(profile.get("skip_receiver_values") or [])
    ref_idx = cols.get("ref")
    if ref_idx is None:
        ref_idx = cols.get("parcel_ref")
    track_idx = cols.get("tracking")
    if ref_idx is None or track_idx is None:
        return []

    refs: list[dict[str, Any]] = []
    for row_num in range(header_row + 1, (ws.max_row or 0) + 1):
        raw_ref = cell_str(ws.cell(row=row_num, column=ref_idx + 1).value)
        if not raw_ref or raw_ref in skip:
            continue
        tracking = cell_str(ws.cell(row=row_num, column=track_idx + 1).value)
        if not tracking:
            continue
        recv = ""
        if "receiver" in cols:
            recv = cell_str(ws.cell(row=row_num, column=cols["receiver"] + 1).value)
            if recv in skip:
                recv = ""
        refs.append(
            {
                "ref": normalize_ref_value(raw_ref, profile),
                "raw_ref": raw_ref,
                "tracking": tracking,
                "receiver": recv,
            }
        )
    return refs


def write_remarks(
    src_path: str,
    dst_path: str,
    layout: dict[str, Any],
    profile: dict[str, Any],
    remarks: list[str],
) -> None:
    wb = openpyxl.load_workbook(src_path)
    sheet_name = layout["sheetName"]
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}")
    ws = wb[sheet_name]

    header_row = int(layout["headerRow"])
    cols = layout["columns"]
    rem_col = int(cols["remarks"]) + 1
    skip = set(profile.get("skip_receiver_values") or [])

    idx = 0
    for row_num in range(header_row + 1, (ws.max_row or 0) + 1):
        recv = cell_str(ws.cell(row=row_num, column=cols["receiver"] + 1).value)
        if not recv or recv in skip:
            continue
        if idx < len(remarks):
            ws.cell(row=row_num, column=rem_col).value = remarks[idx]
        idx += 1

    wb.save(dst_path)


def load_json_file(path: str) -> Any:
    with open(path, encoding="utf-8-sig") as handle:
        return json.load(handle)


def emit_json(payload: Any) -> None:
    """Print JSON for PHP shell_exec (ASCII-safe on Windows consoles)."""
    print(json.dumps(payload, ensure_ascii=True))


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except Exception:
            pass

    try:
        if len(sys.argv) < 4:
            emit_json(
                {
                    "error": (
                        "Usage: order_remarks_workbook.py <workbook> "
                        "<profiles.json> <mode> [args...]"
                    )
                }
            )
            sys.exit(1)

        path = sys.argv[1]
        profiles = load_json_file(sys.argv[2])
        mode = sys.argv[3]

        wb = openpyxl.load_workbook(path, data_only=False)

        if mode == "detect":
            override = sys.argv[4] if len(sys.argv) > 4 and sys.argv[4] else None
            layout = detect_layout(wb, profiles, override or None, mode="detect")
            if layout is None:
                emit_json({"error": "Could not detect a supported courier export format."})
                sys.exit(2)
            profile = profile_by_id(profiles, layout["profileId"])
            if profile is None:
                emit_json({"error": "Profile missing after detection."})
                sys.exit(2)
            ws = wb[layout["sheetName"]]
            orders = read_orders(ws, layout, profile)
            sample = orders[:3]
            emit_json(
                {
                    "layout": layout,
                    "orders": orders,
                    "sampleRows": sample,
                    "totalRows": len(orders),
                }
            )
            return

        if mode == "tracking_map":
            override = sys.argv[4] if len(sys.argv) > 4 and sys.argv[4] else None
            layout = detect_layout(wb, profiles, override or None, mode="tracking_map")
            if layout is None:
                emit_json({"error": "Could not detect a supported tracking export format."})
                sys.exit(2)
            profile = profile_by_id(profiles, layout["profileId"])
            if profile is None:
                emit_json({"error": "Profile missing after detection."})
                sys.exit(2)
            ws = wb[layout["sheetName"]]
            refs = read_tracking_refs(ws, layout, profile)
            emit_json(
                {
                    "profileId": layout["profileId"],
                    "profileLabel": layout.get("profileLabel", layout["profileId"]),
                    "refs": refs,
                    "totalRefs": len(refs),
                }
            )
            return

        if mode == "write":
            if len(sys.argv) < 7:
                emit_json(
                    {
                        "error": (
                            "write requires dst layout.json remarks.json file paths"
                        )
                    }
                )
                sys.exit(1)
            dst = sys.argv[4]
            layout = load_json_file(sys.argv[5])
            remarks = load_json_file(sys.argv[6])
            profile = profile_by_id(profiles, layout.get("profileId", ""))
            if profile is None:
                emit_json({"error": "Unknown profile for export."})
                sys.exit(2)
            write_remarks(path, dst, layout, profile, remarks)
            emit_json({"ok": True})
            return

        emit_json({"error": f"Unknown mode: {mode}"})
        sys.exit(1)
    except Exception as exc:
        emit_json({"error": str(exc)})
        sys.exit(1)


if __name__ == "__main__":
    main()
