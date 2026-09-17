from pathlib import Path
from docx import Document
from docx.shared import Inches
import hashlib,math
exec((Path(__file__).parent/'diagrams.py').read_text(encoding='utf-8-sig').split('# IPO')[0])
A=W/'five_assets';old=(A/'use_five.png').read_bytes()
c=Chart('use_clean',2250);c.im=Image.new('RGB',(2600,2250),'white');c.d=ImageDraw.Draw(c.im)
c.d.rectangle((285,20,2315,2200),outline='black',width=4)
c.text((320,60),'PYX TRACKER',48,True,1960,'center');c.text((320,125),'Web-based Enterprise Operations Management System',35,True,1960,'center')
rx,ry=205,82
def oval(x,y,label):
 c.d.ellipse((x-rx,y-ry,x+rx,y+ry),outline='#555',width=4);ls=wrap(label,360,font(33));c.text((x-180,y-len(ls)*20.6),label,33,w=360,align='center')
def edge(x,y,tx,ty):
 dx,dy=tx-x,ty-y;k=1/math.sqrt((dx/rx)**2+(dy/ry)**2);return(x+dx*k,y+dy*k)
def actor(x,y,label):
 c.d.ellipse((x-31,y-120,x+31,y-58),outline='black',width=4);c.line([(x,y-58),(x,y+85)]);c.line([(x-70,y),(x+70,y)]);c.line([(x,y+85),(x-55,y+180)]);c.line([(x,y+85),(x+55,y+180)]);c.text((x-130,y+215),label,34,True,260,'center')
actor(130,1030,'Admin/\nExecutive');actor(2470,1030,'Staff')
left=[(330,'Maintain costing and prices','Calculate configured costs'),(670,'Review inventory','Read stock movement history'),(1010,'Review order batches','Read order and waybill records'),(1350,'Review parcel records','Read packing and photo history'),(1710,'Manage RTS case','Read case and condition history')]
right=[(330,'Record inventory movement','Update relevant stock balance'),(620,'Finalize order batch','Validate prepared order data'),(910,'Confirm packing movement','Record applicable stock change'),(1200,'Save shipment photo','Store tracking and image'),(1490,'Record return case','Record item condition'),(1780,'Confirm RTS stock-in','Credit reviewed sellable stock')]
for leftside,items in [(True,left),(False,right)]:
 mx,dx=(565,1080) if leftside else (2035,1520)
 ax=200 if leftside else 2400
 for y,main,sub in items:
  oval(mx,y,main);oval(dx,y,sub)
  c.line([(ax,1030),(mx-rx if leftside else mx+rx,y)])
  a=(mx+rx,y) if leftside else (mx-rx,y);b=(dx-rx,y) if leftside else(dx+rx,y)
  c.line([a,b],True,True);c.text((min(a[0],b[0])-15,y-52),'«include»',26,w=abs(a[0]-b[0])+30,align='center')
# Optional administrative claim update extends RTS case management.
oval(1080,2070,'Update claim status')
a=edge(1080,2070,565,1710);b=edge(565,1710,1080,2070);c.line([a,b],True,True)
c.text((525,1905),'«extend»',30,w=320,align='center')
c.text((1250,1990),'Extension condition:\nAdmin/Executive chooses\nto update a claim.',31,w=940,align='center')
c.save()
src=W.parent/'documents/Pyx_Tracker_Capstone_Five_Features_Final.docx';d=Document(src)
for rel in d.part.rels.values():
 if 'image' in rel.reltype and rel.target_part.blob==old:rel.target_part._blob=(A/'use_clean.png').read_bytes()
for s in d.inline_shapes:
 rid=s._inline.graphic.graphicData.pic.blipFill.blip.embed
 if d.part.related_parts[rid].blob==(A/'use_clean.png').read_bytes():s.width=Inches(6);s.height=int(s.width*2250/2600)
for p in d.paragraphs:
 if p.text.startswith('The use case diagram defines exactly two actors:'):
  p.text='The use case diagram defines two actors: Admin/Executive and Staff. Solid lines connect each actor directly to an available use case. Dashed «include» arrows point to required behavior. The dashed «extend» arrow points from Update claim status to Manage RTS case: updating a claim is conditional on the Admin/Executive choosing that action. Case review does not require a claim update. Costing and pricing remain restricted to Admin/Executive. Packing movements and shipment-photo documentation are separate operations.'
# Keep the case matrix consistent with the conditional branch shown in the diagram.
for tb in d.tables:
 if any('RTS case and sellable stock-in' in cell.text for row in tb.rows for cell in row.cells):
  for row in tb.rows:
   if row.cells[0].text=='Exceptions':row.cells[1].text+=' Admin/Executive may optionally update claim status while managing a case; Staff does not perform that administrative action.'
out=src.with_name('Pyx_Tracker_Capstone_Use_Case_Final.docx');d.save(out);print(out)
