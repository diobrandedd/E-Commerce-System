from docx import Document
from zipfile import ZipFile
from lxml import etree
from pathlib import Path
for key,name in [('proposal','pyxtrackerdocs (2).docx'),('format','OUTLINE-AND-SAMPLE-DOCU-CHPATER-123 (2).docx')]:
 p=Path('C:/Users/Admin/Downloads')/name
 d=Document(p)
 out=[]
 for i,e in enumerate(d.element.body):
  texts=e.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
  s=''.join(x.text or '' for x in texts)
  colors=[x.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') for x in e.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color')]
  out.append(f'[{i}] [{e.tag.split("}")[-1]}] colors={colors} {s}')
 Path(f'output/capstone_work/{key}.txt').write_text('\n'.join(out),encoding='utf8')
 print(key, 'paragraphs',len(d.paragraphs),'tables',len(d.tables),'sections',len(d.sections))
 for s in d.sections: print('geometry',s.page_width.inches,s.page_height.inches,s.top_margin.inches,s.bottom_margin.inches,s.left_margin.inches,s.right_margin.inches)
 with ZipFile(p) as z:
  for n in z.namelist():
   if n.startswith('word/media/'):
    dst=Path('output/capstone_work')/key/Path(n).name;dst.parent.mkdir(exist_ok=True);dst.write_bytes(z.read(n))

