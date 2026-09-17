from pathlib import Path
from docx import Document
from docx.shared import Inches
exec((Path(__file__).parent/'diagrams.py').read_text(encoding='utf-8-sig').split('# IPO')[0])
A=W/'five_assets'
old_fdd=(A/'fdd.png').read_bytes();old_use=(A/'use_clean.png').read_bytes()
mods=['Integrated Inventory\nManagement','Product Costing\nand Pricing','Multi-Platform Orders\nand Waybills','Parcel Processing\nand Inventory Sync','Return-to-Sender\nManagement']
ss=[['Receive materials','Track production','Confirm finished stock','Review movements'],['Maintain formulas','Calculate costs','Define parcel recipes','Maintain prices'],['Prepare order files','Validate order data','Stamp waybills','Finalize batches'],['Record packing scans','Confirm movements','Save shipment photos','Review parcel history'],['Record return case','Classify condition','Review stock-in','Confirm sellable stock']]
c=Chart('fdd_no_numbers',1370);c.box(600,10,600,110,'Pyx Tracker',size=45);c.line([(900,120),(900,200),(180,200),(1620,200)])
for i,(m,subs) in enumerate(zip(mods,ss)):
 x=10+i*360;mid=x+170;c.line([(mid,200),(mid,260)])
 c.d.rectangle((x,260,x+335,460),fill='#fff4bb',outline='black',width=3);c.text((x+14,285),m,35,True,307,'center')
 c.line([(mid,460),(mid,510),(x+15,510),(x+15,1235)])
 for j,s in enumerate(subs):
  y=550+j*190;c.line([(x+15,y+60),(x+40,y+60)]);c.box(x+40,y,295,135,s,size=33)
c.save()
# Every oval is an exact FDD leaf; headings identify its FDD parent.
c=Chart('use_fdd_aligned',2520);c.im=Image.new('RGB',(2100,2520),'white');c.d=ImageDraw.Draw(c.im)
c.d.rectangle((265,20,1835,2475),outline='black',width=4)
c.text((300,60),'PYX TRACKER',48,True,1500,'center')
c.text((300,130),'Use cases derived from the FDD',34,w=1500,align='center')
def actor(x,y,label):
 c.d.ellipse((x-30,y-120,x+30,y-60),outline='black',width=4);c.line([(x,y-60),(x,y+80)]);c.line([(x-60,y),(x+60,y)]);c.line([(x,y+80),(x-50,y+170)]);c.line([(x,y+80),(x+50,y+170)]);c.text((x-120,y+210),label,34,True,240,'center')
actor(125,1200,'Admin/\nExecutive');actor(1975,1200,'Staff')
labels=[]
def oval(x,y,label):
 labels.append(label);c.d.ellipse((x-325,y-43,x+325,y+43),outline='#555',width=3);c.text((x-300,y-22),label,36,w=600,align='center')
c.text((325,620),mods[1],37,True,650,'center')
for j,label in enumerate(ss[1]):
 y=800+j*210;oval(650,y,label);c.line([(185,1200),(325,y)])
c.text((340,1710),'Admin/Executive maintains\nformulas, costs, parcel recipes,\nand selling prices.',33,w=620,align='center')
c.text((340,1920),'Solid lines show actor\nparticipation. Each oval\nmatches a function in the FDD.',33,w=620,align='center')
for g,i in enumerate([0,2,3,4]):
 top=240+g*535;c.text((1135,top),mods[i],36,True,650,'center')
 for j,label in enumerate(ss[i]):
  y=top+130+j*103;oval(1460,y,label);c.line([(1915,1200),(1785,y)])
assert len(labels)==20 and set(labels)==set(sum(ss,[]))
c.save()
src=W.parent/'documents/Pyx_Tracker_Capstone_Use_Case_Final.docx';d=Document(src)
mapping={old_fdd:(A/'fdd_no_numbers.png').read_bytes(),old_use:(A/'use_fdd_aligned.png').read_bytes()}
replaced=0
for rel in d.part.rels.values():
 if 'image' in rel.reltype and rel.target_part.blob in mapping:rel.target_part._blob=mapping[rel.target_part.blob];replaced+=1
assert replaced==2
for s in d.inline_shapes:
 rid=s._inline.graphic.graphicData.pic.blipFill.blip.embed
 if d.part.related_parts[rid].blob==(A/'use_fdd_aligned.png').read_bytes():
  s.width=Inches(6);s.height=Inches(7.2)
  p=s._inline.getparent().getparent().getparent();pPr=p.get_or_add_pPr()
  from docx.oxml import OxmlElement
  el=OxmlElement('w:pageBreakBefore');pPr.append(el)
for p in d.paragraphs:
 if p.text.startswith('The use case diagram defines two actors:'):
  p.text='The use case diagram follows the Functional Decomposition Diagram directly. Its twenty use cases use the same names as the twenty FDD functions and are grouped under the same five features. Admin/Executive is the primary actor for costing and pricing; Staff is the primary actor for the inventory, order, parcel, and return functions. These primary responsibilities do not replace the permission rules described in the use case matrices. Solid lines identify actor participation. The FDD specifies functional decomposition, so its branches are not automatically treated as «include» or «extend» relationships.'
 if p.text=='Use Case Matrix':p.paragraph_format.page_break_before=True
for tb in d.tables:
 for row in tb.rows:
  for cell in row.cells:
   suffix=' Admin/Executive may optionally update claim status while managing a case; Staff does not perform that administrative action.'
   if suffix in cell.text:cell.text=cell.text.replace(suffix,'')
out=src.with_name('Pyx_Tracker_Capstone_FDD_Aligned.docx');d.save(out);print(out)

