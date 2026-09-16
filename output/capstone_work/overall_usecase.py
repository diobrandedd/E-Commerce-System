from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
exec((Path(__file__).parent/'diagrams.py').read_text(encoding='utf-8-sig').split('# IPO')[0])
c=Chart('use_overall',2200);c.im=Image.new('RGB',(2400,2200),'white');c.d=ImageDraw.Draw(c.im)
c.d.rectangle((270,20,2130,2150),outline='black',width=4)
c.text((320,65),'PYX TRACKER',48,True,1760,'center')
c.text((320,130),'Web-based Enterprise Operations Management System',36,True,1760,'center')
def oval(x,y,label):
 c.d.ellipse((x-190,y-85,x+190,y+85),outline='#555',width=4)
 ls=wrap(label,330,font(32));c.text((x-165,y-len(ls)*20),label,32,w=330,align='center')
def actor(x,y,label):
 c.d.ellipse((x-30,y-110,x+30,y-50),outline='#555',width=4)
 c.line([(x,y-50),(x,y+85)]);c.line([(x-65,y),(x+65,y)])
 c.line([(x,y+85),(x-55,y+180)]);c.line([(x,y+85),(x+55,y+180)])
 c.text((x-120,y+205),label,31,True,240,'center')
def include(a,b,y):
 c.line([(a,y),(b,y)],True,True);c.text((min(a,b),y-52),'«include»',27,w=abs(a-b),align='center')
# Each cluster has an actor-facing goal and required subordinate behavior.
groups=[(True,420,'Executive','Calculate product cost','Read product formula','Compute component costs'),(True,1050,'Production\nManager','Confirm finished stock','Validate stock input','Record stock movement'),(True,1720,None,'Save shipment photo','Validate tracking input','Store photo record'),(False,420,'CSR Staff /\nTeam Leader','Finalize order batch','Validate order lines','Create parcel records'),(False,1050,'POS Cashier','Complete checkout','Validate sale basket','Deduct selected stock'),(False,1720,'Authorized\nRTS Operator','Confirm RTS stock-in','Validate draft session','Credit sellable inventory')]
for left,y,role,main,a,b in groups:
 mx,dx=(500,970) if left else (1900,1430)
 oval(mx,y,main);oval(dx,y-130,a);oval(dx,y+130,b)
 for yy in [y-130,y+130]:
  start=mx+190 if left else mx-190;end=dx-190 if left else dx+190
  c.line([(start,y),(start+35 if left else start-35,y),(start+35 if left else start-35,yy),(end,yy)],True,True)
  c.text((min(start,end)-10,yy-53),'«include»',27,w=abs(end-start)+20,align='center')
 ax=125 if left else 2275
 if role:actor(ax,y,role)
 if left and role is None:
  c.line([(190,1050),(235,1050),(235,y),(310,y)])
 else:c.line([(ax+65 if left else ax-65,y),(mx-190 if left else mx+190,y)])
# An optional checkout behavior extends the checkout goal.
oval(1870,1400,'Apply approved discount')
c.line([(2060,1400),(2095,1400),(2095,1050),(2090,1050)],True,True)
c.text((1820,1260),'«extend»',27,w=250)
c.save()
src=W.parent/'documents/Pyx_Tracker_Capstone_Chapters_I_III_Revised.docx';doc=Document(src)
paras=list(doc.paragraphs);start=next(i for i,p in enumerate(paras) if p.text.startswith('Human actors appear outside'))
end=next(i for i,p in enumerate(paras[start:],start) if p.text=='Use Case Matrix')
anchor=paras[end]
for p in paras[start:end]:p._element.getparent().remove(p._element)
p=anchor.insert_paragraph_before('The overall use case diagram places the operational actors outside one Pyx Tracker system boundary. Solid lines associate roles with their goals; dashed «include» arrows point to required behavior, while the optional approved discount «extend» arrow points to checkout. Production personnel also save shipment photos. Permissions remain subject to the role and department checks described in the use case matrices.')
p=anchor.insert_paragraph_before();p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(A/'use_overall.png'),width=Inches(6))
p=anchor.insert_paragraph_before('Figure 3.5  Overall Use Case Diagram of Pyx Tracker');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
anchor.paragraph_format.page_break_before=True
for p in doc.paragraphs:
 m=re.match(r'Figure 3\.(\d+)(.*)',p.text)
 if m and int(m[1])>=11:p.text=f'Figure 3.{int(m[1])-5}'+m[2]
out=src.with_name('Pyx_Tracker_Capstone_Updated_Use_Case.docx');doc.save(out)
print(out)
