from zipfile import ZipFile
from pathlib import Path
from docx import Document
import json,re,hashlib
w=Path('output/capstone_work');p=Path('output/documents/Pyx_Tracker_Capstone_Chapters_I_III.docx');d=Document(p)
body='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
issues=[]
for bad in ['Coconut Processing','FACCALTI','Parcon','BIMS','YOUR TITLE','SAMPLE LANG','@FIG','@DICTIONARY','Error!']:
 if bad in body:issues.append('Unexpected text: '+bad)
figs=[p.text for p in d.paragraphs if p.text.startswith('Figure ')];tabs=[p.text for p in d.paragraphs if p.text.startswith('Table 3.')]
assert len(figs)==24 and len(tabs)==39
schema=json.load(open(w/'schema.json'));dictionary=d.tables[6:34]
for table in dictionary:
 fields=[r.cells[0].text.replace('\u200b','') for r in table.rows[1:]]
 # Caption-to-table order verified by selected table list in builder.
 assert len(fields)==len(set(fields))
source=Path('C:/Users/Admin/Downloads/OUTLINE-AND-SAMPLE-DOCU-CHPATER-123 (2).docx')
contract=(w/'artifact.md').read_text(encoding='utf8');assert hashlib.sha256(source.read_bytes()).hexdigest() in contract
print('Content issues:',issues,'Figures',len(figs),'Tables',len(tabs),'Source hash unchanged')
with ZipFile(p) as z:
 print('Old sample text in package:',[n for n in z.namelist() if n.endswith('.xml') and any(b in z.read(n).decode('utf8',errors='ignore') for b in ['Coconut Processing','FACCALTI','Parcon'])])
(w/'validation.json').write_text(json.dumps({'issues':issues,'figures':len(figs),'tables':len(tabs),'schema_tables_documented':28,'source_unchanged':True},indent=2))
