"""Collapse ERD section to a single diagram and save under Desktop/capstone."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

SRC_CANDIDATES = [
    Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx.new"),
    Path(r"C:/Users/Admin/Desktop/capstone/Pyx_Tracker_Capstone_SQL_ERD.docx"),
    Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx.bak"),
]
ORIGINAL = Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx")
OUT_NEW = Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx.new")
DESKTOP = Path(r"C:/Users/Admin/Desktop/capstone")
FINAL_NAME = "Pyx_Tracker_Capstone_SQL_ERD.docx"

DESC = (
    "Figure 3.12 presents the single Entity Relationship Diagram of Pyx Tracker derived from the "
    "project SQL schema (full_schema.sql, applied migrations, and live rts_system metadata). "
    "Entity boxes list primary keys and principal attributes. Crow's-foot notation marks "
    "one-to-many relationships among users, inventory, honey production, products, POS sales, "
    "online orders, deliveries, and RTS cases. The diagram is not split by feature area."
)


def set_para_text(paragraph, text: str) -> None:
    for r in paragraph.runs:
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def remove_element(el) -> None:
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def main() -> None:
    src = next((p for p in SRC_CANDIDATES if p.exists()), None)
    if src is None:
        raise FileNotFoundError("No source DOCX found")

    d = Document(str(src))
    print("opened", src)

    paras = list(d.paragraphs)
    remove_els = []

    for idx, p in enumerate(paras):
        text = p.text.strip()
        has_blip = bool(p._p.findall(".//" + qn("a:blip")))

        if text.startswith("Figure 3.12") and "Entity Relationship" in text:
            set_para_text(p, "Figure 3.12  Entity Relationship Diagram of Pyx Tracker")
            print("kept", p.text)
            continue

        if text.startswith("Figure 3.13") or text.startswith("Figure 3.14") or text.startswith(
            "Figure 3.15"
        ):
            # remove this caption and the preceding image paragraph if empty+blip
            if idx > 0:
                prev = paras[idx - 1]
                if prev._p.findall(".//" + qn("a:blip")) and not prev.text.strip():
                    remove_els.append(prev._element)
            remove_els.append(p._element)
            print("queue remove", text[:70])
            continue

        if (
            text.startswith("Figure 3.12 presents")
            or text.startswith("The ERDs are derived")
            or ("Figures 3.13" in text and "Entity Relationship" in text)
        ):
            set_para_text(p, DESC)
            print("updated description")

    # unique preserve order
    seen = set()
    for el in remove_els:
        if id(el) in seen:
            continue
        seen.add(id(el))
        remove_element(el)

    DESKTOP.mkdir(parents=True, exist_ok=True)
    desktop_out = DESKTOP / FINAL_NAME
    d.save(str(OUT_NEW))
    shutil.copy2(OUT_NEW, desktop_out)

    # verify
    v = Document(str(desktop_out))
    caps = [
        p.text.strip()
        for p in v.paragraphs
        if "Entity Relationship Diagram of Pyx Tracker" in p.text
        or p.text.strip().startswith("Figure 3.1")
        and "Entity Relationship" in p.text
    ]
    cont = [p.text.strip() for p in v.paragraphs if "(continued)" in p.text]
    print("erd-related captions:", caps)
    print("continued leftover:", cont)
    print("saved", desktop_out, desktop_out.stat().st_size)
    print("saved", OUT_NEW, OUT_NEW.stat().st_size)

    try:
        ORIGINAL.unlink(missing_ok=True)
        shutil.copy2(OUT_NEW, ORIGINAL)
        print("overwrote", ORIGINAL)
    except Exception as e:
        print("ORIGINAL_STILL_LOCKED", e)


if __name__ == "__main__":
    main()
