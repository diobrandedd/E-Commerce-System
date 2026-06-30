#!/usr/bin/env python3
"""Build SPX tracking export subset from Online Orders Prep + full SPX export."""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

import openpyxl

PREP_DEFAULT = Path(r"c:\Users\Admin\Downloads\Online_Orders_Prep_2026-06-27 (2).xlsx")
SPX_SRC_DEFAULT = Path(r"c:\Users\Admin\Downloads\0f315d50425041509450ebb42012eb6a_remarks.xlsx")
OUT_DEFAULT = Path(
    r"c:\Users\Admin\Downloads\Online_Orders_Prep_2026-06-27_SPX_tracking_export.xlsx"
)

HEADERS = [
    "Tracking No.",
    "Tracking No. link",
    "Customer Reference No.",
    "Customer Reference No. link",
    "Create Time",
    "Tracking Status",
    "Account ID",
    "Original pickup option",
    "Actual pickup option",
    "Scheduled Pickup Time",
    "Actual Pickup/Drop Off Time",
    "Delivered Time",
    "Delivery OnHold Times",
    "Delivery OnHold Reason",
    "Returning Start Time",
]


def load_prep_rows(prep_path: Path) -> list[dict[str, str]]:
    wb = openpyxl.load_workbook(prep_path, read_only=True, data_only=True)
    rows: list[dict[str, str]] = []
    for row in wb["Orders"].iter_rows(min_row=6, values_only=True):
        if row[2]:
            rows.append(
                {
                    "name": str(row[2]).strip(),
                    "items": str(row[3] or "").strip(),
                    "tracking": str(row[7] or "").strip(),
                }
            )
    wb.close()
    return rows


def load_spx_by_tracking(spx_path: Path) -> dict[str, list]:
    wb = openpyxl.load_workbook(spx_path, read_only=True, data_only=True)
    ws = wb.active
    out: dict[str, list] = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 2 or not row[0]:
            continue
        out[str(row[0]).strip()] = list(row)
    wb.close()
    return out


def main() -> int:
    prep_path = Path(sys.argv[1]) if len(sys.argv) > 1 else PREP_DEFAULT
    spx_src = Path(sys.argv[2]) if len(sys.argv) > 2 else SPX_SRC_DEFAULT
    out_path = Path(sys.argv[3]) if len(sys.argv) > 3 else OUT_DEFAULT

    prep_rows = load_prep_rows(prep_path)
    spx_map = load_spx_by_tracking(spx_src)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(1, 1, f"Report Download Time: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    for col, header in enumerate(HEADERS, 1):
        ws.cell(2, col, header)

    row_num = 3
    for prep in prep_rows:
        src = spx_map.get(prep["tracking"])
        if not src:
            print(f"ERROR: no SPX row for tracking {prep['tracking']} ({prep['name']})", file=sys.stderr)
            return 1
        for col, val in enumerate(src, 1):
            if col <= len(HEADERS):
                ws.cell(row_num, col, val)
        row_num += 1

    wb.save(out_path)
    print(f"Wrote {out_path} ({len(prep_rows)} rows)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
