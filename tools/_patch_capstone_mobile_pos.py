"""Add Point of Sale / Mobile POS feature coverage to the capstone DOCX.

Evidence:
- Approved proposal: Point of Sale System
- RTS: api/v1 auth pos_login, pos.php sales_history/receipt_get/sale_refund,
  pos_mobile_staff, migration_phase80_pos_mobile_tokens.sql
- Mobile app: https://github.com/diobrandedd/POS-PrinterDriver-Mobile
  (Flutter Android, PIN/Bearer login, sell/AR/history/refund, ESC/POS print)

Does not modify ERD (image14) or Database Design (image13).
Regenerates Use Case (image7) to include Mobile POS flows.
"""

from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

import importlib.util

SRC = Path(r"C:/Users/Admin/Desktop/capstone/Pyx_Tracker_Capstone_SQL_ERD.docx")
ALT = Path(r"C:/Users/Admin/Desktop/capstone/Pyx_Tracker_Capstone_SQL_ERD_updated.docx")
WORK = Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx")
DESKTOP = Path(r"C:/Users/Admin/Desktop/capstone")
NOTE = DESKTOP / "Mobile_POS_Feature_Addition.md"
UC_SCRIPT = Path(r"C:/xampp/htdocs/rts/tools/_patch_capstone_use_case.py")

MOBILE_POS_SCOPE = (
    "- Point of Sale System — digitizes walk-in sales through the web POS and the companion "
    "Android Mobile POS application (Pyx POS), including staff authentication, barcode sell "
    "checkout, accounts-receivable credit, sales history with reprint and full-sale refund, "
    "and ESC/POS thermal receipt printing over Bluetooth, USB, or Wi-Fi (port 9100). Sales "
    "synchronize to the same Pyx Tracker / RTS inventory and receipt records."
)

BUILD6 = (
    "Build 6 Point of Sale System (Web and Mobile)\n"
)  # heading only; body separate


def set_text(paragraph, text: str) -> None:
    for r in paragraph.runs:
        r.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def insert_paragraph_after(paragraph, text: str, style=None):
    """Insert a new paragraph after the given one."""
    from copy import deepcopy

    from docx.text.paragraph import Paragraph

    new_p = deepcopy(paragraph._p)
    # clear runs in the cloned element
    for child in list(new_p):
        if child.tag == qn("w:r"):
            new_p.remove(child)
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style is not None:
        try:
            new_para.style = style
        except Exception:
            pass
    set_text(new_para, text)
    return new_para


def replace_five_with_six(text: str) -> str:
    replacements = [
        ("the five features", "the six features"),
        ("these five features", "these six features"),
        ("all five modules", "all six modules"),
        ("the five connected modules", "the six connected modules"),
        ("Each of the five features", "Each of the six features"),
        ("five-feature workflow", "six-feature workflow"),
        ("oversees the five features", "oversees the six features"),
        ("oversees all five modules", "oversees all six modules"),
    ]
    out = text
    for a, b in replacements:
        out = out.replace(a, b)
    return out


def patch_body(doc: Document) -> list[str]:
    log: list[str] = []
    paras = list(doc.paragraphs)

    # Specific objective 6
    for i, p in enumerate(paras):
        if p.text.strip().startswith("5. Implement an RTS module"):
            # insert objective 6 after if not present
            nxt = paras[i + 1].text.strip() if i + 1 < len(paras) else ""
            if not nxt.startswith("6. Implement"):
                # Add after para i by editing following evaluation sentence? Insert new para
                new = (
                    "6. Implement a Point of Sale module that records walk-in sales through the web POS "
                    "and the Android Mobile POS application, deducts inventory in real time, prints or "
                    "stores receipts, and maintains sales history including reprint and refund paths."
                )
                # Put objective 6 before the evaluation sentence
                if "module-specific tests" in nxt.lower() or "Chapter III evaluate" in nxt:
                    set_text(paras[i + 1], new)
                    # re-add evaluation after — insert
                    insert_paragraph_after(
                        paras[i + 1],
                        "The module-specific tests in Chapter III evaluate these objectives through "
                        "observable balances, cost calculations, generated documents, transaction "
                        "records, confirmed stock changes, and POS sale or refund outcomes.",
                    )
                    log.append("added specific objective 6 (POS)")
                else:
                    insert_paragraph_after(p, new)
                    log.append("inserted specific objective 6")
            break

    # Scope: insert POS feature after Parcel Processing bullet if missing
    has_pos_scope = any("Point of Sale" in p.text for p in doc.paragraphs)
    if not has_pos_scope:
        for p in doc.paragraphs:
            if p.text.strip().startswith("- Parcel Processing and Inventory Synchronization"):
                insert_paragraph_after(p, MOBILE_POS_SCOPE)
                log.append("added POS scope bullet")
                break

    # Scope intro: mention POS + mobile
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("The study covers a role-based web application"):
            new = (
                "The study covers a role-based web application for the internal operations of Pyx Food "
                "Products, together with a companion Android Mobile POS terminal for walk-in sales. "
                "Its primary user categories are Admin/Executive and Staff. Admin/Executive oversees the "
                "six features and exclusively maintains costing and pricing. Staff performs assigned "
                "inventory, order-processing, parcel-processing, Point of Sale (web or Mobile POS), and "
                "return operations according to role permissions. Customers are indirect beneficiaries "
                "of completed sales and fulfilled orders."
            )
            set_text(p, new)
            log.append("updated scope intro for Mobile POS")
            break

    # Auth sentence
    for p in doc.paragraphs:
        if p.text.strip().startswith("Authentication and permissions support these"):
            set_text(
                p,
                "Authentication and permissions support these six features, including web login and "
                "Mobile POS staff authentication against the RTS API. Supporting records are documented "
                "only where they are needed for an included feature. The same module names are used in Chapter III.",
            )
            log.append("updated auth/six features sentence")
            break

    # Limitations: add mobile POS limits
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Users need access to the application server"):
            new = (
                "Users need access to the application server through a working network and a compatible "
                "browser for the web application. The Mobile POS application targets Android devices and "
                "requires network reachability to the RTS API endpoint, assigned Mobile POS staff access, "
                "and a configured thermal printer (Bluetooth, USB, or Wi-Fi port 9100) for on-device "
                "receipt printing. Scanning depends on readable barcodes and suitable camera or scanner "
                "hardware. Reliable records require correct starting balances, mappings, units, prices, "
                "and confirmation of physical movements when stock is affected."
            )
            set_text(p, new)
            log.append("updated limitations for Mobile POS hardware/network")
            break

    # Significance staff sentence
    for p in doc.paragraphs:
        if "five-feature workflow" in p.text or "six-feature workflow" in p.text:
            set_text(
                p,
                replace_five_with_six(p.text).replace(
                    "shipment evidence, and returns",
                    "shipment evidence, walk-in POS sales, and returns",
                ),
            )
            log.append("updated significance for POS")
            break

    # Generic five→six replacements in remaining body (conservative)
    for p in doc.paragraphs:
        t = p.text
        if not t:
            continue
        # skip chapter II "five themes"
        if "five themes" in t:
            continue
        nt = replace_five_with_six(t)
        if nt != t:
            set_text(p, nt)
            log.append(f"five→six: {t[:60]}...")

    # Definition of terms — insert Mobile POS alphabetically near M/P
    terms_to_add = [
        (
            "Mobile POS — The Android companion application (Pyx POS) that authenticates assigned "
            "staff against Pyx Tracker / RTS, records walk-in sales and accounts-receivable credits, "
            "prints ESC/POS thermal receipts, and synchronizes sale, reprint, and refund actions to "
            "the shared server records.",
            "Inventory pool",
        ),
        (
            "Point of Sale — The walk-in sales capability of Pyx Tracker, delivered through the web "
            "POS interface and the Mobile POS application, including authentication, checkout, "
            "inventory deduction, receipts, and sales history.",
            "Parcel recipe",
        ),
        (
            "ESC/POS — The command dialect used by the Mobile POS thermal printer path for Bluetooth, "
            "USB, or network (port 9100) receipt printing without watermark overlays.",
            "Cost of goods sold",
        ),
    ]
    existing = " ".join(p.text for p in doc.paragraphs)
    for term_text, after_marker in terms_to_add:
        term_name = term_text.split(" — ")[0]
        if term_name in existing:
            continue
        for p in doc.paragraphs:
            if p.text.strip().startswith(after_marker):
                insert_paragraph_after(p, term_text)
                log.append(f"added term: {term_name}")
                existing += term_text
                break

    # Build 6 after Build 5
    has_build6 = any("Build 6" in p.text for p in doc.paragraphs)
    if not has_build6:
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip() == "Build 5 Return-to-Sender Management":
                # find end of build 5 body (next Planning heading)
                # insert heading + 2 paras before Planning
                pass
        for i, p in enumerate(list(doc.paragraphs)):
            if p.text.strip().startswith("Planning and Conception Initiation Phase"):
                # Insert before this paragraph by adding after previous
                prev = doc.paragraphs[i - 1]
                h = insert_paragraph_after(prev, "Build 6 Point of Sale System (Web and Mobile)")
                # style heading if possible
                try:
                    h.style = doc.styles["Heading 3"]
                except Exception:
                    pass
                b1 = insert_paragraph_after(
                    h,
                    "Design and development implement walk-in sales on the web POS and on the Android "
                    "Mobile POS application documented in the POS-PrinterDriver-Mobile repository. Staff "
                    "authenticate through POS credentials or assigned Mobile POS staff login with Bearer "
                    "tokens, scan barcodes, complete cash or accounts-receivable checkout, and print "
                    "thermal receipts through Bluetooth, USB, or Wi-Fi ESC/POS connections. Sales, "
                    "reprints, and full-sale refunds update the shared RTS inventory and receipt records.",
                )
                insert_paragraph_after(
                    b1,
                    "Testing verifies successful login for authorized Mobile POS staff, sale posting with "
                    "inventory deduction, receipt generation or print path availability, sales-history "
                    "retrieval, and refund restock behavior. Unauthorized staff and unreachable API "
                    "endpoints are treated as exception cases. Claims of completed field evaluation remain "
                    "limited to available test records.",
                )
                log.append("added Build 6 Point of Sale")
                break

    # Technology stack paragraph
    for p in doc.paragraphs:
        if p.text.strip().startswith("The stack follows the proposal"):
            set_text(
                p,
                "The stack follows the proposal and reviewed project: React and TypeScript for the web "
                "interface, Tailwind CSS for presentation, PHP for business actions, MySQL for persistent "
                "records, and Python for document utilities. Vite supports frontend builds. The companion "
                "Mobile POS terminal is implemented in Flutter (Dart) for Android and connects to the same "
                "RTS API for authentication, checkout, sales history, and refunds, with a native ESC/POS "
                "print bridge for Bluetooth, USB, and Wi-Fi printers. Defined API actions connect these components.",
            )
            log.append("updated technology stack for Flutter Mobile POS")
            break

    # TOC entries — soft update if present as plain paragraphs
    for p in doc.paragraphs:
        if p.text.strip() == "Build 5 Return-to-Sender Management\t27":
            insert_paragraph_after(p, "Build 6 Point of Sale System (Web and Mobile)\t27")
            log.append("TOC build 6 line")
            break

    return log


def patch_tables(doc: Document) -> list[str]:
    log: list[str] = []
    for t in doc.tables:
        # Software requirements
        if len(t.rows) >= 2 and t.rows[0].cells[0].text.strip().startswith("Software Requirements"):
            for r in t.rows:
                if r.cells[0].text.strip() == "Language":
                    r.cells[1].text = "PHP, React, TypeScript, Flutter (Dart)"
                    log.append("software language +Flutter")
                if r.cells[0].text.strip() == "Operating system":
                    r.cells[1].text = "Microsoft Windows 11 (web/server); Android (Mobile POS)"
                    log.append("software OS +Android")

        # Hardware — extend if small table
        if len(t.rows) >= 1 and t.rows[0].cells[0].text.strip() == "Processor":
            # append rows if not already
            blob = " ".join(c.text for r in t.rows for c in r.cells)
            if "Android" not in blob:
                # add rows by cloning last row pattern
                def add_row(label, value):
                    row = t.add_row()
                    row.cells[0].text = label
                    if len(row.cells) > 1:
                        row.cells[1].text = value

                add_row("Android device (Mobile POS)", "Android phone or tablet with camera and Bluetooth")
                add_row("Thermal printer (Mobile POS)", "ESC/POS printer (Bluetooth / USB / Wi-Fi port 9100)")
                log.append("hardware +Android +thermal printer")

        # List of Modules
        if (
            len(t.rows) >= 2
            and t.rows[0].cells[0].text.strip() == "Module"
            and "Functions" in t.rows[0].cells[1].text
        ):
            blob = " ".join(c.text for r in t.rows for c in r.cells)
            if "Point of Sale" not in blob:
                # insert before RTS row if present
                # python-docx can't easily insert mid-table; append then note order,
                # or replace RTS row and re-add. Simpler: append POS as new row with Obj 6
                # and renumber - actually add before last by editing: append POS
                row = t.add_row()
                row.cells[0].text = "Point of Sale System (Web and Mobile POS)"
                row.cells[1].text = (
                    "Web and Android Mobile POS authentication, barcode sell and AR credit checkout, "
                    "ESC/POS receipt printing, sales history, reprint, and full-sale refund synchronized "
                    "to RTS inventory."
                )
                if len(row.cells) > 2:
                    row.cells[2].text = "6"
                log.append("modules table +POS")

        # Use case matrix assumptions / primary actor — table 0 style
        if len(t.rows) >= 8 and "Primary Actor" in t.rows[7].cells[0].text:
            for r in t.rows:
                label = r.cells[0].text.strip()
                if label == "Primary Actor":
                    r.cells[1].text = (
                        "Admin/Executive, who oversees all six modules, maintains product costing and "
                        "selling prices, and may authorize Mobile POS staff assignments."
                    )
                    log.append("UC matrix primary actor")
                if label == "Secondary Actor":
                    r.cells[1].text = (
                        "Staff, who can access the inventory, order-processing, parcel-processing, "
                        "Point of Sale (web or Mobile POS), and RTS features according to role permissions."
                    )
                    log.append("UC matrix secondary actor")
                if label == "Assumptions":
                    r.cells[1].text = replace_five_with_six(r.cells[1].text)
                    if "Point of Sale" not in r.cells[1].text and "POS" not in r.cells[1].text:
                        r.cells[1].text = r.cells[1].text.replace(
                            "Orders, Parcel, and RTS",
                            "Orders, Parcel, Point of Sale, and RTS",
                        )
                    log.append("UC matrix assumptions")

    return log


def patch_use_case_diagram(doc_path: Path) -> None:
    """Extend UC generator data via monkeypatch-style regenerate with POS ovals, then replace image7 only."""
    # Load and call generate with extended content by writing a focused POS-aware UC image
    from PIL import Image, ImageDraw, ImageFont
    import math

    # Import drawing helpers by executing prior script's generate after patching — simpler to
    # run a local regenerate that includes POS, then swap image7 and preserve 13/14.
    spec = importlib.util.spec_from_file_location("uc_patch", UC_SCRIPT)
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)

    # Monkeypatch TABLES-like structure by regenerating with custom generate_with_pos
    out = Path(r"C:/xampp/htdocs/rts/output/documents/_use_case_pyx_tracker.png")

    # Reuse fonts/colors from mod
    W, H = 2200, 1680
    img = Image.new("RGB", (W, H), mod.BG)
    draw = ImageDraw.Draw(img)

    title = "Use Case Diagram for Pyx Tracker"
    tw, _ = mod.text_size(draw, title, mod.font_title)
    draw.text(((W - tw) / 2, 14), title, font=mod.font_title, fill=mod.INK)
    draw.text(
        (40, 46),
        "Includes Android Mobile POS (Pyx POS) synchronized to RTS  |  Solid = association  |  Dashed <<include>> / <<extend>>",
        font=mod.font_note,
        fill=mod.MUTED,
    )

    bx1, by1, bx2, by2 = 260, 85, 1940, 1600
    draw.rectangle((bx1, by1, bx2, by2), outline=mod.BORDER, width=3)
    btitle = "Pyx Tracker: Web-based Enterprise Operations Management System"
    btw, _ = mod.text_size(draw, btitle, mod.font_boundary)
    draw.text((bx1 + (bx2 - bx1 - btw) / 2, by1 + 8), btitle, font=mod.font_boundary, fill=mod.INK)

    exec_pt = mod.draw_actor(draw, 110, 260, "Executive")
    hr_pt = mod.draw_actor(draw, 110, 560, "HR\nPersonnel")
    staff_pt = mod.draw_actor(draw, 2090, 420, "Staff")
    for i, line in enumerate("(CSR / Production /\nPOS / Delivery /\nMobile POS)".split("\n")):
        tw2, _ = mod.text_size(draw, line, mod.font_note)
        draw.text((2090 - tw2 / 2, 530 + i * 14), line, font=mod.font_note, fill=mod.MUTED)
    cust_pt = mod.draw_actor(draw, 2090, 1100, "Customer")

    boxes = {}

    def uc(key, x, y, label, bold=False, max_chars=20):
        boxes[key] = mod.draw_oval(draw, x, y, label, bold=bold, max_chars=max_chars)
        return boxes[key]

    for label, x, y in [
        ("User & Access", 380, 110),
        ("Costing", 620, 110),
        ("Inventory / Orders", 980, 110),
        ("Fulfillment / RTS", 1400, 110),
        ("POS / Mobile POS", 1720, 110),
    ]:
        draw.text((x, y), label, font=mod.font_mod, fill="#2F80ED")

    # Left cluster (compressed)
    uc("manage_users", 400, 210, "Manage Users\nand Access", True)
    uc("assign_roles", 400, 310, "Assign Roles /\nMobile POS Staff")
    uc("costing", 400, 450, "Product Costing\nand Pricing", True)
    uc("formulas", 400, 560, "Maintain Formulas\nand Prices")
    uc("dashboards", 400, 700, "Monitor Dashboards\nand Reports", True)
    uc("approve_corr", 400, 820, "Approve Inventory\nCorrections")
    uc("review_rts", 400, 940, "Review RTS Claims\nand Dashboard")

    # Center
    uc("inventory", 820, 230, "Integrated Inventory\nManagement", True, 24)
    uc("receive", 700, 350, "Receive Materials\nand Stock-In")
    uc("honey", 940, 350, "Track Honey\nProduction")
    uc("orders", 820, 500, "Multi-Platform\nOrder Processing", True, 22)
    uc("prepare", 700, 620, "Prepare VIP /\nShopee Orders")
    uc("waybills", 940, 620, "Stamp Waybills /\nFinalize Batches")
    uc("parcel", 820, 780, "Parcel Processing\nand Inventory Sync", True, 24)
    uc("scan", 700, 900, "Scan Tracking /\nCapture Photos")
    uc("status", 940, 900, "Update Parcel\nStatus to Shipped")
    uc("rts", 820, 1080, "Return-to-Sender\nManagement", True)
    uc("rts_case", 700, 1200, "Record RTS Case\nand Sync Pools")
    uc("confirm_stock", 940, 1200, "Confirm Sellable\nStock Restoration")

    # Right: POS web + Mobile POS (sample-style includes)
    uc("pos", 1550, 230, "Point of Sale\nSystem", True)
    uc("web_pos", 1400, 350, "Use Web POS\nCheckout")
    uc("mobile_pos", 1700, 350, "Use Mobile POS\n(Android Pyx POS)", True, 22)
    uc("pos_auth", 1550, 470, "Authenticate POS\nStaff / Bearer Token")
    uc("barcode_sell", 1400, 590, "Scan Barcode\nand Build Cart")
    uc("checkout", 1700, 590, "Cash or AR\nCredit Checkout")
    uc("print_receipt", 1400, 710, "Print ESC/POS\nThermal Receipt")
    uc("history", 1700, 710, "Sales History,\nReprint, Refund")
    uc("printer_cfg", 1550, 830, "Configure Printer\n(BT / USB / Wi-Fi)")
    uc("receive_goods", 1550, 1000, "Receive Purchased\nGoods", False, 22)

    for key in ["manage_users", "costing", "dashboards", "approve_corr", "review_rts", "inventory"]:
        mod.draw_assoc(draw, exec_pt, boxes[key])
    mod.draw_assoc(draw, hr_pt, boxes["manage_users"])
    for key in ["inventory", "orders", "parcel", "rts", "pos", "mobile_pos", "web_pos"]:
        mod.draw_assoc(draw, staff_pt, boxes[key])
    mod.draw_assoc(draw, cust_pt, boxes["receive_goods"])

    includes = [
        ("manage_users", "assign_roles"),
        ("costing", "formulas"),
        ("inventory", "receive"),
        ("inventory", "honey"),
        ("orders", "prepare"),
        ("orders", "waybills"),
        ("parcel", "scan"),
        ("parcel", "status"),
        ("rts", "rts_case"),
        ("rts", "confirm_stock"),
        ("pos", "web_pos"),
        ("pos", "mobile_pos"),
        ("mobile_pos", "pos_auth"),
        ("mobile_pos", "barcode_sell"),
        ("mobile_pos", "checkout"),
        ("mobile_pos", "print_receipt"),
        ("mobile_pos", "history"),
        ("mobile_pos", "printer_cfg"),
        ("web_pos", "pos_auth"),
        ("web_pos", "checkout"),
    ]
    for a, b in includes:
        mod.draw_dependency(draw, boxes[a], boxes[b], "include")

    extends = [
        ("approve_corr", "inventory"),
        ("review_rts", "rts"),
        ("receive_goods", "checkout"),
        ("history", "checkout"),  # reprint/refund extends completed sale path
    ]
    for a, b in extends:
        mod.draw_dependency(draw, boxes[a], boxes[b], "extend")

    draw.rectangle((40, 1540, 240, 1660), outline=mod.BORDER, width=1)
    draw.text((50, 1548), "Legend", font=mod.font_uc_b, fill=mod.INK)
    draw.line((50, 1580, 100, 1580), fill=mod.INK, width=2)
    draw.text((110, 1572), "Association", font=mod.font_note, fill=mod.INK)
    mod.dash_line(draw, (50, 1605), (100, 1605))
    draw.text((110, 1597), "<<include>>", font=mod.font_note, fill=mod.INK)
    mod.dash_line(draw, (50, 1630), (100, 1630))
    draw.text((110, 1622), "<<extend>>", font=mod.font_note, fill=mod.INK)
    draw.text(
        (280, 1580),
        "Mobile POS evidence: github.com/diobrandedd/POS-PrinterDriver-Mobile (Flutter Android) + RTS pos_login / pos sales API.",
        font=mod.font_note,
        fill=mod.MUTED,
    )

    img.save(out)
    print("regenerated UC with Mobile POS", out, out.stat().st_size)

    # Patch image7 only
    png = out.read_bytes()
    with zipfile.ZipFile(doc_path) as z:
        erd = z.read("word/media/image14.png")
        db = z.read("word/media/image13.png")
    tmp = doc_path.parent / (doc_path.name + ".posuc.tmpzip")
    with zipfile.ZipFile(doc_path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            name = info.filename.replace("\\", "/")
            if name == "word/media/image7.png":
                data = png
                print("replaced image7 Use Case")
            zout.writestr(info, data)
    tmp.replace(doc_path)
    with zipfile.ZipFile(doc_path) as z:
        assert z.read("word/media/image14.png") == erd
        assert z.read("word/media/image13.png") == db
        assert z.read("word/media/image7.png") == png
    print("ERD + DB Design preserved")
    shutil.copy2(out, DESKTOP / "Pyx_Tracker_Use_Case_Diagram.png")


def write_note() -> None:
    NOTE.write_text(
        """# Mobile POS Feature Addition

## Source evidence

- Capstone proposal approved feature: **Point of Sale System**
- Companion Android app: [POS-PrinterDriver-Mobile](https://github.com/diobrandedd/POS-PrinterDriver-Mobile)
- RTS backend support in this repository: `auth.php?action=pos_login` (Bearer), `pos.php` actions `sales_history` / `receipt_get` / `sale_refund`, `pos_mobile_staff`, `sql/migration_phase80_pos_mobile_tokens.sql`

## Documented POS flow (Mobile)

1. Assigned staff signs in on Android Pyx POS (employee credentials / Bearer token against RTS API).
2. Sell: barcode lookup → cart → cash checkout → auto thermal print (ESC/POS).
3. Optional AR credit checkout.
4. Sales history: reprint receipt; full-sale refund updates RTS website/inventory data.
5. Printer connection: Bluetooth / USB / Wi-Fi port 9100.

## Manuscript updates

- Scope feature bullet for Point of Sale (web + Mobile POS)
- Specific objective 6
- Build 6 methodology section
- Technology stack, software/hardware specs, list of modules
- Definition of terms (Mobile POS, Point of Sale, ESC/POS)
- Use Case Diagram Figure 3.5 regenerated with Mobile POS includes
- ERD and Database Design figures were not modified
""",
        encoding="utf-8",
    )
    print("wrote", NOTE)


def main() -> None:
    src = SRC if SRC.exists() else ALT
    if not src.exists():
        raise FileNotFoundError(src)
    DESKTOP.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, WORK)

    with zipfile.ZipFile(WORK) as z:
        erd_before = z.read("word/media/image14.png")
        db_before = z.read("word/media/image13.png")

    doc = Document(str(WORK))
    log = patch_body(doc) + patch_tables(doc)
    doc.save(str(WORK))

    patch_use_case_diagram(WORK)

    with zipfile.ZipFile(WORK) as z:
        assert z.read("word/media/image14.png") == erd_before
        assert z.read("word/media/image13.png") == db_before

    write_note()

    saved = DESKTOP / "Pyx_Tracker_Capstone_SQL_ERD.docx"
    try:
        shutil.copy2(WORK, saved)
        print("saved", saved, saved.stat().st_size)
    except PermissionError:
        saved = ALT
        shutil.copy2(WORK, saved)
        print("primary locked; saved", saved, saved.stat().st_size)
    try:
        shutil.copy2(WORK, ALT)
    except Exception:
        pass

    # verify POS text present
    v = Document(str(saved))
    blob = "\n".join(p.text for p in v.paragraphs)
    assert "Mobile POS" in blob or "Point of Sale System" in blob
    assert "Build 6" in blob
    print("LOG:")
    for line in log:
        print(" -", line)
    print("DONE", saved)


if __name__ == "__main__":
    main()
