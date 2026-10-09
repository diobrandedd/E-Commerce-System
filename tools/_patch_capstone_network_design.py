"""Regenerate Network Design (Figure 3.16) with Cisco-style blue icons.

Replaces only word/media/image15.png.
Does NOT touch Network Topology (image16.png) or other figures.
"""

from __future__ import annotations

import math
import shutil
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document

SRC_CANDIDATES = [
    Path(r"C:/Users/Admin/Desktop/capstone/Pyx_Tracker_Capstone_SQL_ERD_with_POS_diagrams.docx"),
    Path(r"C:/Users/Admin/Desktop/capstone/Pyx_Tracker_Capstone_SQL_ERD.docx"),
    Path(r"C:/Users/Admin/Desktop/capstone/Pyx_Tracker_Capstone_SQL_ERD_updated.docx"),
]
WORK = Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx")
DESKTOP = Path(r"C:/Users/Admin/Desktop/capstone")
OUT_IMG = DESKTOP / "Pyx_Tracker_Network_Design.png"

# Cisco-like palette
BG = "#FFFFFF"
INK = "#1A3A5C"
LINE = "#2B5A8A"
FILL = "#E8F2FC"
FILL2 = "#D4E6F7"
MUTED = "#5A6A7A"
CLOUD = "#B8D4F0"
ARROW = "#4A6A8A"
PILL = "#F3E8D0"

font_title = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 26)
font_lab = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 12)
font_sm = ImageFont.truetype(r"C:/Windows/Fonts/arial.ttf", 11)
font_tiny = ImageFont.truetype(r"C:/Windows/Fonts/arial.ttf", 10)


def tw(draw, text, font):
    b = draw.textbbox((0, 0), text, font=font)
    return b[2] - b[0], b[3] - b[1]


def label_under(draw, cx, y, lines, font=font_lab):
    for i, ln in enumerate(lines):
        w, h = tw(draw, ln, font)
        draw.text((cx - w / 2, y + i * (h + 2)), ln, font=font, fill=INK)


def arrow_line(draw, x1, y1, x2, y2, width=3):
    draw.line((x1, y1, x2, y2), fill=ARROW, width=width)
    ang = math.atan2(y2 - y1, x2 - x1)
    size = 10
    for d in (-0.4, 0.4):
        a = ang + math.pi + d
        draw.line(
            (x2, y2, x2 + size * math.cos(a), y2 + size * math.sin(a)),
            fill=ARROW,
            width=width,
        )


# --- Icon drawers (blue outline / light fill) ---

def icon_pc(draw, cx, cy, scale=1.0):
    """Desktop workstation: monitor + stand + base."""
    s = scale
    # monitor
    draw.rounded_rectangle(
        (cx - 28 * s, cy - 26 * s, cx + 28 * s, cy + 8 * s),
        radius=3,
        fill=FILL,
        outline=LINE,
        width=2,
    )
    draw.rectangle(
        (cx - 22 * s, cy - 20 * s, cx + 22 * s, cy + 2 * s),
        fill="#F7FBFF",
        outline=LINE,
        width=1,
    )
    # stand
    draw.line((cx, cy + 8 * s, cx, cy + 18 * s), fill=LINE, width=2)
    draw.rectangle(
        (cx - 16 * s, cy + 18 * s, cx + 16 * s, cy + 24 * s),
        fill=FILL2,
        outline=LINE,
        width=2,
    )


def icon_phone(draw, cx, cy, scale=1.0):
    s = scale
    draw.rounded_rectangle(
        (cx - 14 * s, cy - 28 * s, cx + 14 * s, cy + 28 * s),
        radius=4,
        fill=FILL,
        outline=LINE,
        width=2,
    )
    draw.rectangle(
        (cx - 10 * s, cy - 22 * s, cx + 10 * s, cy + 16 * s),
        fill="#F7FBFF",
        outline=LINE,
        width=1,
    )
    draw.ellipse(
        (cx - 4 * s, cy + 20 * s, cx + 4 * s, cy + 26 * s),
        outline=LINE,
        width=2,
    )


def icon_users(draw, cx, cy, scale=1.0):
    """Admin / executive as two people (Cisco-ish people symbol simplified)."""
    s = scale
    for dx in (-12 * s, 12 * s):
        draw.ellipse(
            (cx + dx - 8 * s, cy - 22 * s, cx + dx + 8 * s, cy - 6 * s),
            fill=FILL,
            outline=LINE,
            width=2,
        )
        draw.pieslice(
            (cx + dx - 14 * s, cy - 4 * s, cx + dx + 14 * s, cy + 24 * s),
            0,
            180,
            fill=FILL,
            outline=LINE,
            width=2,
        )


def icon_app(draw, cx, cy, scale=1.0):
    """Use-system / application terminal."""
    s = scale
    draw.rounded_rectangle(
        (cx - 26 * s, cy - 22 * s, cx + 26 * s, cy + 22 * s),
        radius=4,
        fill=FILL,
        outline=LINE,
        width=2,
    )
    draw.arc(
        (cx - 16 * s, cy - 14 * s, cx + 16 * s, cy + 10 * s),
        200,
        340,
        fill=LINE,
        width=3,
    )
    draw.rectangle(
        (cx - 8 * s, cy + 2 * s, cx + 8 * s, cy + 14 * s),
        fill=FILL2,
        outline=LINE,
        width=2,
    )


def icon_servers(draw, cx, cy, scale=1.0, n=3):
    s = scale
    total_w = n * 22 * s + (n - 1) * 6 * s
    x0 = cx - total_w / 2
    for i in range(n):
        x = x0 + i * (22 * s + 6 * s)
        draw.rounded_rectangle(
            (x, cy - 28 * s, x + 22 * s, cy + 28 * s),
            radius=3,
            fill=FILL,
            outline=LINE,
            width=2,
        )
        for k in range(3):
            yy = cy - 18 * s + k * 8 * s
            draw.line((x + 4 * s, yy, x + 18 * s, yy), fill=LINE, width=2)


def icon_session_box(draw, cx, cy, text="sess_…php"):
    draw.rounded_rectangle(
        (cx - 48, cy - 20, cx + 48, cy + 20),
        radius=4,
        fill="#FFFFFF",
        outline=LINE,
        width=2,
    )
    w, h = tw(draw, text, font_tiny)
    draw.text((cx - w / 2, cy - h / 2), text, font=font_tiny, fill=INK)


def icon_lock(draw, cx, cy, scale=1.0):
    s = scale
    # shackle
    draw.arc(
        (cx - 12 * s, cy - 28 * s, cx + 12 * s, cy - 4 * s),
        0,
        180,
        fill=LINE,
        width=3,
    )
    draw.rounded_rectangle(
        (cx - 18 * s, cy - 8 * s, cx + 18 * s, cy + 22 * s),
        radius=3,
        fill=FILL,
        outline=LINE,
        width=2,
    )
    draw.ellipse(
        (cx - 5 * s, cy + 2 * s, cx + 5 * s, cy + 12 * s),
        fill=FILL2,
        outline=LINE,
        width=2,
    )


def icon_document(draw, cx, cy, scale=1.0):
    s = scale
    draw.rounded_rectangle(
        (cx - 18 * s, cy - 24 * s, cx + 18 * s, cy + 24 * s),
        radius=2,
        fill=FILL,
        outline=LINE,
        width=2,
    )
    for k in range(4):
        yy = cy - 12 * s + k * 8 * s
        draw.line((cx - 10 * s, yy, cx + 10 * s, yy), fill=LINE, width=1)
    # pen accent
    draw.line(
        (cx + 8 * s, cy - 20 * s, cx + 20 * s, cy - 6 * s),
        fill=LINE,
        width=2,
    )


def icon_network_cloud(draw, cx, cy, scale=1.0):
    """Cloud with Wi-Fi router inside."""
    s = scale
    # cloud blobs
    for ox, oy, r in [(-40, 5, 28), (0, -8, 36), (40, 5, 28), (-15, 18, 24), (15, 18, 24)]:
        draw.ellipse(
            (
                cx + ox * s - r * s,
                cy + oy * s - r * s,
                cx + ox * s + r * s,
                cy + oy * s + r * s,
            ),
            fill=CLOUD,
            outline=LINE,
            width=2,
        )
    # router body
    draw.rounded_rectangle(
        (cx - 22 * s, cy - 6 * s, cx + 22 * s, cy + 16 * s),
        radius=3,
        fill="#FFFFFF",
        outline=LINE,
        width=2,
    )
    # antennas
    draw.line((cx - 10 * s, cy - 6 * s, cx - 10 * s, cy - 22 * s), fill=LINE, width=2)
    draw.line((cx + 10 * s, cy - 6 * s, cx + 10 * s, cy - 22 * s), fill=LINE, width=2)
    draw.ellipse((cx - 12 * s, cy - 26 * s, cx - 8 * s, cy - 22 * s), fill=LINE)
    draw.ellipse((cx + 8 * s, cy - 26 * s, cx + 12 * s, cy - 22 * s), fill=LINE)
    # signal arcs
    for r in (10, 16, 22):
        draw.arc(
            (cx - r * s, cy - 38 * s, cx + r * s, cy - 18 * s),
            220,
            320,
            fill=LINE,
            width=2,
        )


def icon_pill(draw, cx, cy, text="USE SYSTEM"):
    w, h = tw(draw, text, font_sm)
    pad_x, pad_y = 14, 8
    draw.rounded_rectangle(
        (cx - w / 2 - pad_x, cy - h / 2 - pad_y, cx + w / 2 + pad_x, cy + h / 2 + pad_y),
        radius=16,
        fill=PILL,
        outline=LINE,
        width=1,
    )
    draw.text((cx - w / 2, cy - h / 2), text, font=font_sm, fill=INK)


def generate(path: Path) -> None:
    W, H = 1500, 920
    img = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(img)

    title = "Network Design for Pyx Tracker"
    tw0, _ = tw(draw, title, font_title)
    draw.text(((W - tw0) / 2, 18), title, font=font_title, fill=INK)
    draw.text(
        (40, 52),
        "Same logical flow as Figure 3.16. Icons restyled to institutional blue network-equipment notation.",
        font=font_sm,
        fill=MUTED,
    )

    # Top row centers
    y_top = 160
    nodes = [
        (120, "admin"),
        (320, "staff_pc"),
        (500, "use"),
        (700, "servers"),
        (900, "session"),
        (1080, "auth"),
        (1280, "saved"),
    ]
    # draw icons
    icon_users(draw, 120, y_top, 1.1)
    label_under(draw, 120, y_top + 40, ["ADMIN /", "EXECUTIVE"])

    icon_pc(draw, 320, y_top, 1.15)
    label_under(draw, 320, y_top + 40, ["STAFF", "WORKSTATION"])

    # mobile under staff (same topology as original)
    icon_phone(draw, 320, y_top + 130, 0.95)
    label_under(draw, 320, y_top + 168, ["STAFF MOBILE PHONE", "(camera / scan)"])
    arrow_line(draw, 320, y_top + 100, 320, y_top + 35, width=2)

    icon_app(draw, 500, y_top, 1.1)
    label_under(draw, 500, y_top + 40, ["USE SYSTEM", "(orders / POS /", "inventory / RTS)"])

    icon_servers(draw, 700, y_top, 1.0, n=3)
    # small monitor accent in front
    icon_pc(draw, 760, y_top + 10, 0.55)
    label_under(draw, 720, y_top + 45, ["APPLICATION", "SERVER / DATA"])

    icon_session_box(draw, 900, y_top)
    label_under(draw, 900, y_top + 36, ["SESSION / API"])

    icon_lock(draw, 1080, y_top, 1.15)
    label_under(draw, 1080, y_top + 40, ["AUTH /", "ENCRYPTION"])

    icon_document(draw, 1280, y_top, 1.15)
    label_under(draw, 1280, y_top + 40, ["SAVED RECORD"])

    # top flow arrows
    pairs = [(150, 280), (360, 450), (560, 640), (780, 850), (960, 1030), (1140, 1230)]
    for x1, x2 in pairs:
        arrow_line(draw, x1, y_top, x2, y_top, width=3)

    # Network cloud center
    cloud_cx, cloud_cy = 750, 430
    arrow_line(draw, 1280, y_top + 70, cloud_cx + 180, cloud_cy - 40, width=3)
    # elbow down-left into cloud
    draw.line((1280, y_top + 70, 1280, 360), fill=ARROW, width=3)
    draw.line((1280, 360, cloud_cx + 90, cloud_cy - 20), fill=ARROW, width=3)
    arrow_line(draw, cloud_cx + 90, cloud_cy - 20, cloud_cx + 55, cloud_cy - 5, width=3)

    icon_network_cloud(draw, cloud_cx, cloud_cy, 1.35)
    label_under(draw, cloud_cx, cloud_cy + 70, ["NETWORK", "(Router / Wi-Fi access point)"])

    # Bottom three branches
    bottoms = [
        (280, "admin"),
        (750, "staff_pc"),
        (1220, "phone"),
    ]
    for bx, kind in bottoms:
        # down from cloud
        arrow_line(draw, cloud_cx + (bx - cloud_cx) * 0.15, cloud_cy + 55, bx, 560, width=3)
        # session/browser document
        icon_document(draw, bx, 600, 0.9)
        w1, _ = tw(draw, "[SESSION]", font_tiny)
        w2, _ = tw(draw, "[BROWSER]", font_tiny)
        draw.text((bx - w1 / 2, 635), "[SESSION]", font=font_tiny, fill=MUTED)
        draw.text((bx - w2 / 2, 648), "[BROWSER]", font=font_tiny, fill=MUTED)

        arrow_line(draw, bx, 670, bx, 720, width=2)

        if kind == "admin":
            icon_users(draw, bx, 760, 1.0)
            label_under(draw, bx, 800, ["ADMIN /", "EXECUTIVE"])
        elif kind == "staff_pc":
            icon_pc(draw, bx, 760, 1.05)
            label_under(draw, bx, 800, ["STAFF", "WORKSTATION"])
        else:
            icon_phone(draw, bx, 760, 1.0)
            label_under(draw, bx, 800, ["STAFF MOBILE", "PHONE"])

        icon_pill(draw, bx + 110, 760, "USE SYSTEM")
        arrow_line(draw, bx + 40, 760, bx + 70, 760, width=2)

    draw.text(
        (40, H - 36),
        "Topology unchanged: users → system → server/API/auth → saved record → WLAN hub → browser clients. Icon style aligned to blue network-equipment sample.",
        font=font_sm,
        fill=MUTED,
    )
    img.save(path)
    print("wrote", path, img.size, path.stat().st_size)


def patch_docx(doc_path: Path, png: Path) -> None:
    data_png = png.read_bytes()
    with zipfile.ZipFile(doc_path) as z:
        topo_before = z.read("word/media/image16.png")
        erd = z.read("word/media/image14.png")
        db = z.read("word/media/image13.png")

    tmp = doc_path.parent / (doc_path.name + ".netdesign.tmpzip")
    replaced = 0
    with zipfile.ZipFile(doc_path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            name = info.filename.replace("\\", "/")
            if name == "word/media/image15.png":
                data = data_png
                replaced += 1
                print("replaced Network Design", name)
            zout.writestr(info, data)
    if replaced != 1:
        raise RuntimeError(f"expected 1 replace, got {replaced}")
    try:
        tmp.replace(doc_path)
    except PermissionError:
        blob = tmp.read_bytes()
        tmp.unlink(missing_ok=True)
        doc_path.write_bytes(blob)

    with zipfile.ZipFile(doc_path) as z:
        assert z.read("word/media/image16.png") == topo_before, "Topology changed!"
        assert z.read("word/media/image14.png") == erd
        assert z.read("word/media/image13.png") == db
        assert z.read("word/media/image15.png") == data_png
    print("Network Topology preserved OK")


def patch_caption(doc_path: Path) -> None:
    d = Document(str(doc_path))

    def set_text(p, text):
        for r in p.runs:
            r.text = ""
        (p.runs[0] if p.runs else p.add_run()).text = text

    for p in d.paragraphs:
        t = p.text.strip()
        if t.startswith("A network design outlines the architecture"):
            set_text(
                p,
                "A network design outlines the architecture that supports communication between "
                "devices used with Pyx Tracker. Figure 3.16 keeps the same logical WLAN design: "
                "Admin/Executive and Staff clients (workstation and mobile) use the system through "
                "the application server, session/API, and authentication layers; saved records are "
                "reached through the central router / Wi-Fi access point; clients reopen the system "
                "in a browser session. Device icons are drawn in blue network-equipment notation "
                "for consistency with institutional sample diagrams.",
            )
        if t.strip() == "This figure shows the network design.":
            set_text(
                p,
                "Figure 3.16 presents the Network Design of Pyx Tracker. The Network Topology "
                "diagram that follows is unchanged.",
            )
    d.save(str(doc_path))
    print("updated Network Design caption")


def main() -> None:
    src = next((p for p in SRC_CANDIDATES if p.exists()), None)
    if src is None:
        raise FileNotFoundError("No source DOCX")
    print("SRC", src)
    DESKTOP.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, WORK)

    generate(OUT_IMG)
    patch_docx(WORK, OUT_IMG)
    patch_caption(WORK)

    targets = [
        DESKTOP / "Pyx_Tracker_Capstone_SQL_ERD.docx",
        DESKTOP / "Pyx_Tracker_Capstone_SQL_ERD_with_POS_diagrams.docx",
        DESKTOP / "Pyx_Tracker_Capstone_SQL_ERD_updated.docx",
    ]
    saved = None
    for t in targets:
        try:
            shutil.copy2(WORK, t)
            print("saved", t, t.stat().st_size)
            saved = t
        except PermissionError:
            print("locked", t.name)
    if saved is None:
        alt = DESKTOP / "Pyx_Tracker_Capstone_SQL_ERD_netdesign.docx"
        shutil.copy2(WORK, alt)
        saved = alt
        print("fallback", alt)

    with zipfile.ZipFile(saved) as z:
        print(
            "verify",
            "design",
            len(z.read("word/media/image15.png")),
            "topo",
            len(z.read("word/media/image16.png")),
        )
    print("DONE", saved)


if __name__ == "__main__":
    main()
