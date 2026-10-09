"""Add POS Device leaf to Network Topology (Figure 3.17 / image16.png).

Does NOT modify Network Design (image15.png).
"""

from __future__ import annotations

import math
import shutil
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document

SRC = Path(r"C:/Users/Admin/Desktop/capstone/Pyx_Tracker_Capstone_SQL_ERD.docx")
WORK = Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD_topo_work.docx")
DESKTOP = Path(r"C:/Users/Admin/Desktop/capstone")
OUT_IMG = DESKTOP / "Pyx_Tracker_Network_Topology.png"

BG = "#FFFFFF"
INK = "#1A1A1A"
LINE = "#333333"
BOX = "#F7F7F7"
BORDER = "#444444"
MUTED = "#555555"
ACCENT = "#2F6FAE"

font_title = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 28)
font_lab = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 13)
font_sm = ImageFont.truetype(r"C:/Windows/Fonts/arial.ttf", 11)
font_tiny = ImageFont.truetype(r"C:/Windows/Fonts/arial.ttf", 10)


def tw(d, t, f):
    b = d.textbbox((0, 0), t, font=f)
    return b[2] - b[0], b[3] - b[1]


def center_text(d, cx, cy, text, font, fill=INK):
    lines = text.split("\n")
    heights = [tw(d, ln, font)[1] for ln in lines]
    total = sum(heights) + 2 * (len(lines) - 1)
    y = cy - total / 2
    for ln, h in zip(lines, heights):
        w, _ = tw(d, ln, font)
        d.text((cx - w / 2, y), ln, font=font, fill=fill)
        y += h + 2


def icon_router(d, cx, cy, s=1.0):
    d.rounded_rectangle((cx - 36 * s, cy - 16 * s, cx + 36 * s, cy + 16 * s), 4, fill="#E8F2FC", outline=ACCENT, width=2)
    d.line((cx - 12 * s, cy - 16 * s, cx - 12 * s, cy - 30 * s), fill=ACCENT, width=2)
    d.line((cx + 12 * s, cy - 16 * s, cx + 12 * s, cy - 30 * s), fill=ACCENT, width=2)
    for r in (8, 14, 20):
        d.arc((cx - r * s, cy - 42 * s, cx + r * s, cy - 22 * s), 220, 320, fill=ACCENT, width=2)


def icon_pc(d, cx, cy, s=1.0):
    d.rounded_rectangle((cx - 22 * s, cy - 20 * s, cx + 10 * s, cy + 6 * s), 3, fill="#F0F6FB", outline=BORDER, width=2)
    d.rectangle((cx - 18 * s, cy - 16 * s, cx + 6 * s, cy + 2 * s), fill="#FFFFFF", outline=BORDER, width=1)
    d.rounded_rectangle((cx + 12 * s, cy - 18 * s, cx + 26 * s, cy + 16 * s), 2, fill="#F0F6FB", outline=BORDER, width=2)
    d.rectangle((cx - 24 * s, cy + 10 * s, cx + 10 * s, cy + 16 * s), fill="#E0E0E0", outline=BORDER, width=1)


def icon_phone(d, cx, cy, s=1.0):
    d.rounded_rectangle((cx - 14 * s, cy - 26 * s, cx + 14 * s, cy + 26 * s), 4, fill="#F0F6FB", outline=BORDER, width=2)
    d.rectangle((cx - 10 * s, cy - 20 * s, cx + 10 * s, cy + 14 * s), fill="#FFFFFF", outline=BORDER, width=1)
    d.ellipse((cx - 4 * s, cy + 18 * s, cx + 4 * s, cy + 24 * s), outline=BORDER, width=2)


def icon_server(d, cx, cy, s=1.0):
    d.rounded_rectangle((cx - 24 * s, cy - 30 * s, cx + 24 * s, cy + 30 * s), 3, fill="#F0F6FB", outline=BORDER, width=2)
    for k in range(4):
        y = cy - 22 * s + k * 12 * s
        d.rectangle((cx - 18 * s, y, cx + 18 * s, y + 8 * s), fill="#FFFFFF", outline=BORDER, width=1)


def icon_pos(d, cx, cy, s=1.0):
    # phone with small POS mark
    icon_phone(d, cx, cy, s)
    d.rounded_rectangle((cx - 8 * s, cy - 8 * s, cx + 8 * s, cy + 4 * s), 2, fill="#D9ECFA", outline=ACCENT, width=1)
    center_text(d, cx, cy - 2 * s, "POS", font_tiny, fill=ACCENT)


def leaf_box(d, cx, cy, w, h, draw_icon, title):
    d.rounded_rectangle((cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2), 10, fill=BOX, outline=BORDER, width=2)
    draw_icon(d, cx, cy - 18, 1.0)
    center_text(d, cx, cy + 28, title, font_lab)


def generate(path: Path) -> None:
    W, H = 1100, 780
    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)

    title = "Network Topology"
    tw0, _ = tw(d, title, font_title)
    d.text(((W - tw0) / 2, 20), title, font=font_title, fill=INK)
    d.text(
        (40, 58),
        "Figure 3.17 style: logical star for local access. POS Device added as a client leaf; printer/scanner stay in Network Design.",
        font=font_sm,
        fill=MUTED,
    )

    cx, cy = W / 2, 400
    # hub
    hub_w, hub_h = 260, 160
    d.rounded_rectangle((cx - hub_w / 2, cy - hub_h / 2, cx + hub_w / 2, cy + hub_h / 2), 12, fill="#FFFFFF", outline=BORDER, width=2)
    icon_router(d, cx, cy - 28, 1.2)
    center_text(d, cx, cy + 28, "Router / access point", font_lab)
    center_text(d, cx, cy + 50, "Logical local access topology", font_tiny, fill=MUTED)
    center_text(d, cx, cy + 64, "Remote hosting requires an internet connection.", font_tiny, fill=MUTED)

    # Five leaves around the hub (star)
    # angles: top-left, top-right, bottom-right, bottom-left, mid-right (POS)
    leaves = [
        (210, 180, icon_pc, "Admin / Executive\nworkstation"),
        (890, 180, icon_pc, "Staff / CSR\nworkstation"),
        (890, 620, icon_server, "Application\nserver"),
        (210, 620, icon_phone, "Staff mobile\nphone"),
        (890, 400, icon_pos, "POS Device\n(Android Pyx POS)"),
    ]

    for lx, ly, icon_fn, title_t in leaves:
        # line from hub edge toward leaf
        dx, dy = lx - cx, ly - cy
        dist = math.hypot(dx, dy) or 1
        # start near hub box
        sx = cx + dx / dist * (hub_w / 2 - 10)
        sy = cy + dy / dist * (hub_h / 2 - 10)
        # end near leaf box
        leaf_w, leaf_h = 200, 120
        ex = lx - dx / dist * (leaf_w / 2 - 8)
        ey = ly - dy / dist * (leaf_h / 2 - 8)
        d.line((sx, sy, ex, ey), fill=LINE, width=2)
        leaf_box(d, lx, ly, leaf_w, leaf_h, icon_fn, title_t)

    d.text(
        (40, H - 40),
        "Star topology: all local clients attach logically to the site router/AP. POS Device is a Mobile POS client; peripherals are shown only in Network Design.",
        font=font_sm,
        fill=MUTED,
    )
    img.save(path)
    print("wrote", path, img.size, path.stat().st_size)


def patch_docx(doc_path: Path, png: Path) -> None:
    png_bytes = png.read_bytes()
    with zipfile.ZipFile(doc_path) as z:
        design_before = z.read("word/media/image15.png")
    tmp = doc_path.parent / (doc_path.name + ".topo.tmpzip")
    with zipfile.ZipFile(doc_path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            name = info.filename.replace("\\", "/")
            if name == "word/media/image16.png":
                data = png_bytes
                print("replaced image16 Network Topology")
            zout.writestr(info, data)
    try:
        tmp.replace(doc_path)
    except PermissionError:
        doc_path.write_bytes(tmp.read_bytes())
        tmp.unlink(missing_ok=True)
    with zipfile.ZipFile(doc_path) as z:
        assert z.read("word/media/image15.png") == design_before
        assert z.read("word/media/image16.png") == png_bytes
    print("Network Design preserved")


def patch_text(doc_path: Path) -> None:
    doc = Document(str(doc_path))

    def set_text(p, text):
        for r in p.runs:
            r.text = ""
        (p.runs[0] if p.runs else p.add_run()).text = text

    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Figure 3.17"):
            set_text(p, "Figure 3.17  Star Topology for Local Access")
        if t.startswith("Clients connect through a router") or t.startswith("Clients connect through a router or access point"):
            set_text(
                p,
                "Clients connect through a router or access point in a logical star. End nodes include "
                "Admin/Executive and Staff/CSR workstations, Staff mobile phones, the Android POS Device "
                "(Mobile POS), and the application server. A remote or hosted server still requires an "
                "internet uplink, as shown in the Network Design. Thermal printer and barcode scanner "
                "peripherals are documented in Network Design only, not as separate topology nodes.",
            )
    doc.save(str(doc_path))
    print("caption updated")


def main() -> None:
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    DESKTOP.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, WORK)
    generate(OUT_IMG)
    patch_docx(WORK, OUT_IMG)
    patch_text(WORK)

    for name in [
        "Pyx_Tracker_Capstone_SQL_ERD.docx",
        "Pyx_Tracker_Capstone_SQL_ERD_updated.docx",
        "Pyx_Tracker_Capstone_SQL_ERD_netdesign.docx",
        "Pyx_Tracker_Capstone_SQL_ERD_with_POS_diagrams.docx",
    ]:
        dest = DESKTOP / name
        try:
            shutil.copy2(WORK, dest)
            print("saved", dest.name, dest.stat().st_size)
        except Exception as e:
            print("skip", name, type(e).__name__)

    with zipfile.ZipFile(DESKTOP / "Pyx_Tracker_Capstone_SQL_ERD.docx") as z:
        print("verify design", len(z.read("word/media/image15.png")), "topo", len(z.read("word/media/image16.png")))
    print("DONE")


if __name__ == "__main__":
    main()
