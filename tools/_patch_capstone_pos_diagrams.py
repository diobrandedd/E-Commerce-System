"""Update Agile model (Fig 3.1), Gantt (Fig 3.3), and FDD (Fig 3.4) to include POS.

Replaces:
- word/media/image3.png  -> Incremental Agile (6 builds)
- word/media/image5.png  -> Gantt with Feature 6 POS
- word/media/image6.png  -> FDD with Point of Sale branch

Preserves ERD (image14) and Database Design (image13).
"""

from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document

SRC = Path(r"C:/Users/Admin/Desktop/capstone/Pyx_Tracker_Capstone_SQL_ERD.docx")
WORK = Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx")
DESKTOP = Path(r"C:/Users/Admin/Desktop/capstone")
OUT_AGILE = DESKTOP / "Pyx_Tracker_Agile_Model.png"
OUT_GANTT = DESKTOP / "Pyx_Tracker_Gantt_Chart.png"
OUT_FDD = DESKTOP / "Pyx_Tracker_FDD.png"

BG = "#FFFFFF"
INK = "#1A1A1A"
MUTED = "#555555"
BORDER = "#333333"
BLUE = "#4A7BB7"
BLUE_DK = "#2F5F8F"
GREEN = "#5B8C5A"
BROWN = "#8B6914"
HEADER = "#F5F5F5"
TAN = "#F5E6C8"
REQ = "#E8EEF5"

font_title = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 28)
font_sub = ImageFont.truetype(r"C:/Windows/Fonts/arial.ttf", 14)
font_h = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 13)
font_b = ImageFont.truetype(r"C:/Windows/Fonts/arial.ttf", 12)
font_sm = ImageFont.truetype(r"C:/Windows/Fonts/arial.ttf", 11)
font_tiny = ImageFont.truetype(r"C:/Windows/Fonts/arial.ttf", 10)


def tw(draw, text, font):
    b = draw.textbbox((0, 0), text, font=font)
    return b[2] - b[0], b[3] - b[1]


def rounded(draw, box, fill, outline=BORDER, width=2, radius=8):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def center_text(draw, cx, cy, text, font, fill=INK):
    w, h = tw(draw, text, font)
    draw.text((cx - w / 2, cy - h / 2), text, font=font, fill=fill)


def multiline_center(draw, cx, cy, lines, font, fill=INK, gap=2):
    heights = [tw(draw, ln, font)[1] for ln in lines]
    total = sum(heights) + gap * (len(lines) - 1)
    y = cy - total / 2
    for ln, h in zip(lines, heights):
        w, _ = tw(draw, ln, font)
        draw.text((cx - w / 2, y), ln, font=font, fill=fill)
        y += h + gap


# ---------------------------------------------------------------------------
# Figure 3.1 — Incremental Agile (6 builds)
# ---------------------------------------------------------------------------
def generate_agile(path: Path) -> None:
    builds = [
        "Build 1\nInventory",
        "Build 2\nCosting and pricing",
        "Build 3\nOrders and waybills",
        "Build 4\nParcel processing",
        "Build 5\nReturn-to-Sender",
        "Build 6\nPoint of Sale\n(Web + Mobile POS)",
    ]
    W, H = 1600, 980
    img = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(img)

    title = "Incremental Agile Software Development Model"
    tw0, _ = tw(draw, title, font_title)
    draw.text(((W - tw0) / 2, 24), title, font=font_title, fill=INK)
    draw.text(
        (40, 64),
        "Six feature builds. Each build: Design and development → Testing → Implementation, with feedback for refinement.",
        font=font_sub,
        fill=MUTED,
    )

    # Requirements box
    rx1, ry1, rx2, ry2 = 40, 380, 200, 480
    rounded(draw, (rx1, ry1, rx2, ry2), REQ, width=2, radius=6)
    center_text(draw, (rx1 + rx2) / 2, (ry1 + ry2) / 2, "Requirements", font_h)

    stages = ["Design and\ndevelopment", "Testing", "Implementation"]
    row_h = 120
    top = 120
    left0 = 280
    box_w, box_h = 150, 56
    gap_x = 40

    for i, label in enumerate(builds):
        cy = top + i * row_h + 40
        # connector from requirements
        draw.line((rx2, (ry1 + ry2) / 2, left0 - 10, cy), fill=BORDER, width=2)

        # build label left of first stage
        lines = label.split("\n")
        multiline_center(draw, left0 - 90, cy, lines, font_tiny, fill=BLUE_DK)

        xs = []
        for s, stage in enumerate(stages):
            x1 = left0 + s * (box_w + gap_x)
            y1 = cy - box_h / 2
            x2 = x1 + box_w
            y2 = y1 + box_h
            xs.append((x1, y1, x2, y2))
            rounded(draw, (x1, y1, x2, y2), "#FAFBFC", width=2, radius=4)
            multiline_center(draw, (x1 + x2) / 2, (y1 + y2) / 2, stage.split("\n"), font_sm)
            if s > 0:
                px2 = xs[s - 1][2]
                draw.line((px2, cy, x1, cy), fill=BORDER, width=2)
                # arrow
                draw.polygon([(x1, cy), (x1 - 8, cy - 5), (x1 - 8, cy + 5)], fill=BORDER)

        # feedback dashed arrow Implementation -> Design
        x_imp = xs[2][2]
        x_des = xs[0][0]
        y_fb = cy + box_h / 2 + 18
        # dashed
        x = x_imp
        on = True
        while x > x_des:
            nx = max(x - 8, x_des)
            if on:
                draw.line((x, y_fb, nx, y_fb), fill=MUTED, width=1)
            x = nx - 4
            on = not on
        draw.line((x_imp, cy + box_h / 2, x_imp, y_fb), fill=MUTED, width=1)
        draw.line((x_des, y_fb, x_des, cy + box_h / 2), fill=MUTED, width=1)
        draw.polygon([(x_des, cy + box_h / 2), (x_des - 5, cy + box_h / 2 + 8), (x_des + 5, cy + box_h / 2 + 8)], fill=MUTED)

    draw.text(
        (40, H - 40),
        "Source: approved Agile model; Build 6 covers web POS and Android Mobile POS (Pyx POS) synchronized to RTS.",
        font=font_sm,
        fill=MUTED,
    )
    img.save(path)
    print("agile", path, img.size, path.stat().st_size)


# ---------------------------------------------------------------------------
# Figure 3.3 — Gantt (sample style + Feature 6)
# ---------------------------------------------------------------------------
def generate_gantt(path: Path) -> None:
    # weeks: May1-4, Jun1-4, Jul1-4, Aug1-4  => 16 columns
    features = [
        {
            "name": "Feature 1 — Integrated Inventory Management",
            "design": (0, 4),  # May 1-4
            "test": (4, 5),  # Jun 1
            "impl": (5, 6),  # Jun 2
        },
        {
            "name": "Feature 2 — Product Costing and Pricing",
            "design": (2, 6),  # May 3-4, Jun 1-2
            "test": (6, 7),
            "impl": (7, 8),
        },
        {
            "name": "Feature 3 — Multi-Platform Orders and Waybills",
            "design": (4, 10),  # Jun 1-4, Jul 1-2
            "test": (10, 11),
            "impl": (11, 12),
        },
        {
            "name": "Feature 4 — Parcel Processing and Inventory Sync",
            "design": (6, 12),  # Jun 3-4, Jul 1-4
            "test": (12, 13),
            "impl": (13, 14),
        },
        {
            "name": "Feature 5 — Return-to-Sender Management",
            "design": (8, 14),  # Jul 1-4, Aug 1-2
            "test": (14, 15),
            "impl": (15, 16),
        },
        {
            "name": "Feature 6 — Point of Sale (Web and Mobile POS)",
            "design": (6, 12),  # Jun 3–Jul 4 (overlap with fulfillment; Mobile POS API/app)
            "test": (12, 13),  # Aug 1
            "impl": (13, 14),  # Aug 2
        },
    ]

    label_w = 420
    col_w = 48
    n_cols = 16
    header_h = 70
    row_h = 78
    top = 110
    left = 40
    W = left + label_w + n_cols * col_w + 40
    H = top + header_h + len(features) * row_h + 90
    img = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(img)

    title = "System Development Schedule | May–August"
    tw0, _ = tw(draw, title, font_title)
    draw.text(((W - tw0) / 2, 28), title, font=font_title, fill=INK)

    # legend
    lx = left
    for color, lab in [(BLUE, "Design and Development"), (GREEN, "Testing"), (BROWN, "Implementation")]:
        draw.rectangle((lx, 70, lx + 18, 88), fill=color, outline=BORDER)
        draw.text((lx + 24, 70), lab, font=font_sm, fill=INK)
        lx += 200

    # header months
    chart_x = left + label_w
    months = [("May", 0, 4), ("June", 4, 4), ("July", 8, 4), ("August", 12, 4)]
    draw.rectangle((left, top, chart_x + n_cols * col_w, top + header_h), outline=BORDER, width=2)
    draw.rectangle((left, top, chart_x, top + header_h), fill=HEADER, outline=BORDER, width=2)
    center_text(draw, left + label_w / 2, top + header_h / 2, "Features and activities", font_h)

    for name, start, span in months:
        x1 = chart_x + start * col_w
        x2 = x1 + span * col_w
        draw.rectangle((x1, top, x2, top + 32), fill=HEADER, outline=BORDER, width=1)
        center_text(draw, (x1 + x2) / 2, top + 16, name, font_h)
        for w in range(span):
            wx1 = x1 + w * col_w
            wx2 = wx1 + col_w
            draw.rectangle((wx1, top + 32, wx2, top + header_h), outline=BORDER, width=1)
            center_text(draw, (wx1 + wx2) / 2, top + 32 + 19, str(w + 1), font_sm)

    phase_color = {"design": BLUE, "test": GREEN, "impl": BROWN}
    phase_label = {
        "design": "Design and Development",
        "test": "Testing",
        "impl": "Implementation",
    }

    for i, feat in enumerate(features):
        y1 = top + header_h + i * row_h
        y2 = y1 + row_h
        draw.rectangle((left, y1, chart_x + n_cols * col_w, y2), outline=BORDER, width=1)
        draw.rectangle((left, y1, chart_x, y2), fill="#FAFAFA", outline=BORDER, width=1)
        # feature name
        draw.text((left + 10, y1 + 8), feat["name"], font=font_h, fill=INK)
        # three mini rows for phases
        sub_h = (row_h - 28) / 3
        for pi, phase in enumerate(["design", "test", "impl"]):
            sy = y1 + 28 + pi * sub_h
            draw.text((left + 16, sy + 2), phase_label[phase], font=font_tiny, fill=MUTED)
            a, b = feat[phase]
            bx1 = chart_x + a * col_w + 3
            bx2 = chart_x + b * col_w - 3
            by1 = sy + 2
            by2 = sy + sub_h - 4
            if bx2 > bx1:
                draw.rectangle((bx1, by1, bx2, by2), fill=phase_color[phase], outline=BORDER)

        # grid lines
        for c in range(n_cols + 1):
            x = chart_x + c * col_w
            draw.line((x, y1, x, y2), fill="#DDDDDD", width=1)

    draw.text(
        (left, H - 36),
        "Feature 6 adds web POS and Android Mobile POS (barcode sell, AR, ESC/POS print, history/refund) synchronized to RTS.",
        font=font_sm,
        fill=MUTED,
    )
    img.save(path)
    print("gantt", path, img.size, path.stat().st_size)


# ---------------------------------------------------------------------------
# Figure 3.4 — FDD with 6 modules
# ---------------------------------------------------------------------------
def generate_fdd(path: Path) -> None:
    groups = [
        (
            "Integrated Inventory\nManagement",
            ["Receive materials", "Track production", "Confirm finished stock", "Review movements"],
        ),
        (
            "Product Costing\nand Pricing",
            ["Maintain formulas", "Calculate costs", "Define parcel recipes", "Maintain prices"],
        ),
        (
            "Multi-Platform Orders\nand Waybills",
            ["Prepare order files", "Validate order data", "Stamp waybills", "Finalize batches"],
        ),
        (
            "Parcel Processing and\nInventory Sync",
            ["Record packing scans", "Confirm movements", "Save shipment photos", "Review parcel history"],
        ),
        (
            "Return-to-Sender\nManagement",
            ["Record return case", "Classify condition", "Review stock-in", "Confirm sellable stock"],
        ),
        (
            "Point of Sale\n(Web + Mobile POS)",
            ["Authenticate POS staff", "Barcode sell / AR checkout", "Print ESC/POS receipt", "History, reprint, refund"],
        ),
    ]

    n = len(groups)
    col_w = 250
    gap = 28
    left = 50
    W = left * 2 + n * col_w + (n - 1) * gap
    root_y = 100
    group_y = 220
    child_h = 44
    child_gap = 12
    H = group_y + 70 + 4 * (child_h + child_gap) + 80
    img = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(img)

    title = "Functional Decomposition Diagram"
    tw0, _ = tw(draw, title, font_title)
    draw.text(((W - tw0) / 2, 24), title, font=font_title, fill=INK)

    # root
    rw, rh = 220, 50
    rcx = W / 2
    rounded(draw, (rcx - rw / 2, root_y, rcx + rw / 2, root_y + rh), "#FFFFFF", width=2, radius=6)
    center_text(draw, rcx, root_y + rh / 2, "Pyx Tracker", font_h)

    # rail under root
    centers = []
    for i in range(n):
        cx = left + i * (col_w + gap) + col_w / 2
        centers.append(cx)
    rail_y = root_y + rh + 30
    draw.line((centers[0], rail_y, centers[-1], rail_y), fill=BORDER, width=2)
    draw.line((rcx, root_y + rh, rcx, rail_y), fill=BORDER, width=2)

    for i, ((gname, children), cx) in enumerate(zip(groups, centers)):
        draw.line((cx, rail_y, cx, group_y), fill=BORDER, width=2)
        # group box
        gx1, gy1 = cx - col_w / 2, group_y
        gx2, gy2 = cx + col_w / 2, group_y + 58
        rounded(draw, (gx1, gy1, gx2, gy2), TAN, width=2, radius=4)
        multiline_center(draw, cx, (gy1 + gy2) / 2, gname.split("\n"), font_sm)

        # children
        prev_y = gy2
        for j, child in enumerate(children):
            cy1 = gy2 + 18 + j * (child_h + child_gap)
            cy2 = cy1 + child_h
            # stem
            draw.line((cx, prev_y if j == 0 else cy1 - child_gap, cx, cy1), fill=BORDER, width=1)
            rounded(draw, (gx1 + 8, cy1, gx2 - 8, cy2), "#FFFFFF", width=1, radius=4)
            # wrap child text
            words = child.split()
            if len(child) > 22:
                mid = len(words) // 2
                lines = [" ".join(words[:mid]), " ".join(words[mid:])]
            else:
                lines = [child]
            multiline_center(draw, cx, (cy1 + cy2) / 2, lines, font_tiny)
            prev_y = cy2

    draw.text(
        (left, H - 36),
        "Figure 3.4 style: system → six modules → subfunctions. POS includes web checkout and Android Mobile POS flows.",
        font=font_sm,
        fill=MUTED,
    )
    img.save(path)
    print("fdd", path, img.size, path.stat().st_size)


def patch_docx(doc_path: Path, agile: Path, gantt: Path, fdd: Path) -> None:
    mapping = {
        "word/media/image3.png": agile.read_bytes(),
        "word/media/image5.png": gantt.read_bytes(),
        "word/media/image6.png": fdd.read_bytes(),
    }
    with zipfile.ZipFile(doc_path) as z:
        erd = z.read("word/media/image14.png")
        db = z.read("word/media/image13.png")
        uc = z.read("word/media/image7.png")

    tmp = doc_path.parent / (doc_path.name + ".fddgantt.tmpzip")
    with zipfile.ZipFile(doc_path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            name = info.filename.replace("\\", "/")
            if name in mapping:
                data = mapping[name]
                print("replaced", name)
            zout.writestr(info, data)
    tmp.replace(doc_path)

    with zipfile.ZipFile(doc_path) as z:
        assert z.read("word/media/image14.png") == erd
        assert z.read("word/media/image13.png") == db
        assert z.read("word/media/image7.png") == uc
        for k, v in mapping.items():
            assert z.read(k) == v
    print("preserved ERD, DB Design, Use Case")


def patch_text(doc_path: Path) -> None:
    d = Document(str(doc_path))

    def set_text(p, text):
        for r in p.runs:
            r.text = ""
        (p.runs[0] if p.runs else p.add_run()).text = text

    for p in d.paragraphs:
        t = p.text.strip()
        if t.startswith("The study uses the Agile Software Development Model"):
            set_text(
                p,
                "The study uses the Agile Software Development Model identified in the proposal. Agile "
                "principles emphasize working software, collaboration, responsiveness to requirements, and "
                "regular reflection. Development is organized into six increments. Each build covers design "
                "and development, testing, and implementation or refinement. Shared authentication and data "
                "definitions support every increment, including the Point of Sale build for web POS and the "
                "Android Mobile POS application.",
            )
        if t.startswith("The diagram follows the sample's requirements-to-build structure"):
            set_text(
                p,
                "The diagram follows the sample's requirements-to-build structure, expanded to the six "
                "modules (Inventory, Costing and pricing, Orders and waybills, Parcel processing, "
                "Return-to-Sender, and Point of Sale — Web and Mobile POS). Return arrows indicate feedback "
                "and revision. This sequence matches the Build sections that follow.",
            )
        if t.startswith("System development starts in May and ends in August"):
            set_text(
                p,
                "System development starts in May and ends in August. Each of the six features has Design "
                "and Development, Testing, and Implementation activities. Feature 6 covers the Point of Sale "
                "System, including the web POS and the Android Mobile POS terminal (barcode sell, AR credit, "
                "ESC/POS printing, sales history, reprint, and refund) synchronized to RTS. The four numbered "
                "intervals in each month organize the schedule. Activity bars show a proposed allocation for "
                "planning and documentation; they do not replace recorded test evidence.",
            )
        if t.startswith("The FDD decomposes the root into five modules"):
            set_text(
                p,
                "The FDD decomposes the root into six modules and their subfunctions, following the sample's "
                "system-to-function-to-subfunction hierarchy. It represents Integrated Inventory Management, "
                "Product Costing and Pricing, Multi-Platform Orders and Waybills, Parcel Processing and "
                "Inventory Sync, Return-to-Sender Management, and Point of Sale (Web and Mobile POS). POS "
                "subfunctions cover staff authentication, barcode sell and AR checkout, ESC/POS receipt "
                "printing, and sales history with reprint and refund.",
            )
        # Fix Build 6 body wrongly set as Heading 3
        if t.startswith("Design and development implement walk-in sales"):
            try:
                p.style = d.styles["Normal"]
            except Exception:
                pass
        if t.startswith("Testing verifies successful login for authorized Mobile POS"):
            try:
                p.style = d.styles["Normal"]
            except Exception:
                pass

    d.save(str(doc_path))
    print("updated Agile/Gantt/FDD captions and Build 6 styles")


def main() -> None:
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    DESKTOP.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, WORK)

    generate_agile(OUT_AGILE)
    generate_gantt(OUT_GANTT)
    generate_fdd(OUT_FDD)

    patch_docx(WORK, OUT_AGILE, OUT_GANTT, OUT_FDD)
    patch_text(WORK)

    desk = DESKTOP / "Pyx_Tracker_Capstone_SQL_ERD.docx"
    try:
        shutil.copy2(WORK, desk)
        print("saved", desk, desk.stat().st_size)
    except PermissionError:
        alt = DESKTOP / "Pyx_Tracker_Capstone_SQL_ERD_updated.docx"
        shutil.copy2(WORK, alt)
        print("primary locked; saved", alt, alt.stat().st_size)
        desk = alt

    # verify
    v = Document(str(desk))
    blob = "\n".join(p.text for p in v.paragraphs)
    assert "six increments" in blob or "six modules" in blob
    assert "Feature 6" in blob or "Build 6" in blob
    with zipfile.ZipFile(desk) as z:
        print(
            "sizes",
            len(z.read("word/media/image3.png")),
            len(z.read("word/media/image5.png")),
            len(z.read("word/media/image6.png")),
            "erd",
            len(z.read("word/media/image14.png")),
        )
    print("DONE")


if __name__ == "__main__":
    main()
