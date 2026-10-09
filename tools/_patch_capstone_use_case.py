"""Generate Pyx Tracker Use Case Diagram in institutional UML sample style.

Visual conventions match the S.I.G.N. sample:
- System boundary rectangle with title
- Stick-figure actors outside left/right
- Oval use cases
- Solid association lines actor→use case
- Dashed <<include>> / <<extend>> arrows

Patches only Figure 3.5 media (image7.png). Does not touch ERD or Database Design.
"""

from __future__ import annotations

import math
import shutil
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

SRC_CANDIDATES = [
    Path(r"C:/Users/Admin/Desktop/capstone/Pyx_Tracker_Capstone_SQL_ERD_updated.docx"),
    Path(r"C:/Users/Admin/Desktop/capstone/Pyx_Tracker_Capstone_SQL_ERD.docx"),
    Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx"),
]
WORK_DOC = Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx")
OUT_IMG = Path(r"C:/xampp/htdocs/rts/output/documents/_use_case_pyx_tracker.png")
DESKTOP_DIR = Path(r"C:/Users/Admin/Desktop/capstone")
DESKTOP_DOC = DESKTOP_DIR / "Pyx_Tracker_Capstone_SQL_ERD.docx"
DESKTOP_ALT = DESKTOP_DIR / "Pyx_Tracker_Capstone_SQL_ERD_updated.docx"
DESKTOP_PNG = DESKTOP_DIR / "Pyx_Tracker_Use_Case_Diagram.png"

BG = "#FFFFFF"
INK = "#1A1A1A"
MUTED = "#555555"
BORDER = "#333333"
OVAL_FILL = "#FFFFFF"
OVAL_OUTLINE = "#222222"
DASH = "#444444"

font_title = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 22)
font_boundary = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 16)
font_actor = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 14)
font_uc = ImageFont.truetype(r"C:/Windows/Fonts/arial.ttf", 12)
font_uc_b = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 12)
font_stereo = ImageFont.truetype(r"C:/Windows/Fonts/ariali.ttf", 10)
font_note = ImageFont.truetype(r"C:/Windows/Fonts/arial.ttf", 12)
font_mod = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 11)


def text_size(draw, text, font):
    b = draw.textbbox((0, 0), text, font=font)
    return b[2] - b[0], b[3] - b[1]


def wrap_lines(text: str, max_chars: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    cur = ""
    for w in words:
        trial = (cur + " " + w).strip()
        if len(trial) <= max_chars:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines or [text]


def draw_actor(draw, cx, cy, label: str):
    """Classic UML stick figure; returns (cx, cy) association anchor."""
    r = 14
    draw.ellipse((cx - r, cy - 58, cx + r, cy - 58 + 2 * r), outline=INK, width=2)
    draw.line((cx, cy - 58 + 2 * r, cx, cy + 8), fill=INK, width=2)
    draw.line((cx - 28, cy - 20, cx + 28, cy - 20), fill=INK, width=2)
    draw.line((cx, cy + 8, cx - 22, cy + 48), fill=INK, width=2)
    draw.line((cx, cy + 8, cx + 22, cy + 48), fill=INK, width=2)
    for i, line in enumerate(label.split("\n")):
        tw, _ = text_size(draw, line, font_actor)
        draw.text((cx - tw / 2, cy + 54 + i * 16), line, font=font_actor, fill=INK)
    return (cx, cy - 10)


def draw_oval(draw, cx, cy, label: str, bold: bool = False, max_chars: int = 22):
    lines = wrap_lines(label, max_chars)
    font = font_uc_b if bold else font_uc
    max_w = 0
    total_h = 0
    line_sizes = []
    for line in lines:
        tw, th = text_size(draw, line, font)
        line_sizes.append((tw, th))
        max_w = max(max_w, tw)
        total_h += th + 2
    pad_x, pad_y = 18, 10
    rw = max(70, max_w / 2 + pad_x)
    rh = max(22, total_h / 2 + pad_y)
    box = (cx - rw, cy - rh, cx + rw, cy + rh)
    draw.ellipse(box, fill=OVAL_FILL, outline=OVAL_OUTLINE, width=2)
    y = cy - total_h / 2
    for (tw, th), line in zip(line_sizes, lines):
        draw.text((cx - tw / 2, y), line, font=font, fill=INK)
        y += th + 2
    return box  # x1,y1,x2,y2


def oval_edge(box, toward):
    """Point on ellipse rim toward an external point."""
    x1, y1, x2, y2 = box
    cx, cy = (x1 + x2) / 2, (y1 + y2) / 2
    rx, ry = (x2 - x1) / 2, (y2 - y1) / 2
    tx, ty = toward
    dx, dy = tx - cx, ty - cy
    if dx == 0 and dy == 0:
        return (cx + rx, cy)
    # parametric scale onto ellipse
    scale = 1.0 / math.sqrt((dx * dx) / (rx * rx) + (dy * dy) / (ry * ry))
    return (cx + dx * scale, cy + dy * scale)


def draw_assoc(draw, p0, box):
    p1 = oval_edge(box, p0)
    draw.line((p0[0], p0[1], p1[0], p1[1]), fill=INK, width=1)


def dash_line(draw, p0, p1, dash=7, gap=5):
    x0, y0 = p0
    x1, y1 = p1
    length = math.hypot(x1 - x0, y1 - y0) or 1
    dx, dy = (x1 - x0) / length, (y1 - y0) / length
    pos = 0.0
    on = True
    while pos < length:
        seg = dash if on else gap
        end = min(pos + seg, length)
        if on:
            draw.line(
                (x0 + dx * pos, y0 + dy * pos, x0 + dx * end, y0 + dy * end),
                fill=DASH,
                width=1,
            )
        pos = end
        on = not on


def arrow_head(draw, tip, tail):
    ang = math.atan2(tip[1] - tail[1], tip[0] - tail[0])
    size = 9
    for delta in (-0.4, 0.4):
        a = ang + math.pi + delta
        draw.line(
            (tip[0], tip[1], tip[0] + size * math.cos(a), tip[1] + size * math.sin(a)),
            fill=DASH,
            width=1,
        )


def draw_dependency(draw, box_from, box_to, stereotype: str):
    """Dashed arrow from base toward included/extended target (UML: arrow points to base for include? 

    UML convention:
    - <<include>>: dashed arrow from base use case TO the included use case
    - <<extend>>: dashed arrow from extending use case TO the base use case

    Sample S.I.G.N. labels arrows near the relationship. We draw arrow toward `box_to`.
    """
    c_from = ((box_from[0] + box_from[2]) / 2, (box_from[1] + box_from[3]) / 2)
    c_to = ((box_to[0] + box_to[2]) / 2, (box_to[1] + box_to[3]) / 2)
    p0 = oval_edge(box_from, c_to)
    p1 = oval_edge(box_to, c_from)
    dash_line(draw, p0, p1)
    arrow_head(draw, p1, p0)
    mx, my = (p0[0] + p1[0]) / 2, (p0[1] + p1[1]) / 2
    label = f"<<{stereotype}>>"
    tw, th = text_size(draw, label, font_stereo)
    draw.rectangle((mx - tw / 2 - 2, my - th / 2 - 8, mx + tw / 2 + 2, my + th / 2 - 6), fill=BG)
    draw.text((mx - tw / 2, my - th / 2 - 7), label, font=font_stereo, fill=INK)


def generate(path: Path) -> None:
    W, H = 2200, 1600
    img = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(img)

    # Page title
    title = "Use Case Diagram for Pyx Tracker"
    tw, _ = text_size(draw, title, font_title)
    draw.text(((W - tw) / 2, 18), title, font=font_title, fill=INK)
    draw.text(
        (40, 50),
        "Web-based Enterprise Operations Management System  |  Solid = association  |  Dashed <<include>> / <<extend>> = UML dependency",
        font=font_note,
        fill=MUTED,
    )

    # System boundary
    bx1, by1, bx2, by2 = 260, 90, 1940, 1520
    draw.rectangle((bx1, by1, bx2, by2), outline=BORDER, width=3)
    btitle = "Pyx Tracker: Web-based Enterprise Operations Management System"
    btw, _ = text_size(draw, btitle, font_boundary)
    draw.text((bx1 + (bx2 - bx1 - btw) / 2, by1 + 10), btitle, font=font_boundary, fill=INK)

    # Actors
    exec_pt = draw_actor(draw, 110, 280, "Executive")
    hr_pt = draw_actor(draw, 110, 620, "HR\nPersonnel")
    staff_pt = draw_actor(draw, 2090, 520, "Staff")
    # Staff note (CSR / Production / POS / Delivery)
    note = "(CSR / Production /\nPOS / Delivery)"
    for i, line in enumerate(note.split("\n")):
        tw, _ = text_size(draw, line, font_note)
        draw.text((2090 - tw / 2, 620 + i * 14), line, font=font_note, fill=MUTED)
    cust_pt = draw_actor(draw, 2090, 1180, "Customer")
    draw.text((2000, 1285), "(beneficiary)", font=font_note, fill=MUTED)

    boxes: dict[str, tuple] = {}

    def uc(key, x, y, label, bold=False, max_chars=20):
        boxes[key] = draw_oval(draw, x, y, label, bold=bold, max_chars=max_chars)
        return boxes[key]

    # ----- Module band labels (inside boundary, subtle) -----
    for label, x, y in [
        ("User & Access", 380, 115),
        ("Costing", 700, 115),
        ("Inventory", 1050, 115),
        ("Orders", 1400, 115),
        ("Fulfillment / POS / RTS", 1700, 115),
    ]:
        draw.text((x, y), label, font=font_mod, fill="#2F80ED")

    # ===== Left / Executive–HR cluster =====
    uc("manage_users", 420, 220, "Manage Users\nand Access", bold=True)
    uc("assign_roles", 420, 330, "Assign Roles\nand Permissions")
    uc("activate_accounts", 420, 430, "Activate /\nDeactivate Accounts")

    uc("costing", 420, 580, "Product Costing\nand Pricing", bold=True)
    uc("formulas", 420, 700, "Maintain Cost\nFormulas")
    uc("cogs", 420, 800, "Calculate\nManufacturing COGS")
    uc("pack_tiers", 420, 900, "Define Parcel\nPack Tiers")
    uc("prices", 420, 1000, "Maintain Selling\nPrices")

    uc("dashboards", 420, 1160, "Monitor Dashboards\nand Reports", bold=True)
    uc("approve_corr", 420, 1280, "Approve Inventory\nCorrections")
    uc("review_rts", 420, 1400, "Review RTS Claims\nand Dashboard")

    # ===== Center inventory / orders =====
    uc("inventory", 900, 260, "Integrated Inventory\nManagement", bold=True, max_chars=24)
    uc("receive", 780, 390, "Receive Materials\nand Stock-In")
    uc("honey", 1020, 390, "Track Honey\nProduction Pipeline")
    uc("finished", 780, 510, "Confirm Finished\nGoods Stock")
    uc("movements", 1020, 510, "Review Stock\nMovements")

    uc("orders", 900, 680, "Multi-Platform Order\nProcessing", bold=True, max_chars=24)
    uc("prepare", 780, 810, "Prepare VIP /\nShopee SPX Orders")
    uc("waybills", 1020, 810, "Match and Stamp\nWaybills")
    uc("finalize", 900, 930, "Finalize Order\nBatches")

    # ===== Right fulfillment / POS / RTS =====
    uc("parcel", 1500, 240, "Parcel Processing\nand Inventory Sync", bold=True, max_chars=24)
    uc("scan", 1380, 370, "Scan Tracking\nNumbers")
    uc("deduct", 1620, 370, "Deduct Finished Goods\nand Packing Materials")
    uc("photos", 1380, 500, "Capture Before\nShipping Photos")
    uc("status", 1620, 500, "Update Parcel Status\n(To Pack → Shipped)")

    uc("pos", 1500, 680, "Point of Sale\nSystem", bold=True)
    uc("pin", 1380, 800, "Authenticate with\nPOS PIN")
    uc("sale", 1620, 800, "Record Walk-in\nSale")
    uc("receipt", 1380, 920, "Generate Receipt\nand History")
    uc("rts_stock", 1620, 920, "Select RTS vs\nNon-RTS Stock")

    uc("rts", 1500, 1100, "Return-to-Sender\nManagement", bold=True)
    uc("rts_case", 1380, 1230, "Record RTS Case\nand Line Items")
    uc("classify", 1620, 1230, "Classify Sellable /\nDamaged Condition")
    uc("sync_pool", 1380, 1360, "Synchronize\nInventory Pools")
    uc("confirm_stock", 1620, 1360, "Confirm Sellable\nStock Restoration")

    # Indirect customer benefit (secondary)
    uc("receive_goods", 1750, 1480, "Receive Ordered /\nPurchased Goods", max_chars=26)

    # ===== Associations =====
    for key in ["manage_users", "costing", "dashboards", "approve_corr", "review_rts", "inventory"]:
        draw_assoc(draw, exec_pt, boxes[key])
    for key in ["manage_users", "activate_accounts"]:
        draw_assoc(draw, hr_pt, boxes[key])
    for key in [
        "inventory",
        "orders",
        "parcel",
        "pos",
        "rts",
        "receive",
        "prepare",
        "scan",
        "sale",
        "rts_case",
    ]:
        draw_assoc(draw, staff_pt, boxes[key])
    draw_assoc(draw, cust_pt, boxes["receive_goods"])

    # ===== <<include>> (base → included) =====
    includes = [
        ("manage_users", "assign_roles"),
        ("manage_users", "activate_accounts"),
        ("costing", "formulas"),
        ("costing", "cogs"),
        ("costing", "pack_tiers"),
        ("costing", "prices"),
        ("inventory", "receive"),
        ("inventory", "honey"),
        ("inventory", "finished"),
        ("inventory", "movements"),
        ("orders", "prepare"),
        ("orders", "waybills"),
        ("orders", "finalize"),
        ("parcel", "scan"),
        ("parcel", "deduct"),
        ("parcel", "photos"),
        ("parcel", "status"),
        ("pos", "pin"),
        ("pos", "sale"),
        ("pos", "receipt"),
        ("rts", "rts_case"),
        ("rts", "classify"),
        ("rts", "sync_pool"),
        ("rts", "confirm_stock"),
    ]
    for a, b in includes:
        draw_dependency(draw, boxes[a], boxes[b], "include")

    # ===== <<extend>> (extending → base) =====
    extends = [
        ("approve_corr", "movements"),  # executive correction approval extends movement review
        ("review_rts", "rts"),  # executive oversight extends RTS module
        ("rts_stock", "sale"),  # optional RTS stock source extends sale
        ("costing", "prices"),  # already included; skip duplicate
        ("dashboards", "inventory"),  # monitoring extends operational inventory view
        ("receive_goods", "status"),  # customer receipt of goods extends shipped status outcome
    ]
    # Remove the bogus costing->prices extend (already include)
    extends = [e for e in extends if e != ("costing", "prices")]
    for a, b in extends:
        draw_dependency(draw, boxes[a], boxes[b], "extend")

    # Legend
    draw.rectangle((40, 1480, 240, 1585), outline=BORDER, width=1)
    draw.text((50, 1488), "Legend", font=font_uc_b, fill=INK)
    draw.line((50, 1520, 100, 1520), fill=INK, width=2)
    draw.text((110, 1512), "Association", font=font_note, fill=INK)
    dash_line(draw, (50, 1545), (100, 1545))
    draw.text((110, 1537), "<<include>>", font=font_note, fill=INK)
    dash_line(draw, (50, 1570), (100, 1570))
    draw.text((110, 1562), "<<extend>>", font=font_note, fill=INK)

    draw.text(
        (280, 1545),
        "Actors follow the approved proposal (Executive, HR Personnel, Staff, Customer). Staff covers CSR, production, POS, and delivery duties under role permissions.",
        font=font_note,
        fill=MUTED,
    )

    img.save(path)
    print(f"wrote {path} {img.size} {path.stat().st_size} bytes")


def patch_use_case_only(doc_path: Path, png: Path) -> dict:
    """Replace image7.png only; verify image13 (DB design) and image14 (ERD) unchanged."""
    png_bytes = png.read_bytes()
    before = {}
    with zipfile.ZipFile(doc_path) as z:
        for name in ("word/media/image7.png", "word/media/image13.png", "word/media/image14.png"):
            if name in z.namelist():
                before[name] = z.read(name)

    tmp = doc_path.parent / (doc_path.name + ".uc.tmpzip")
    replaced = 0
    with zipfile.ZipFile(doc_path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            name = info.filename.replace("\\", "/")
            if name == "word/media/image7.png":
                data = png_bytes
                replaced += 1
                print("replaced Use Case media:", name)
            zout.writestr(info, data)
    if replaced != 1:
        raise RuntimeError(f"Expected 1 image7 replacement, got {replaced}")
    try:
        tmp.replace(doc_path)
    except PermissionError:
        blob = tmp.read_bytes()
        tmp.unlink(missing_ok=True)
        doc_path.write_bytes(blob)

    after = {}
    with zipfile.ZipFile(doc_path) as z:
        for name in ("word/media/image7.png", "word/media/image13.png", "word/media/image14.png"):
            if name in z.namelist():
                after[name] = z.read(name)

    if before.get("word/media/image13.png") != after.get("word/media/image13.png"):
        raise RuntimeError("Database Design image13.png changed unexpectedly")
    if before.get("word/media/image14.png") != after.get("word/media/image14.png"):
        raise RuntimeError("ERD image14.png changed unexpectedly")
    if after["word/media/image7.png"] != png_bytes:
        raise RuntimeError("Use Case image7.png not updated")
    print("preserved Database Design + ERD: OK")
    return {"image7": len(after["word/media/image7.png"]), "image13": len(after.get("word/media/image13.png", b"")), "image14": len(after.get("word/media/image14.png", b""))}


def update_caption(doc_path: Path) -> None:
    from docx import Document

    d = Document(str(doc_path))
    for p in d.paragraphs:
        t = p.text.strip()
        if t.startswith("Figure 3.5"):
            new = "Figure 3.5  Overall Use Case Diagram of Pyx Tracker"
            for r in p.runs:
                r.text = ""
            if p.runs:
                p.runs[0].text = new
            else:
                p.add_run(new)
        # Optional explanatory paragraph immediately after figure if present
        if t.startswith("Figure 3.5 presents") or (
            "Overall Use Case" in t and t.startswith("The")
        ):
            pass
    # Insert/update description after Figure 3.5 caption if next para is Use Case Matrix
    paras = list(d.paragraphs)
    for i, p in enumerate(paras):
        if p.text.strip().startswith("Figure 3.5  Overall Use Case"):
            # If next is Use Case Matrix, we may leave matrix as-is; caption is enough.
            # Add description into a following short para if it already explains the figure.
            if i + 1 < len(paras):
                nxt = paras[i + 1].text.strip()
                if nxt == "Use Case Matrix" or nxt.startswith("Table 3.1"):
                    # leave structure; description can sit in caption area via prior narrative
                    pass
                elif "use case" in nxt.lower() or nxt.startswith("Figure 3.5 presents"):
                    new = (
                        "Figure 3.5 presents the overall Use Case Diagram of Pyx Tracker using UML notation "
                        "consistent with the institutional sample. Actors outside the system boundary are "
                        "Executive, HR Personnel, Staff (CSR, production, POS, and delivery duties), and Customer "
                        "as beneficiary. Solid lines show associations. Dashed arrows mark <<include>> "
                        "sub-flows for inventory, costing, multi-platform orders, parcel processing, POS, and "
                        "RTS, and <<extend>> relationships for optional or oversight behaviors such as "
                        "inventory-correction approval, RTS stock selection at POS, and executive RTS review."
                    )
                    for r in paras[i + 1].runs:
                        r.text = ""
                    if paras[i + 1].runs:
                        paras[i + 1].runs[0].text = new
                    else:
                        paras[i + 1].add_run(new)
            break
    d.save(str(doc_path))
    print("updated Figure 3.5 caption")


def save_desktop(work: Path) -> Path:
    DESKTOP_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT_IMG, DESKTOP_PNG)
    try:
        shutil.copy2(work, DESKTOP_DOC)
        print(f"saved {DESKTOP_DOC} ({DESKTOP_DOC.stat().st_size} bytes)")
        # keep updated copy in sync
        try:
            shutil.copy2(work, DESKTOP_ALT)
        except Exception:
            pass
        return DESKTOP_DOC
    except PermissionError:
        shutil.copy2(work, DESKTOP_ALT)
        print(f"primary locked; saved {DESKTOP_ALT} ({DESKTOP_ALT.stat().st_size} bytes)")
        return DESKTOP_ALT


def main() -> None:
    src = next((p for p in SRC_CANDIDATES if p.exists()), None)
    if src is None:
        raise FileNotFoundError("No source DOCX found")
    print("SRC", src)
    shutil.copy2(src, WORK_DOC)

    generate(OUT_IMG)
    sizes = patch_use_case_only(WORK_DOC, OUT_IMG)
    update_caption(WORK_DOC)
    saved = save_desktop(WORK_DOC)
    print("media sizes", sizes)
    print("PNG", DESKTOP_PNG, DESKTOP_PNG.stat().st_size)
    print("DOCX", saved, saved.stat().st_size)
    print("DONE")


if __name__ == "__main__":
    main()
