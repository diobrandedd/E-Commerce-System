"""Create Pyx Tracker Network Design (icon/arrow style) and replace Figure 3.16 media."""

from __future__ import annotations

import shutil
import zipfile
from collections import deque
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageFilter

ASSETS = Path(r"C:/Users/Admin/.cursor/projects/c-xampp-htdocs-rts/assets")
ICONS = {
    "server": ASSETS
    / "c__Users_Admin_AppData_Roaming_Cursor_User_workspaceStorage_c7c434809b0e3f6b99f4071ead395cf1_images_image-a95ef318-edef-4c76-a546-2d6e727bc9d6.png",
    "mobile": ASSETS
    / "c__Users_Admin_AppData_Roaming_Cursor_User_workspaceStorage_c7c434809b0e3f6b99f4071ead395cf1_images_image-58211f65-cc43-4953-b9f0-7c53bc55eed7.png",
    "staff": ASSETS
    / "c__Users_Admin_AppData_Roaming_Cursor_User_workspaceStorage_c7c434809b0e3f6b99f4071ead395cf1_images_image-ca41de92-4028-4dbd-92ed-3335586114aa.png",
    "admin": ASSETS
    / "c__Users_Admin_AppData_Roaming_Cursor_User_workspaceStorage_c7c434809b0e3f6b99f4071ead395cf1_images_image-06d1e19a-b3a5-4c8b-ac47-8405cc284175.png",
    "router": ASSETS
    / "c__Users_Admin_AppData_Roaming_Cursor_User_workspaceStorage_c7c434809b0e3f6b99f4071ead395cf1_images_image-4539b8ca-1aaa-4455-88b8-82934bf69d44.png",
}

OUT = Path(r"C:/xampp/htdocs/rts/output/documents/_network_design_pyx.png")
MEDIA = "word/media/image15.png"
DOCS = [
    Path(r"C:/Users/Admin/Desktop/capstone/Pyx_Tracker_Capstone_SQL_ERD.docx"),
    Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx.new"),
    Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx"),
]


def load_icon(path: Path, size: int = 90, phone: bool = False, crop_bottom: float = 0.0) -> Image.Image:
    im = Image.open(path).convert("RGBA")
    w, h = im.size
    if crop_bottom:
        im = im.crop((0, 0, w, int(h * (1 - crop_bottom))))
    if phone:
        # keep black handset visible on white tile
        canvas = Image.new("RGBA", (size, size), (255, 255, 255, 255))
        im.thumbnail((size - 8, size - 8), Image.Resampling.LANCZOS)
        ox = (size - im.size[0]) // 2
        oy = (size - im.size[1]) // 2
        canvas.paste(im, (ox, oy), im)
        return canvas
    # transparent bg for line icons
    px = im.load()
    ww, hh = im.size
    visited = [[False] * ww for _ in range(hh)]
    q = deque([(0, 0), (ww - 1, 0), (0, hh - 1), (ww - 1, hh - 1)])
    # only clear near-white corners for line icons; skip for solid black silhouettes
    # detect if mostly white background
    corner = sum(im.getpixel(p)[0] + im.getpixel(p)[1] + im.getpixel(p)[2] for p in [(0, 0), (ww - 1, 0)]) / 2
    if corner > 600:
        while q:
            x, y = q.popleft()
            if x < 0 or y < 0 or x >= ww or y >= hh or visited[y][x]:
                continue
            r, g, b, _a = px[x, y]
            if r + g + b < 720:
                continue
            visited[y][x] = True
            px[x, y] = (255, 255, 255, 0)
            q.extend([(x + 1, y), (x - 1, y), (x, y + 1), (x, y - 1)])
    im.thumbnail((size, size), Image.Resampling.LANCZOS)
    canvas = Image.new("RGBA", (size, size), (255, 255, 255, 0))
    ox = (size - im.size[0]) // 2
    oy = (size - im.size[1]) // 2
    canvas.paste(im, (ox, oy), im)
    return canvas


def arrow(draw: ImageDraw.ImageDraw, x0, y0, x1, y1, width=3, color=(90, 90, 90)):
    draw.line((x0, y0, x1, y1), fill=color, width=width)
    # arrow head
    import math

    ang = math.atan2(y1 - y0, x1 - x0)
    s = 12
    for dlt in (-0.4, 0.4):
        a = ang + 3.14159 + dlt
        draw.line((x1, y1, x1 + s * math.cos(a), y1 + s * math.sin(a)), fill=color, width=width)


def thick_arrow(draw, x0, y0, x1, y1, color=(120, 120, 120)):
    # blocky arrow like sample
    import math

    ang = math.atan2(y1 - y0, x1 - x0)
    length = math.hypot(x1 - x0, y1 - y0)
    if length < 1:
        return
    ux, uy = (x1 - x0) / length, (y1 - y0) / length
    px, py = -uy, ux
    body = length - 22
    half = 7
    tip = 18
    pts = [
        (x0 + px * half, y0 + py * half),
        (x0 + ux * body + px * half, y0 + uy * body + py * half),
        (x0 + ux * body + px * tip, y0 + uy * body + py * tip),
        (x1, y1),
        (x0 + ux * body - px * tip, y0 + uy * body - py * tip),
        (x0 + ux * body - px * half, y0 + uy * body - py * half),
        (x0 - px * half, y0 - py * half),
    ]
    draw.polygon(pts, fill=color)


def label_under(draw, cx, y, text, font, fill=(30, 30, 30)):
    lines = text.split("\n")
    yy = y
    for line in lines:
        b = draw.textbbox((0, 0), line, font=font)
        tw = b[2] - b[0]
        draw.text((cx - tw / 2, yy), line, font=font, fill=fill)
        yy += b[3] - b[1] + 2


def draw_cloud(img: Image.Image, cx, cy, w=280, h=140):
    """Simple blue network cloud."""
    overlay = Image.new("RGBA", img.size, (0, 0, 0, 0))
    d = ImageDraw.Draw(overlay)
    # overlapping ellipses
    color = (120, 180, 230, 230)
    outline = (60, 120, 180, 255)
    ellipses = [
        (cx - w // 2, cy - 20, cx - w // 2 + 120, cy + 70),
        (cx - 40, cy - h // 2, cx + 80, cy + 40),
        (cx + 20, cy - 30, cx + w // 2, cy + 65),
        (cx - 90, cy - 5, cx + 90, cy + h // 2),
    ]
    for e in ellipses:
        d.ellipse(e, fill=color, outline=outline, width=2)
    img.paste(Image.alpha_composite(img.convert("RGBA"), overlay).convert("RGB"))


def main() -> None:
    admin = load_icon(ICONS["admin"], 88)
    staff = load_icon(ICONS["staff"], 88, crop_bottom=0.08)
    mobile = load_icon(ICONS["mobile"], 88, phone=True)
    server = load_icon(ICONS["server"], 95)
    router = load_icon(ICONS["router"], 100)

    W, H = 1400, 980
    img = Image.new("RGB", (W, H), (255, 255, 255))
    draw = ImageDraw.Draw(img)

    font_title = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 30)
    font_lab = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 12)
    font_small = ImageFont.truetype(r"C:/Windows/Fonts/arial.ttf", 12)
    font_tag = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 11)
    font_cloud = ImageFont.truetype(r"C:/Windows/Fonts/arialbd.ttf", 22)

    draw.text((40, 24), "Network Design", font=font_title, fill=(25, 25, 25))

    # ---- TOP FLOW: users / devices -> application server -> services ----
    top_y = 130
    steps = [
        (160, admin, "ADMIN /\nEXECUTIVE\nWORKSTATION"),
        (400, staff, "STAFF\nWORKSTATION"),
        (640, mobile, "STAFF\nMOBILE\nPHONE"),
        (880, server, "APPLICATION\nSERVER"),
    ]
    for i, (x, icon, lab) in enumerate(steps):
        # card
        draw.rounded_rectangle((x - 70, top_y - 20, x + 70, top_y + 150), radius=8, outline=(50, 50, 50), width=2)
        img.paste(icon, (x - icon.size[0] // 2, top_y), icon)
        label_under(draw, x, top_y + 100, lab, font_lab)
        if i < len(steps) - 1:
            thick_arrow(draw, x + 78, top_y + 50, steps[i + 1][0] - 78, top_y + 50, color=(150, 150, 150))

    # service badges after server
    draw.rounded_rectangle((1000, top_y + 10, 1180, top_y + 55), radius=6, outline=(40, 100, 180), width=2, fill=(230, 242, 255))
    draw.text((1015, top_y + 22), "React + PHP APIs", font=font_small, fill=(20, 60, 120))
    draw.rounded_rectangle((1000, top_y + 70, 1180, top_y + 115), radius=6, outline=(40, 120, 70), width=2, fill=(232, 248, 236))
    draw.text((1025, top_y + 82), "MySQL + files", font=font_small, fill=(20, 80, 40))
    thick_arrow(draw, 960, top_y + 50, 995, top_y + 50, color=(150, 150, 150))
    thick_arrow(draw, 960, top_y + 90, 995, top_y + 90, color=(150, 150, 150))

    # big arrow down into network
    thick_arrow(draw, 1090, top_y + 125, 1090, 360, color=(130, 130, 130))
    thick_arrow(draw, 880, top_y + 160, 1000, 360, color=(130, 130, 130))

    # ---- CENTER: NETWORK cloud + router ----
    cloud_cx, cloud_cy = W // 2, 460
    # cloud
    overlay = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay)
    cloud_fill = (140, 195, 235, 220)
    cloud_line = (50, 110, 170, 255)
    for e in [
        (cloud_cx - 200, cloud_cy - 40, cloud_cx - 40, cloud_cy + 90),
        (cloud_cx - 100, cloud_cy - 90, cloud_cx + 80, cloud_cy + 50),
        (cloud_cx - 20, cloud_cy - 50, cloud_cx + 200, cloud_cy + 95),
        (cloud_cx - 130, cloud_cy - 10, cloud_cx + 130, cloud_cy + 110),
    ]:
        od.ellipse(e, fill=cloud_fill, outline=cloud_line, width=3)
    img = Image.alpha_composite(img.convert("RGBA"), overlay).convert("RGB")
    draw = ImageDraw.Draw(img)

    # router icon on cloud
    img.paste(router, (cloud_cx - router.size[0] // 2, cloud_cy - 55), router)
    nb = draw.textbbox((0, 0), "NETWORK", font=font_cloud)
    draw.text((cloud_cx - (nb[2] - nb[0]) / 2, cloud_cy + 50), "NETWORK", font=font_cloud, fill=(20, 50, 100))
    rb = draw.textbbox((0, 0), "Router / Wi-Fi access point", font=font_small)
    draw.text((cloud_cx - (rb[2] - rb[0]) / 2, cloud_cy + 78), "Router / Wi-Fi access point", font=font_small, fill=(40, 40, 40))

    # ---- BOTTOM: three client branches ----
    bottoms = [
        (250, admin, "ADMIN / EXECUTIVE\nWORKSTATION", "Access dashboards,\napprovals, costing"),
        (700, staff, "STAFF\nWORKSTATION", "Orders, inventory,\nPOS, RTS"),
        (1150, mobile, "STAFF\nMOBILE PHONE", "Browser access &\ncamera / scan"),
    ]
    by = 780
    for cx, icon, lab, note in bottoms:
        # document/session chip above (like signed document in sample)
        draw.rounded_rectangle((cx - 55, 600, cx + 55, 640), radius=5, outline=(80, 80, 80), width=2, fill=(250, 250, 250))
        draw.text((cx - 48, 610), "SESSION / API", font=font_tag, fill=(40, 40, 40))
        thick_arrow(draw, cloud_cx + (cx - cloud_cx) * 0.25, cloud_cy + 115, cx, 595, color=(140, 140, 140))
        # tags
        draw.text((cx - 90, 650), "[WIFI]", font=font_tag, fill=(0, 100, 160))
        draw.text((cx + 20, 650), "[BROWSER]", font=font_tag, fill=(0, 100, 160))
        thick_arrow(draw, cx, 645, cx, 700, color=(140, 140, 140))
        # endpoint card
        draw.rounded_rectangle((cx - 95, 700, cx + 95, 920), radius=8, outline=(40, 40, 40), width=2)
        img.paste(icon, (cx - icon.size[0] // 2, 715), icon)
        label_under(draw, cx, 815, lab, font_lab)
        label_under(draw, cx, 870, note, font_small, fill=(70, 70, 70))

    # footer
    foot = (
        "WLAN-style logical design: client devices connect through the company router / Wi-Fi access point "
        "to the application server (React interface, PHP APIs, MySQL, and files)."
    )
    draw.text((40, H - 40), foot, font=font_small, fill=(80, 80, 80))

    img.save(OUT)
    print("wrote", OUT, img.size, OUT.stat().st_size)

    png = OUT.read_bytes()
    for doc in DOCS:
        if not doc.exists():
            continue
        tmp = doc.parent / (doc.name + ".tmp_net")
        try:
            with zipfile.ZipFile(doc, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
                for info in zin.infolist():
                    data = zin.read(info.filename)
                    if info.filename.replace("\\", "/") == MEDIA:
                        data = png
                        print("replaced", MEDIA, "in", doc.name)
                    zout.writestr(info, data)
            try:
                tmp.replace(doc)
                print("saved", doc)
            except PermissionError:
                try:
                    doc.write_bytes(tmp.read_bytes())
                    print("saved write", doc)
                except PermissionError:
                    print("LOCKED", doc)
                tmp.unlink(missing_ok=True)
        except Exception as e:
            print("ERR", doc, e)


from docx import Document  # noqa: E402


def update_captions(doc_path: Path) -> None:
    d = Document(str(doc_path))
    changed = False
    for p in d.paragraphs:
        t = p.text.strip()
        if t.startswith("Figure 3.16"):
            new = "Figure 3.16  Network Design of Pyx Tracker"
            for r in p.runs:
                r.text = ""
            if p.runs:
                p.runs[0].text = new
            else:
                p.add_run(new)
            changed = True
        if t.startswith("Browser clients connect"):
            new = (
                "A network design outlines the architecture that supports communication between devices. "
                "The diagram depicts a Wireless LAN (WLAN) style configuration in which Admin/Executive "
                "workstations, Staff workstations, and Staff mobile phones connect through a central "
                "router / Wi-Fi access point to the application server. The server hosts the React interface "
                "and PHP session APIs, with MySQL and controlled file storage for operational records."
            )
            for r in p.runs:
                r.text = ""
            if p.runs:
                p.runs[0].text = new
            else:
                p.add_run(new)
            changed = True
        if t.startswith("The diagram does not assert"):
            new = (
                "This figure shows the logical network design. Exact hosting address, router model, and "
                "installed hardware require deployment records."
            )
            for r in p.runs:
                r.text = ""
            if p.runs:
                p.runs[0].text = new
            else:
                p.add_run(new)
            changed = True
    if not changed:
        return
    tmp = doc_path.parent / (doc_path.name + ".tmp_cap")
    d.save(str(tmp))
    try:
        tmp.replace(doc_path)
        print("captions updated", doc_path.name)
    except PermissionError:
        try:
            doc_path.write_bytes(tmp.read_bytes())
            print("captions write", doc_path.name)
        except PermissionError:
            print("captions LOCKED", doc_path.name)
        tmp.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
    for p in DOCS:
        if p.exists():
            try:
                update_captions(p)
            except Exception as e:
                print("caption err", p, e)
    desk, new = DOCS[0], DOCS[1]
    if new.exists():
        try:
            shutil.copy2(new, desk)
            print("desktop synced")
        except Exception as e:
            print("desk sync", e)
