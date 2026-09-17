"""Build Use Case Matrix exactly like coconut sample (2-col, blue header, actor sections)."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DESK = Path(r"C:/Users/Admin/Desktop/capstone/Pyx_Tracker_Capstone_SQL_ERD.docx")
MAIN = Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx")
NEW = Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx.new")

AUTHOR = "Cayunda, Klint Russel R.; Delos Santos, Brent R.; Fin, Noli III M."
CAPTION = (
    "Table 3.1  Use Case Matrix of Pyx Tracker: "
    "Web-based Enterprise Operations Management System."
)

# Sample pattern: Primary Actor (Admin) block then Secondary Actor (Staff) block
USE_CASES = [
    {
        "name": "Inventory stock movement",
        "primary": {
            "actor": "Admin/Executive",
            "trigger": "When Admin/Executive records a material receipt or finished-stock operation.",
            "flow": [
                "1. Select category and product.",
                "2. Enter or scan the movement.",
                "3. Review quantity and source.",
                "4. Confirm the allowed operation.",
            ],
            "post": "The relevant balance and movement record reflect the confirmed quantity.",
        },
        "secondary": {
            "actor": "Staff",
            "trigger": "When Staff enters or scans an allowed inventory movement.",
            "flow": [
                "1. Select category and product.",
                "2. Enter or scan the movement.",
                "3. Review quantity and source.",
                "4. Confirm the allowed operation.",
            ],
            "post": "Confirmed stock movement is recorded for the permitted operation.",
        },
    },
    {
        "name": "Product and parcel costing",
        "primary": {
            "actor": "Admin/Executive",
            "trigger": "When Admin/Executive reviews or updates product or parcel costs.",
            "flow": [
                "1. Select product or tier.",
                "2. Review quantities, prices and yield.",
                "3. Maintain formula lines.",
                "4. Review calculated cost and price.",
            ],
            "post": "The cost is reproducible from the stored inputs.",
        },
        "secondary": None,  # Admin/Executive only
    },
    {
        "name": "Order preparation and finalization",
        "primary": {
            "actor": "Admin/Executive",
            "trigger": "When Admin/Executive receives supported order files for preparation.",
            "flow": [
                "1. Choose preparation tool.",
                "2. Upload and review data.",
                "3. Run separate PDF/Excel stamper.",
                "4. Download results and finalize batch.",
            ],
            "post": "Prepared outputs and finalized order/parcel records are available.",
        },
        "secondary": {
            "actor": "Staff",
            "trigger": "When Staff prepares and finalizes a supported order batch.",
            "flow": [
                "1. Choose preparation tool.",
                "2. Upload and review data.",
                "3. Run separate PDF/Excel stamper.",
                "4. Download results and finalize batch.",
            ],
            "post": "Prepared outputs and finalized batch records are available.",
        },
    },
    {
        "name": "Before-shipment documentation",
        "primary": {
            "actor": "Admin/Executive",
            "trigger": "When a parcel is presented for before-shipment photo recording.",
            "flow": [
                "1. Scan or type tracking.",
                "2. Capture or upload photo.",
                "3. Review and save.",
                "4. Continue to next record.",
            ],
            "post": "Photo is retrievable by tracking number; stock is unchanged.",
        },
        "secondary": {
            "actor": "Staff",
            "trigger": "When Staff records a before-shipment photo for a parcel.",
            "flow": [
                "1. Scan or type tracking.",
                "2. Capture or upload photo.",
                "3. Review and save.",
                "4. Continue to next record.",
            ],
            "post": "Photo is stored against the tracking number without changing stock.",
        },
    },
    {
        "name": "RTS case and sellable stock-in",
        "primary": {
            "actor": "Admin/Executive",
            "trigger": "When a returned parcel is received for RTS case recording.",
            "flow": [
                "1. Record case and condition.",
                "2. Open stock-in session.",
                "3. Scan received sellable goods.",
                "4. Review and confirm session.",
            ],
            "post": "Case history is retained; confirmed sellable quantity is credited once.",
        },
        "secondary": {
            "actor": "Staff",
            "trigger": "When Staff records an RTS case and confirms sellable stock-in.",
            "flow": [
                "1. Record case and condition.",
                "2. Open stock-in session.",
                "3. Scan received sellable goods.",
                "4. Review and confirm session.",
            ],
            "post": "Confirmed sellable units are credited; discarded drafts leave inventory unchanged.",
        },
    },
]


def _tc(text: str, *, bold: bool = False, fill: str | None = None, white: bool = False, width: str = "2500"):
    tc = OxmlElement("w:tc")
    tcPr = OxmlElement("w:tcPr")
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), width)
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)
    tc.append(tcPr)
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")
    rPr.append(rf)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")
    rPr.append(sz)
    if bold:
        rPr.append(OxmlElement("w:b"))
    if white:
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "FFFFFF")
        rPr.append(color)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    p.append(r)
    tc.append(p)
    return tc


def _row(left: str, right: str, *, header: bool = False, alt: bool = False):
    tr = OxmlElement("w:tr")
    if header:
        fill, white = "2F80ED", True
        bold_l = bold_r = True
    else:
        fill = "D6EAF8" if alt else "FFFFFF"
        white = False
        bold_l, bold_r = True, False
    tr.append(_tc(left, bold=bold_l, fill=fill, white=white, width="2000"))
    tr.append(_tc(right, bold=bold_r, fill=fill, white=white, width="3000"))
    return tr


def add_actor_block(tbl, actor_label: str, actor_name: str, block: dict, alt: list):
    """actor_label is 'Primary Actor' or 'Secondary Actor'."""
    tbl.append(_row(actor_label, actor_name, alt=alt[0]))
    alt[0] = not alt[0]
    tbl.append(_row("Trigger", block["trigger"], alt=alt[0]))
    alt[0] = not alt[0]
    for i, step in enumerate(block["flow"]):
        left = "Basic Flow" if i == 0 else ""
        tbl.append(_row(left, step, alt=alt[0]))
        alt[0] = not alt[0]
    tbl.append(_row("Post Condition", block["post"], alt=alt[0]))
    alt[0] = not alt[0]


def build_table():
    tbl = OxmlElement("w:tbl")
    tblPr = OxmlElement("w:tblPr")
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), "5000")
    tblW.set(qn("w:type"), "pct")
    tblPr.append(tblW)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
    tbl.append(tblPr)
    grid = OxmlElement("w:tblGrid")
    for w in ("2000", "3000"):
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), w)
        grid.append(gc)
    tbl.append(grid)

    alt = [False]
    for uc in USE_CASES:
        # Sample header: Use Case | <Use case name>
        tbl.append(_row("Use Case", uc["name"], header=True))
        tbl.append(_row("Author", AUTHOR, alt=alt[0]))
        alt[0] = not alt[0]
        add_actor_block(tbl, "Primary Actor", uc["primary"]["actor"], uc["primary"], alt)
        if uc["secondary"]:
            add_actor_block(tbl, "Secondary Actor", uc["secondary"]["actor"], uc["secondary"], alt)
    return tbl


def is_uc_table(table) -> bool:
    blob = " ".join(c.text for r in table.rows[:8] for c in r.cells).lower()
    return ("uc01" in blob or "inventory stock" in blob) and (
        "author" in blob or "primary" in blob or "element" in blob
    )


def set_para(p, text: str) -> None:
    for r in p.runs:
        r.text = ""
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def make_caption_para():
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")
    rPr.append(rf)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rPr.append(sz)
    i = OxmlElement("w:i")
    rPr.append(i)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = CAPTION
    r.append(t)
    p.append(r)
    return p


def main() -> None:
    src = DESK if DESK.exists() else (NEW if NEW.exists() else MAIN)
    d = Document(str(src))

    for table in list(d.tables):
        if is_uc_table(table):
            table._tbl.getparent().remove(table._tbl)

    # Remove old Table 3.1 UC captions (we place caption under table like sample)
    for p in d.paragraphs:
        t = p.text.strip()
        if t.startswith("Table 3.1") and "Use Case Matrix" in t:
            set_para(p, "")
        if t.startswith("Table 3.2  Use Case") or t.startswith("Table 3.3  Use Case") or t.startswith(
            "Table 3.4  Use Case"
        ) or t.startswith("Table 3.5  Use Case"):
            set_para(p, "")

    new_tbl = build_table()
    caption = make_caption_para()
    placed = False
    for p in d.paragraphs:
        if p.text.strip() == "Use Case Matrix":
            # Insert table then caption after heading (sample has caption below table)
            p._element.addnext(new_tbl)
            new_tbl.addnext(caption)
            placed = True
            break
    if not placed:
        d.element.body.append(new_tbl)
        new_tbl.addnext(caption)

    out = DESK.parent / "_uc_coconut_style.docx"
    DESK.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out))

    v = Document(str(out))
    uc = [t for t in v.tables if is_uc_table(t)]
    print("UC tables:", len(uc), "rows", len(uc[0].rows) if uc else 0, "cols", len(uc[0].columns) if uc else 0)
    if uc:
        print("header:", uc[0].rows[0].cells[0].text, "|", uc[0].rows[0].cells[1].text)
        print("row1:", uc[0].rows[1].cells[0].text, "|", uc[0].rows[1].cells[1].text[:50])
        print("row2:", uc[0].rows[2].cells[0].text, "|", uc[0].rows[2].cells[1].text)
    caps = [p.text.strip() for p in v.paragraphs if "Use Case Matrix" in p.text]
    print("captions:", caps)

    for dest in (DESK, NEW, MAIN):
        try:
            shutil.copy2(out, dest)
            print("SAVED", dest, dest.stat().st_size)
        except Exception as e:
            print("LOCKED", dest, e)
            # fallback write bytes
            try:
                dest.write_bytes(out.read_bytes())
                print("SAVED write_bytes", dest)
            except Exception as e2:
                print("FAIL", dest, e2)
    out.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
