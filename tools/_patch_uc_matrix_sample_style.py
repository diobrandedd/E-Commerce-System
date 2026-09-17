"""Restyle Use Case Matrix as one 2-column table matching coconut sample look."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DESK = Path(r"C:/Users/Admin/Desktop/capstone/Pyx_Tracker_Capstone_SQL_ERD.docx")
MAIN = Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx")
NEW = Path(r"C:/xampp/htdocs/rts/output/documents/Pyx_Tracker_Capstone_SQL_ERD.docx.new")

AUTHOR = "Cayunda, Klint Russel R.; Delos Santos, Brent R.; Fin, Noli III M."

# Each block becomes: blue header Use Case | <name>, then body rows like sample
USE_CASES = [
    {
        "name": "UC01 Inventory stock movement",
        "rows": [
            ("Author", AUTHOR),
            ("Primary Actor", "Admin/Executive and Staff"),
            ("Trigger", "Material receipt or finished-stock operation"),
            ("Preconditions", "Valid product, stock unit, permission and source record."),
            (
                "Basic Flow",
                [
                    "1. Select category and product.",
                    "2. Enter or scan the movement.",
                    "3. Review quantity and source.",
                    "4. Confirm the allowed operation.",
                ],
            ),
            ("Post Condition", "The relevant balance and movement record reflect the confirmed quantity."),
            (
                "Exceptions",
                "Reject invalid input; preserve balances when an unconfirmed action is cancelled.",
            ),
        ],
    },
    {
        "name": "UC02 Product and parcel costing",
        "rows": [
            ("Author", AUTHOR),
            ("Primary Actor", "Admin/Executive"),
            ("Trigger", "Review or update of product costs"),
            ("Preconditions", "Configured material prices, units and formulas."),
            (
                "Basic Flow",
                [
                    "1. Select product or tier.",
                    "2. Review quantities, prices and yield.",
                    "3. Maintain formula lines.",
                    "4. Review calculated cost and price.",
                ],
            ),
            ("Post Condition", "The cost is reproducible from the stored inputs."),
            (
                "Exceptions",
                "Deny non-executive access; identify missing configuration before relying on the result.",
            ),
        ],
    },
    {
        "name": "UC03 Order preparation and finalization",
        "rows": [
            ("Author", AUTHOR),
            ("Primary Actor", "Admin/Executive and Staff"),
            ("Trigger", "Receipt of supported order files"),
            ("Preconditions", "Supported source files and an active processing batch."),
            (
                "Basic Flow",
                [
                    "1. Choose preparation tool.",
                    "2. Upload and review data.",
                    "3. Run separate PDF/Excel stamper.",
                    "4. Download results and finalize batch.",
                ],
            ),
            ("Post Condition", "Prepared outputs and finalized order/parcel records are available."),
            (
                "Exceptions",
                "Report duplicates, unmatched records or missing fields; do not describe partial preparation as finalization.",
            ),
        ],
    },
    {
        "name": "UC04 Before-shipment documentation",
        "rows": [
            ("Author", AUTHOR),
            ("Primary Actor", "Admin/Executive and Staff"),
            ("Trigger", "Parcel presented for photo recording"),
            ("Preconditions", "Readable or manually entered tracking identifier and an image."),
            (
                "Basic Flow",
                [
                    "1. Scan or type tracking.",
                    "2. Capture or upload photo.",
                    "3. Review and save.",
                    "4. Continue to next record.",
                ],
            ),
            ("Post Condition", "Photo is retrievable by tracking number; stock is unchanged."),
            (
                "Exceptions",
                "Provide duplicate feedback and request valid image/tracking input. A CSR parcel is not required.",
            ),
        ],
    },
    {
        "name": "UC05 RTS case and sellable stock-in",
        "rows": [
            ("Author", AUTHOR),
            ("Primary Actor", "Admin/Executive and Staff"),
            ("Trigger", "Returned parcel received"),
            ("Preconditions", "Tracking or case information, item condition and readable product codes."),
            (
                "Basic Flow",
                [
                    "1. Record case and condition.",
                    "2. Open stock-in session.",
                    "3. Scan received sellable goods.",
                    "4. Review and confirm session.",
                ],
            ),
            ("Post Condition", "Case history is retained; confirmed sellable quantity is credited once."),
            (
                "Exceptions",
                "Case save alone adds no sellable stock; discarding a draft leaves inventory unchanged.",
            ),
        ],
    },
]


def _tc(text: str, *, bold: bool = False, fill: str | None = None, white: bool = False, width: str = "2500"):
    tc = OxmlElement("w:tc")
    tcPr = OxmlElement("w:tcPr")
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), width)
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)
    tc.append(tcPr)

    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")
    rPr.append(rf)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")  # 10pt
    rPr.append(sz)
    if bold:
        rPr.append(OxmlElement("w:b"))
    if white:
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "FFFFFF")
        rPr.append(color)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    p.append(r)
    tc.append(p)
    return tc


def _row(left: str, right: str, *, header: bool = False, alt: bool = False):
    tr = OxmlElement("w:tr")
    if header:
        fill, white, bold_l, bold_r = "2F80ED", True, True, True
    else:
        fill = "D6EAF8" if alt else "FFFFFF"
        white, bold_l, bold_r = False, True, False
    tr.append(_tc(left, bold=bold_l, fill=fill, white=white, width="2000"))
    tr.append(_tc(right, bold=bold_r, fill=fill, white=white, width="3000"))
    return tr


def build_table():
    tbl = OxmlElement("w:tbl")
    tblPr = OxmlElement("w:tblPr")
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), "5000")
    tblW.set(qn("w:type"), "pct")
    tblPr.append(tblW)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
    tbl.append(tblPr)

    grid = OxmlElement("w:tblGrid")
    for w in ("2000", "3000"):
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), w)
        grid.append(gc)
    tbl.append(grid)

    alt = False
    for uc in USE_CASES:
        # Blue header like sample: Use Case | <UC name>
        tbl.append(_row("Use Case", uc["name"], header=True))
        for label, value in uc["rows"]:
            if isinstance(value, list):
                # First basic-flow row keeps label; continuation rows blank left cell (sample style)
                for i, step in enumerate(value):
                    left = label if i == 0 else ""
                    tbl.append(_row(left, step, alt=alt))
                    alt = not alt
            else:
                tbl.append(_row(label, value, alt=alt))
                alt = not alt
    return tbl


def is_uc_table(table) -> bool:
    blob = " ".join(c.text for r in table.rows[:5] for c in r.cells).lower()
    if "uc01" in blob and ("inventory stock" in blob or "author" in blob):
        return True
    h0 = table.rows[0].cells[0].text.strip().lower()
    if h0 in ("use case", "element") and ("uc01" in blob or "pyx tracker" in blob or "primary" in blob):
        return True
    return False


def set_para(p, text: str) -> None:
    for r in p.runs:
        r.text = ""
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def main() -> None:
    src = next(p for p in (DESK, NEW, MAIN) if p.exists())
    d = Document(str(src))

    removed = 0
    for table in list(d.tables):
        if is_uc_table(table):
            table._tbl.getparent().remove(table._tbl)
            removed += 1
    print("removed", removed)

    new_tbl = build_table()
    placed = False
    for p in d.paragraphs:
        if p.text.strip() == "Use Case Matrix":
            target = p._element
            probe = target.getnext()
            while probe is not None:
                if probe.tag == qn("w:p"):
                    joined = "".join((t.text or "") for t in probe.findall(".//" + qn("w:t"))).strip()
                    if joined.startswith("Table 3.1"):
                        target = probe
                        break
                    if joined.startswith("Storyboard"):
                        break
                probe = probe.getnext()
            target.addnext(new_tbl)
            placed = True
            break
    if not placed:
        d.element.body.append(new_tbl)
    print("placed", placed)

    # One caption only; remove 3.2-3.5 UC captions if any remain
    uc_caption_done = False
    for p in d.paragraphs:
        t = p.text.strip()
        if t.startswith("Table 3.1") and ("Use Case Matrix" in t or "Inventory stock" in t):
            if not uc_caption_done:
                set_para(
                    p,
                    "Table 3.1  Use Case Matrix of Pyx Tracker: Web-based Enterprise Operations Management System.",
                )
                uc_caption_done = True
            elif "Data Dictionary" not in t:
                # duplicate wrong caption
                pass
        if t.startswith("Table 3.2  Use Case Matrix") or t.startswith("Table 3.3  Use Case Matrix") or t.startswith(
            "Table 3.4  Use Case Matrix"
        ) or t.startswith("Table 3.5  Use Case Matrix"):
            set_para(p, "")

    out = DESK.parent / "_uc_sample_style.docx"
    DESK.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out))

    v = Document(str(out))
    uc = [t for t in v.tables if is_uc_table(t)]
    print("UC tables:", len(uc))
    if uc:
        t = uc[0]
        print("rows x cols:", len(t.rows), "x", len(t.columns))
        print("first header:", t.rows[0].cells[0].text, "|", t.rows[0].cells[1].text)
        headers = [r.cells[1].text for r in t.rows if r.cells[0].text.strip() == "Use Case"]
        print("UC headers in one table:", headers)

    for dest in (DESK, NEW, MAIN):
        try:
            shutil.copy2(out, dest)
            print("saved", dest, dest.stat().st_size)
        except Exception as e:
            print("LOCKED", dest, e)
    out.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
