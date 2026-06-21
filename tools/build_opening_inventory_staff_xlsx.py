#!/usr/bin/env python3
"""
Build staff-friendly opening inventory workbook (Option B).
- Staff_* sheets: plain English for printing and data entry
- _Import_* sheets: machine columns for import (hidden in Excel)
"""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "docs/opening_inventory/export_data.json"
OUT = ROOT / "docs/opening_inventory" / f"PyxTracker_Opening_Inventory_{date.today().isoformat()}.xlsx"

GRAY = PatternFill("solid", fgColor="D9D9D9")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREEN_HDR = PatternFill("solid", fgColor="548235")
WHITE_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=12, color="1F4E79")
TITLE_FILL = PatternFill("solid", fgColor="D9E1F2")
BDR = Border(
    left=Side(style="thin", color="999999"),
    right=Side(style="thin", color="999999"),
    top=Side(style="thin", color="999999"),
    bottom=Side(style="thin", color="999999"),
)

ROLE_FRIENDLY = {
    "executive": "Executive",
    "prod_manager": "Production manager",
    "tl": "Team leader",
    "staff": "Staff",
    "hr": "HR",
    "pos": "POS terminal",
}
ROLE_SYSTEM = {v: k for k, v in ROLE_FRIENDLY.items()}

DEPT_FRIENDLY = {
    "honey": "Honey",
    "paragis": "Paragis Tea",
    "cacao_brc": "Cacao / Black Rice / Coffee",
    "csr": "CSR",
}
DEPT_SYSTEM = {v: k for k, v in DEPT_FRIENDLY.items()}

SUPPLIER_CAT_FRIENDLY = {
    "honey": "Honey supplier",
    "paragis": "Paragis supplier",
    "black_rice": "Black rice supplier",
    "cacao": "Cacao supplier",
    "honey_comb": "Honey comb supplier",
}
SUPPLIER_CAT_SYSTEM = {v: k for k, v in SUPPLIER_CAT_FRIENDLY.items()}

HONEY_VARIANT_FRIENDLY = {"light": "Light", "amber": "Amber", "dark": "Dark"}
HONEY_VARIANT_SYSTEM = {v: k for k, v in HONEY_VARIANT_FRIENDLY.items()}

DEPT_PRODUCT_FRIENDLY = {
    "pt": "Paragis Tea (leaf bag)",
    "brc": "Black Rice Coffee",
    "jar_cacao": "Jar Cacao",
}

CATEGORY_FRIENDLY = {
    "raw_material": "Raw material",
    "packaging": "Packaging",
    "supply": "Supplies",
    "costing_only": "Costing only",
}


def friendly_role(system: str) -> str:
    return ROLE_FRIENDLY.get(system or "", system or "")


def friendly_dept(system: str) -> str:
    if not system:
        return ""
    return DEPT_FRIENDLY.get(system, system)


def header_row(ws, row: int, titles: list[str], fill=GREEN_HDR) -> None:
    for c, title in enumerate(titles, 1):
        cell = ws.cell(row, c, title)
        cell.fill = fill
        cell.font = WHITE_FONT
        cell.border = BDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 32


def data_row(ws, row: int, values: list, locked_cols: set[int]) -> None:
    for c, val in enumerate(values, 1):
        cell = ws.cell(row, c, val)
        cell.border = BDR
        cell.fill = GRAY if c in locked_cols else YELLOW
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def section_title(ws, row: int, text: str, max_col: int = 8) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_col)
    cell = ws.cell(row, 1, text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 28


def add_dropdown(ws, col: str, start: int, end: int, options: list[str]) -> None:
    joined = ",".join(options)
    dv = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}{start}:{col}{end}")


def autosize(ws, max_w: int = 44) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        w = min(max(max(len(str(c.value or "")) for c in col) + 2, 12), max_w)
        ws.column_dimensions[letter].width = w


def hide_sheet(ws) -> None:
    ws.sheet_state = "hidden"


def build_import_keys(ws, data: dict) -> None:
    ws.title = "_Import_Keys"
    header_row(ws, 1, ["map_type", "friendly_label", "system_value", "notes"])
    rows = [
        ("role", "Executive", "executive", "SPA login role"),
        ("role", "Production manager", "prod_manager", ""),
        ("role", "Team leader", "tl", ""),
        ("role", "Staff", "staff", ""),
        ("role", "HR", "hr", ""),
        ("department", "Honey", "honey", ""),
        ("department", "Paragis Tea", "paragis", ""),
        ("department", "Cacao / Black Rice / Coffee", "cacao_brc", ""),
        ("department", "CSR", "csr", ""),
        ("honey_variant", "Light", "light", ""),
        ("honey_variant", "Amber", "amber", ""),
        ("honey_variant", "Dark", "dark", ""),
        ("dept_product", "Paragis Tea (leaf bag)", "pt", "Supplier dept price"),
        ("dept_product", "Black Rice Coffee", "brc", ""),
        ("dept_product", "Jar Cacao", "jar_cacao", ""),
    ]
    r = 2
    for row in rows:
        data_row(ws, r, list(row), {1, 2, 3, 4})
        r += 1
    section_title(ws, r + 1, "Stock ID lookup (match by item name + department)", 4)
    header_row(ws, r + 2, ["stock_id", "item_name", "department_system", "category_system"])
    r = r + 3
    for s in data["stocks"]:
        data_row(ws, r, [s[0], s[1], s[2], s[3]], {1, 2, 3, 4})
        r += 1
    autosize(ws, 50)
    hide_sheet(ws)


def build_import_users_spa(ws, data: dict) -> None:
    ws.title = "_Import_Users_SPA"
    header_row(ws, 1, ["action", "user_id", "username", "full_name", "employee_id", "role", "department", "is_active", "notes"])
    r = 2
    for u in data["users_spa"]:
        uid, user, name, eid, role, dept, active = u
        if role == "pos":
            continue
        data_row(ws, r, ["keep", uid, user, name, eid, role, dept or "", active, ""], {2, 3, 4, 5, 6, 7, 8})
        r += 1
    for _ in range(10):
        data_row(ws, r, ["add", "", "", "", "", "", "", "1", ""], set())
        r += 1
    autosize(ws)
    hide_sheet(ws)


def build_staff_users(ws, data: dict) -> None:
    ws.title = "Staff - Users"
    section_title(ws, 1, "PYX TRACKER LOGINS - Fill yellow columns. Employee ID becomes their first password.", 7)
    headers = [
        "What to do",
        "Login username",
        "Full name",
        "Employee ID #",
        "Job role",
        "Department",
        "Notes",
    ]
    header_row(ws, 2, headers)
    r = 3
    locked = {2, 3, 4, 5, 6}
    for u in data["users_spa"]:
        uid, user, name, eid, role, dept, active = u
        if role == "pos":
            continue
        data_row(
            ws,
            r,
            ["Keep this person", user, name, eid, friendly_role(role), friendly_dept(dept), ""],
            locked,
        )
        r += 1
    for _ in range(15):
        data_row(ws, r, ["Add new person", "", "", "", "", "", ""], set())
        r += 1
    add_dropdown(ws, "A", 3, r - 1, ["Add new person", "Keep this person", "Remove person"])
    add_dropdown(ws, "E", 3, r - 1, list(ROLE_FRIENDLY.values())[:-1])  # no POS on SPA sheet
    add_dropdown(ws, "F", 3, r - 1, list(DEPT_FRIENDLY.values()))
    autosize(ws)


def build_staff_delivery(ws, data: dict) -> None:
    ws.title = "Staff - Delivery"
    section_title(ws, 1, "DELIVERY PHOTO APP — Separate from Pyx Tracker; usually role = Delivery", 6)
    header_row(ws, 2, ["What to do", "Login username", "Full name", "Employee ID #", "Role", "Active?", "Notes"])
    r = 3
    locked = {2, 3, 4, 5, 6}
    for u in data["app_users"]:
        aid, user, name, eid, role, status = u
        active = "Yes" if status == "active" else "No"
        role_f = "Delivery" if role == "delivery" else "Executive"
        data_row(ws, r, ["Keep this person", user, name, eid, role_f, active, ""], locked)
        r += 1
    for _ in range(5):
        data_row(ws, r, ["Add new person", "", "", "", "Delivery", "Yes", ""], set())
        r += 1
    add_dropdown(ws, "A", 3, r - 1, ["Add new person", "Keep this person", "Remove person"])
    add_dropdown(ws, "E", 3, r - 1, ["Delivery", "Executive"])
    add_dropdown(ws, "F", 3, r - 1, ["Yes", "No"])
    autosize(ws)


def build_staff_suppliers(ws) -> None:
    ws.title = "Staff - Suppliers"
    section_title(ws, 1, "SUPPLIER LIST — Make up a short code (example H001) and use the same code in the price tables below", 8)
    header_row(
        ws,
        2,
        [
            "Supplier code (you choose)",
            "First name",
            "Last name",
            "Mobile number",
            "Payment (GCash/Bank/etc.)",
            "Account number",
            "Address",
            "Supplier type",
            "Notes",
        ],
    )
    r = 3
    for _ in range(12):
        data_row(ws, r, [""] * 9, set())
        r += 1
    add_dropdown(ws, "H", 3, r - 1, list(SUPPLIER_CAT_FRIENDLY.values()))

    section_title(ws, r + 1, "HONEY PRICES — Price per gallon (US gallon) for each honey type", 6)
    header_row(ws, r + 2, ["Supplier code", "Honey type", "Price per gallon (PHP)", "Price starts on (YYYY-MM-DD)", "Notes"])
    r = r + 3
    for _ in range(15):
        data_row(ws, r, [""] * 5, set())
        r += 1
    add_dropdown(ws, "B", r - 15, r - 1, list(HONEY_VARIANT_FRIENDLY.values()))

    section_title(ws, r + 1, "RAW MATERIAL PRICES — PT / BRC / Jar Cacao price per unit", 6)
    header_row(ws, r + 2, ["Supplier code", "Product", "Price per unit (PHP)", "Price starts on (YYYY-MM-DD)", "Notes"])
    r = r + 3
    for _ in range(12):
        data_row(ws, r, [""] * 5, set())
        r += 1
    add_dropdown(ws, "B", r - 12, r - 1, list(DEPT_PRODUCT_FRIENDLY.values()))
    autosize(ws)


def build_staff_stock(ws, data: dict) -> None:
    ws.title = "Staff - Stock Count"
    section_title(
        ws,
        1,
        "STOCK COUNT — Write how many you have NOW. Do not change item names (gray). Fill yellow only.",
        9,
    )
    headers = [
        "Item name (do not change)",
        "Department (do not change)",
        "Type (do not change)",
        "Unit (do not change)",
        "How many on hand?",
        "Cost per unit (PHP)",
        "Supplier code (if any)",
        "Batch / lot # (optional)",
        "Date counted (YYYY-MM-DD)",
        "Notes",
    ]
    header_row(ws, 2, headers)
    r = 3
    locked = {1, 2, 3, 4}
    for s in data["stocks"]:
        sid, name, dept, cat, unit, sup_code = s
        data_row(
            ws,
            r,
            [name, friendly_dept(dept), CATEGORY_FRIENDLY.get(cat, cat), unit, "", "", sup_code, "", "", ""],
            locked,
        )
        r += 1
    autosize(ws, 48)


def build_staff_finished(ws, data: dict, inv: dict) -> None:
    ws.title = "Staff - Finished Goods"
    section_title(ws, 1, "FINISHED PRODUCTS ON SHELF — Units ready to sell + selling price", 6)
    header_row(
        ws,
        2,
        ["Product (do not change)", "Department (do not change)", "How many on shelf?", "Selling price (PHP)", "Notes"],
    )
    r = 3
    locked = {1, 2}
    for p in data["dept_products"]:
        pid, pname, dept, pkey = p
        qty = inv.get(pkey, "")
        data_row(ws, r, [pname, friendly_dept(dept), qty, "", ""], locked)
        r += 1

    section_title(ws, r + 1, "CSR PACKING STAGING (optional) — Materials issued to packing bench", 6)
    header_row(ws, r + 2, ["Item (do not change)", "Issued qty", "Used / consumed qty", "Damaged qty", "Notes"])
    r = r + 3
    locked = {1}
    for f in data["fulfillment"]:
        data_row(ws, r, [f[1], "", "", "", ""], locked)
        r += 1

    section_title(ws, r + 1, "FINISHED HONEY BOTTLES — Count bottles by color and size", 5)
    header_row(ws, r + 2, ["Honey color", "Bottle size", "How many bottles?", "Notes"])
    r = r + 3
    locked = set()
    for p in data["honey_products"]:
        pid, variant, size = p
        data_row(ws, r, [HONEY_VARIANT_FRIENDLY.get(variant, variant), size, "", ""], locked)
        r += 1
    add_dropdown(ws, "A", r - len(data["honey_products"]), r - 1, list(HONEY_VARIANT_FRIENDLY.values()))
    add_dropdown(ws, "B", r - len(data["honey_products"]), r - 1, ["250ml", "350ml", "liter", "gallon"])
    autosize(ws)


def build_readme(ws) -> None:
    ws.title = "START HERE"
    cut = date.today().isoformat()
    lines = [
        ("Pyx Tracker - Opening inventory (staff forms)", ""),
        ("", ""),
        ("Who fills what", ""),
        ("Staff - Users", "HR / executive — Pyx Tracker logins"),
        ("Staff - Delivery", "Delivery team — photo app logins"),
        ("Staff - Suppliers", "Purchasing — suppliers + prices"),
        ("Staff - Stock Count", "Each dept — count materials & packaging"),
        ("Staff - Finished Goods", "Production — finished products + honey bottles"),
        ("", ""),
        ("Count date", cut),
        ("Month (YYYY-MM)", date.today().strftime("%Y-%m")),
        ("", ""),
        ("Colors", ""),
        ("Gray", "Already filled — do not change"),
        ("Yellow", "YOU fill in here"),
        ("Blue title rows", "Instructions only — do not write here"),
        ("", ""),
        ("Workflow", "Print → count on paper → type into yellow cells → send file to IT"),
        ("", ""),
        ("Hidden sheets", "Sheets starting with _Import_ are for IT import only — ignore them"),
    ]
    for i, (a, b) in enumerate(lines, 1):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
        if i == 1:
            ws.cell(i, 1).font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 62


def main() -> None:
    data = json.loads(DATA.read_text(encoding="utf-8"))
    inv = data.get("dept_inv", {})
    wb = Workbook()
    build_readme(wb.active)

    build_staff_users(wb.create_sheet(), data)
    build_staff_delivery(wb.create_sheet(), data)
    build_staff_suppliers(wb.create_sheet())
    build_staff_stock(wb.create_sheet(), data)
    build_staff_finished(wb.create_sheet(), data, inv)

    build_import_keys(wb.create_sheet(), data)
    build_import_users_spa(wb.create_sheet(), data)

    wb.save(OUT)
    print(f"Wrote {OUT}")
    print("Staff sheets:", [s for s in wb.sheetnames if s.startswith("Staff")])
    print("Hidden:", [s for s in wb.sheetnames if s.startswith("_Import")])


if __name__ == "__main__":
    main()
