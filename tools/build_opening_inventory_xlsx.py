import json
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(r"C:/xampp/htdocs/rts")
DATA = ROOT / "docs/opening_inventory/export_data.json"
OUT = ROOT / "docs/opening_inventory" / f"PyxTracker_Opening_Inventory_{date.today().isoformat()}.xlsx"
GRAY = PatternFill("solid", fgColor="D9D9D9")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
HDR = PatternFill("solid", fgColor="4472C4")
HFONT = Font(bold=True, color="FFFFFF")
BDR = Border(left=Side(style="thin", color="999999"), right=Side(style="thin", color="999999"), top=Side(style="thin", color="999999"), bottom=Side(style="thin", color="999999"))

def hdr(ws, titles):
    for c, t in enumerate(titles, 1):
        cell = ws.cell(1, c, t)
        cell.fill, cell.font, cell.border = HDR, HFONT, BDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 28

def wrow(ws, r, vals, prefill):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v)
        cell.border = BDR
        cell.fill = GRAY if c in prefill else YELLOW
        cell.alignment = Alignment(vertical="center", wrap_text=True)

def autosize(ws, m=44):
    for col in ws.columns:
        w = min(max(max(len(str(x.value or "")) for x in col) + 2, 10), m)
        ws.column_dimensions[get_column_letter(col[0].column)].width = w

def dv(ws, col, s, e, opts):
    d = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
    ws.add_data_validation(d)
    d.add(f"{col}{s}:{col}{e}")

data = json.loads(DATA.read_text(encoding="utf-8"))
wb = Workbook()
cut = date.today().isoformat()
my = date.today().strftime("%Y-%m")

ws = wb.active
ws.title = "README"
readme = [
    ("Pyx Tracker Opening Inventory", ""),
    ("Cutover date", cut),
    ("Month YYYY-MM", my),
    ("", ""),
    ("Gray columns", "Pre-filled - do not change"),
    ("Yellow columns", "Fill in counts, prices, dates"),
    ("Workflow", "Print -> paper -> type in Excel -> send back"),
]
for i, (a, b) in enumerate(readme, 1):
    ws.cell(i, 1, a)
    ws.cell(i, 2, b)
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 60

ws = wb.create_sheet("Lookups")
hdr(ws, ["list_name", "value"])
lookups = [
    ("department", ["honey", "paragis", "cacao_brc", "csr"]),
    ("spa_role", ["executive", "prod_manager", "tl", "staff", "hr"]),
    ("user_action", ["keep", "add", "deactivate"]),
    ("supplier_category", ["honey", "paragis", "black_rice", "cacao", "honey_comb"]),
    ("honey_variant", ["light", "amber", "dark"]),
    ("dept_price_product_key", ["pt", "brc", "jar_cacao"]),
]
r = 2
for name, vals in lookups:
    for v in vals:
        ws.cell(r, 1, name)
        ws.cell(r, 2, v)
        r += 1
autosize(ws)

ws = wb.create_sheet("Users_SPA")
hdr(ws, ["action", "user_id", "username", "full_name", "employee_id", "role", "department", "is_active", "notes"])
pf = {2, 3, 4, 5, 6, 7, 8}
rn = 2
for u in data["users_spa"]:
    wrow(ws, rn, ["keep", *u, ""], pf)
    rn += 1
for _ in range(5):
    wrow(ws, rn, ["add", "", "", "", "", "", "", "1", ""], pf)
    rn += 1
dv(ws, "A", 2, rn - 1, "keep,add,deactivate")
dv(ws, "F", 2, rn - 1, "executive,prod_manager,tl,staff,hr")
dv(ws, "G", 2, rn - 1, "honey,paragis,cacao_brc,csr")
dv(ws, "H", 2, rn - 1, "1,0")
autosize(ws)

ws = wb.create_sheet("Users_Delivery")
hdr(ws, ["action", "app_user_id", "username", "full_name", "employee_id", "role", "status", "notes"])
pf = {2, 3, 4, 5, 6, 7}
rn = 2
for u in data["app_users"]:
    wrow(ws, rn, ["keep", *u, ""], pf)
    rn += 1
for _ in range(3):
    wrow(ws, rn, ["add", "", "", "", "", "delivery", "active", ""], pf)
    rn += 1
autosize(ws)

ws = wb.create_sheet("Suppliers_Master")
hdr(ws, ["supplier_code", "first_name", "last_name", "contact_number", "payment_account", "payment_account_number", "address", "category", "valid_id_filename", "is_active", "notes"])
rn = 2
for _ in range(15):
    wrow(ws, rn, [""] * 11, set())
    rn += 1
dv(ws, "H", 2, rn - 1, "honey,paragis,black_rice,cacao,honey_comb")
dv(ws, "J", 2, rn - 1, "1,0")
autosize(ws)

ws = wb.create_sheet("Supplier_Prices_Honey")
hdr(ws, ["supplier_code", "variant", "price_per_gallon", "effective_from", "notes"])
rn = 2
for _ in range(20):
    wrow(ws, rn, [""] * 5, set())
    rn += 1
dv(ws, "B", 2, rn - 1, "light,amber,dark")
autosize(ws)

ws = wb.create_sheet("Supplier_Prices_Dept")
hdr(ws, ["supplier_code", "product_key", "price_per_unit", "effective_from", "notes"])
rn = 2
for _ in range(15):
    wrow(ws, rn, [""] * 5, set())
    rn += 1
dv(ws, "B", 2, rn - 1, "pt,brc,jar_cacao")
autosize(ws)

ws = wb.create_sheet("Stock_OnHand")
hdr(ws, ["stock_id", "name", "department", "category", "unit", "on_hand", "current_price", "supplier_code", "supplier_batch_no", "batching_date", "notes"])
pf = {1, 2, 3, 4, 5}
rn = 2
for s in data["stocks"]:
    wrow(ws, rn, [s[0], s[1], s[2], s[3], s[4], "", "", s[5], "", "", ""], pf)
    rn += 1
autosize(ws, 48)

ws = wb.create_sheet("Honey_Finished_Simple")
hdr(ws, ["product_id", "variant", "size", "qty_on_hand", "notes"])
pf = {1, 2, 3}
rn = 2
for p in data["honey_products"]:
    wrow(ws, rn, [p[0], p[1], p[2], "", ""], pf)
    rn += 1
autosize(ws)

inv = data.get("dept_inv", {})
ws = wb.create_sheet("Dept_Finished_Products")
hdr(ws, ["product_id", "product_name", "department", "dept_product_key", "current_qty", "selling_price", "baseline_cost", "notes"])
pf = {1, 2, 3, 4}
rn = 2
for p in data["dept_products"]:
    qty = inv.get(p[3], "")
    wrow(ws, rn, [p[0], p[1], p[2], p[3], qty, "", "", ""], pf)
    rn += 1
autosize(ws)

ws = wb.create_sheet("Fulfillment_Staging")
hdr(ws, ["stock_id", "name", "issued_qty", "consumed_qty", "damaged_qty", "notes"])
pf = {1, 2}
rn = 2
for f in data["fulfillment"]:
    wrow(ws, rn, [f[0], f[1], "", "", "", ""], pf)
    rn += 1
autosize(ws)

wb.save(OUT)
print(f"Wrote {OUT}")
print("Sheets:", ", ".join(wb.sheetnames))
