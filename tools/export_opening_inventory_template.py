#!/usr/bin/env python3
"""Export Pyx Tracker opening-inventory Excel template from live rts_system DB."""
from __future__ import annotations
import subprocess, sys
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(r"C:\xampp\htdocs\rts")
OUT_DIR = ROOT / "docs" / "opening_inventory"
MYSQL = Path(r"C:\xampp\mysql\bin\mysql.exe")
DB = "rts_system"
FILL_PREFILLED = PatternFill("solid", fgColor="D9D9D9")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_HEADER = PatternFill("solid", fgColor="4472C4")
FONT_HEADER = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def mysql_query(sql: str) -> list[list[str]]:
    cmd = [str(MYSQL), "-u", "root", "--default-character-set=utf8mb4", "-N", "-B", DB, "-e", sql]
    proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.strip() or "mysql query failed")
    return [line.split("\t") for line in proc.stdout.strip().splitlines() if line.strip()]

def style_header_row(ws, headers, row=1):
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill, cell.font, cell.border = FILL_HEADER, FONT_HEADER, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 28

def write_row(ws, row, values, prefill_cols):
    for col, val in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.border = BORDER
        cell.fill = FILL_PREFILLED if col in prefill_cols else FILL_INPUT
        cell.alignment = Alignment(vertical="center", wrap_text=True)

def autosize(ws, max_width=42):
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        length = max(len(str(c.value or "")) for c in col_cells)
        ws.column_dimensions[letter].width = min(max(length + 2, 10), max_width)

def add_list_validation(ws, col_letter, start_row, end_row, options):
    dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{start_row}:{col_letter}{end_row}")

def build_workbook():
    wb = Workbook()
    cutover = date.today().isoformat()
    month_year = date.today().strftime("%Y-%m")
    ws = wb.active
    ws.title = "README"
    for r, row in enumerate([
        ["Pyx Tracker - Opening Inventory Template"], [""],
        ["Cutover date (YYYY-MM-DD):", cutover],
        ["Beginning inventory month (YYYY-MM):", month_year], [""],
        ["WORKFLOW"], ["1. Print sheets per team"], ["2. Floor staff write on paper (yellow columns only)"],
        ["3. Data-entry team types into Excel"], ["4. Send file back for import"], [""],
        ["Gray = pre-filled. Yellow = fill in."],
    ], 1):
        ws.cell(row=r, column=1, value=row[0])
        if len(row) > 1: ws.cell(row=r, column=2, value=row[1])
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 64

    ws = wb.create_sheet("Lookups")
    style_header_row(ws, ["list_name", "value"])
    lookup_rows = [
        ("department", ["honey","paragis","cacao_brc","csr"]),
        ("spa_role", ["executive","prod_manager","tl","staff","hr"]),
        ("user_action", ["keep","add","deactivate"]),
        ("supplier_category", ["honey","paragis","black_rice","cacao","honey_comb"]),
        ("honey_variant", ["light","amber","dark","other"]),
        ("honey_size", ["250ml","350ml","liter","gallon"]),
        ("dept_price_product_key", ["pt","brc","jar_cacao"]),
    ]
    r = 2
    for name, vals in lookup_rows:
        for v in vals:
            ws.cell(row=r, column=1, value=name)
            ws.cell(row=r, column=2, value=v)
            r += 1
    autosize(ws)

    ws = wb.create_sheet("Users_SPA")
    headers = ["action","user_id","username","full_name","employee_id","role","department","is_active","notes"]
    style_header_row(ws, headers)
    prefill = {2,3,4,5,6,7,8}
    row_num = 2
    for row in mysql_query("SELECT u.id,u.username,u.name,u.employee_id,r.name,IFNULL(d.name,''),u.is_active FROM users u JOIN roles r ON r.id=u.role_id LEFT JOIN departments d ON d.id=u.dept_id ORDER BY u.id"):
        write_row(ws, row_num, ["keep", *row, ""], prefill); row_num += 1
    for _ in range(5):
        write_row(ws, row_num, ["add","","","","","","1",""], prefill); row_num += 1
    add_list_validation(ws,"A",2,row_num-1,"keep,add,deactivate")
    add_list_validation(ws,"F",2,row_num-1,"executive,prod_manager,tl,staff,hr")
    add_list_validation(ws,"G",2,row_num-1,"honey,paragis,cacao_brc,csr")
    add_list_validation(ws,"H",2,row_num-1,"1,0")
    autosize(ws)

    ws = wb.create_sheet("Users_Delivery")
    style_header_row(ws, ["action","app_user_id","username","full_name","employee_id","role","status","notes"])
    prefill = {2,3,4,5,6,7}
    row_num = 2
    for row in mysql_query("SELECT id,username,full_name,IFNULL(employee_id,''),role,status FROM app_users ORDER BY id"):
        write_row(ws, row_num, ["keep", *row, ""], prefill); row_num += 1
    for _ in range(3):
        write_row(ws, row_num, ["add","","","","","delivery","active",""], prefill); row_num += 1
    autosize(ws)

    ws = wb.create_sheet("Suppliers_Master")
    style_header_row(ws, ["supplier_code","first_name","last_name","contact_number","payment_account","payment_account_number","address","category","valid_id_filename","is_active","notes"])
    row_num = 2
    for _ in range(15):
        write_row(ws, row_num, [""]*11, set()); row_num += 1
    add_list_validation(ws,"H",2,row_num-1,"honey,paragis,black_rice,cacao,honey_comb")
    add_list_validation(ws,"J",2,row_num-1,"1,0")
    autosize(ws)

    ws = wb.create_sheet("Supplier_Prices_Honey")
    style_header_row(ws, ["supplier_code","variant","price_per_gallon","effective_from","notes"])
    row_num = 2
    for _ in range(20):
        write_row(ws, row_num, ["","","","",""], set()); row_num += 1
    add_list_validation(ws,"B",2,row_num-1,"light,amber,dark")
    autosize(ws)

    ws = wb.create_sheet("Supplier_Prices_Dept")
    style_header_row(ws, ["supplier_code","product_key","price_per_unit","effective_from","notes"])
    row_num = 2
    for _ in range(15):
        write_row(ws, row_num, ["","","","",""], set()); row_num += 1
    add_list_validation(ws,"B",2,row_num-1,"pt,brc,jar_cacao")
    autosize(ws)

    ws = wb.create_sheet("Stock_OnHand")
    style_header_row(ws, ["stock_id","name","department","category","unit","on_hand","current_price","supplier_code","supplier_batch_no","batching_date","notes"])
    prefill = {1,2,3,4,5}
    row_num = 2
    for row in mysql_query("SELECT s.id,s.name,d.name,s.category,s.unit,s.on_hand,s.current_price,IFNULL(s.supplier_code,'''') FROM stocks s JOIN departments d ON d.id=s.dept_id ORDER BY d.name, FIELD(s.category,''raw_material'',''packaging'',''supply'',''costing_only''), s.name"):
        while len(row) < 8:
            row.append("")
        write_row(ws, row_num, [row[0],row[1],row[2],row[3],row[4],"","",row[7],"","",""], prefill); row_num += 1
    autosize(ws, 48)

    ws = wb.create_sheet("Honey_Finished_Simple")
    style_header_row(ws, ["product_id","variant","size","qty_on_hand","notes"])
    prefill = {1,2,3}
    row_num = 2
    for row in mysql_query("SELECT p.id,IFNULL(p.variant,''),IFNULL(p.size,'') FROM products p JOIN departments d ON d.id=p.dept_id WHERE d.name='honey' AND p.name='Honey' ORDER BY FIELD(p.variant,'light','amber','dark'), FIELD(p.size,'250ml','350ml','liter','gallon')"):
        write_row(ws, row_num, [row[0],row[1],row[2],"",""], prefill); row_num += 1
    autosize(ws)

    ws = wb.create_sheet("Dept_Finished_Products")
    style_header_row(ws, ["product_id","product_name","department","dept_product_key","current_qty","selling_price","baseline_cost","notes"])
    prefill = {1,2,3,4}
    inv_qty = {}
    for row in mysql_query("SELECT dir.product_key, dir.current_qty FROM dept_inventory_rows dir"):
        inv_qty[row[0]] = row[1]
    row_num = 2
    for row in mysql_query("SELECT p.id,p.name,d.name,IFNULL(p.dept_product_key,'') FROM products p JOIN departments d ON d.id=p.dept_id WHERE p.dept_product_key IN ('pt','brc','jar_cacao','honey_comb') OR p.name IN ('Paragis Tea','Black Rice Coffee','Jar Cacao','Honey Comb') ORDER BY d.name,p.name"):
        qty = inv_qty.get(row[3], "")
        write_row(ws, row_num, [row[0],row[1],row[2],row[3],qty,"","",""], prefill); row_num += 1
    autosize(ws)

    ws = wb.create_sheet("Fulfillment_Staging")
    style_header_row(ws, ["stock_id","name","issued_qty","consumed_qty","damaged_qty","notes"])
    prefill = {1,2}
    row_num = 2
    for row in mysql_query("SELECT fpb.stock_id,s.name,fpb.issued_qty,fpb.consumed_qty,fpb.damaged_qty FROM fulfillment_packing_balances fpb JOIN stocks s ON s.id=fpb.stock_id ORDER BY s.name"):
        write_row(ws, row_num, [row[0],row[1],"","","",""], prefill); row_num += 1
    autosize(ws)
    return wb

def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / f"PyxTracker_Opening_Inventory_{date.today().isoformat()}.xlsx"
    wb = build_workbook()
    wb.save(out)
    print(f"Wrote {out}")
    print("Sheets:", ", ".join(wb.sheetnames))

if __name__ == "__main__":
    main()

